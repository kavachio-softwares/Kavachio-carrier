"""Export engine — produce a BDX-shaped Excel workbook from canonical data.

Flow:
  1. User uploads a sample output BDX. `parse_template` reads every sheet,
     finds the header row, and extracts column headers + sample values.
  2. `propose_template_mapping` asks an LLM to map each excel column to a
     canonical data-model key, and to infer per-sheet `row_strategy`
     (which entity each row represents: policy / claim / coverage / etc.).
     Calls are made in parallel — one per sheet — to avoid output-token
     truncation and to reduce wall-clock time.
  3. The user reviews/corrects the structure via the API and approves it.
  4. `generate_workbook` takes assembled policy JSONs (same shape as /dwh)
     and writes an xlsx that matches the approved template — single sheet
     or multi-sheet, whatever the template defines.

Key accuracy fixes in this version:
  - Thinking disabled (thinking_budget=0): Gemini 2.5 Flash was spending
    ~15k "thinking" tokens before writing a single output token, leaving
    almost no room for the actual JSON. Disabling thinking frees the full
    output budget for candidates.
  - Field list filtered by table relevance: instead of dumping all 600+
    fields, only fields whose table matches the sheet's row_strategy (plus
    a small set of cross-cutting join fields) are sent. Smaller field list
    → model focuses correctly and produces better matches.
  - Compact field format: send "key|table.column|type|desc" as plain text
    lines instead of JSON objects, saving ~40% of prompt tokens.
  - Stronger heuristics: date-shaped columns (Effective Date, Expiry Date,
    Inception Date) now match the correct specific canonical key, not a
    generic one.
  - Explicit disambiguation examples in the prompt for common mistake pairs
    (program_name vs program_valid_from, policy_number vs claim_number, etc.)
  - Reduced chunk size to 30 (from 60): each chunk now fits comfortably in
    the output budget even with 5 candidates per column.
"""
from __future__ import annotations

import io
import json
import logging
import os
import re
import tempfile
import time
from copy import copy
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime
from typing import Any

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import MergedCell

from data_model import DATA_MODEL
from contract_upload_services.gemini_service import invoke_with_retry as _gateway_invoke

log = logging.getLogger("bdx.exporter")

# ---- Domain constants -------------------------------------------------------

SCALAR_TABLES = {
    "policy", "policyholder", "program", "contract", "tenant",
}

COLLECTION_TABLES = {
    "coverage", "premium_transaction", "risk_location",
    "claim", "claim_transaction", "claim_fee_line", "claim_reserve",
    "tax_line", "commission_line", "coverage_participation",
    "ingested_party", "party_license",
}

ROW_STRATEGY_TABLES = {
    "policy": None,
    "claim": "claim",
    "coverage": "coverage",
    "premium_transaction": "premium_transaction",
    "risk_location": "risk_location",
}

# Tables always included regardless of row_strategy (join keys live here).
_ALWAYS_INCLUDED_TABLES = {"policy", "policyholder", "program", "contract", "tenant"}

# row_strategy → primary tables to include in the field catalog sent to LLM.
_STRATEGY_TABLES: dict[str, set[str]] = {
    "policy":               {"policy", "policyholder", "program", "contract", "tenant",
                             "risk_location", "coverage", "ingested_party",
                             "party_license", "commission_line",
                             "premium_transaction"},
    "claim":                {"claim", "claim_transaction", "claim_fee_line",
                             "claim_reserve", "policy", "policyholder", "program",
                             "risk_location"},
    "coverage":             {"coverage", "coverage_participation", "policy",
                             "policyholder", "program", "premium_transaction",
                             "commission_line"},
    "premium_transaction":  {"premium_transaction", "policy", "policyholder",
                             "program", "coverage", "commission_line", "tax_line"},
    "risk_location":        {"risk_location", "policy", "policyholder", "program"},
}

MAX_SAMPLES = 5
_LLM_MAX_WORKERS = 4

# Output tokens per chunk. With thinking disabled each column+5 candidates
# costs ~120 tokens, so 30 cols × 120 = 3600 output tokens — well within 8k.
_LLM_MAX_OUTPUT_TOKENS = 8_192

TOP_N_CANDIDATES = 5

# Columns per Gemini call. Kept small so output never hits the token ceiling.
_CHUNK_SIZE = 30

_EXPORTABLE_FIELDS: set[str] = {
    k for k, v in DATA_MODEL.items() if v.get("source") != "system"
}


# ---- Template parsing -------------------------------------------------------

def hidden_sheet_names(file_bytes: bytes) -> set[str]:
    """Names of sheets flagged HIDDEN or VERY-HIDDEN in the workbook. These are
    working / scratch tabs the author deliberately hid — never a delivery or
    validation target — so we drop them from ALL processing (template parsing,
    input reading, sheet pickers). Best-effort: returns an empty set for a
    non-xlsx payload or on any read error."""
    try:
        wb = load_workbook(io.BytesIO(file_bytes), read_only=True)
        try:
            return {ws.title for ws in wb.worksheets if ws.sheet_state != "visible"}
        finally:
            wb.close()
    except Exception:
        return set()


def parse_template(file_bytes: bytes, filename: str | None = None) -> dict[str, Any]:
    """Read sample BDX → structural skeleton ready for LLM mapping."""
    sheets_out: list[dict[str, Any]] = []

    try:
        xl = pd.ExcelFile(io.BytesIO(file_bytes))
        hidden = hidden_sheet_names(file_bytes)
        raw_by_sheet: dict[str, "pd.DataFrame"] = {}
        for sheet_name in xl.sheet_names:
            if sheet_name in hidden:          # skip deliberately-hidden tabs
                continue
            raw = pd.read_excel(xl, sheet_name=sheet_name, header=None, dtype=object)
            if not raw.empty:
                raw_by_sheet[sheet_name] = raw

        # A workbook may include DATA-DICTIONARY / spec sheets (Field_Name +
        # Description columns) that document each data column's meaning and its
        # allowed values. Parse those first so we can (a) NOT treat them as data
        # sheets and (b) enrich the real data columns with description/allowed
        # values/type/required.
        specs: dict[str, dict] = {}
        for sheet_name, raw in raw_by_sheet.items():
            spec = _parse_spec_sheet(raw)
            if spec:
                specs[sheet_name] = spec

        for sheet_name, raw in raw_by_sheet.items():
            if sheet_name in specs:
                continue        # dictionary sheet — not a validation target
            header_row = _detect_header_row(raw)
            raw_hdrs = raw.iloc[header_row].tolist()
            valid_idx = [
                i for i, v in enumerate(raw_hdrs)
                if pd.notna(v) and str(v).strip()
                and not _UNNAMED_PAT.match(str(v).strip())
            ]
            headers = [str(raw_hdrs[i]).strip() for i in valid_idx]
            data = raw.iloc[header_row + 1:, valid_idx].reset_index(drop=True)
            columns = _build_columns(headers, data, valid_idx)
            if columns:
                sheets_out.append({
                    "sheet_name": sheet_name,
                    "header_row": header_row,
                    "data_start_row": header_row + 1,
                    "columns": columns,
                    "row_strategy": "policy",
                })

        if specs:
            _enrich_columns_from_specs(sheets_out, specs)
        return {"sheets": sheets_out}
    except Exception:
        log.debug("Not an Excel workbook — trying CSV/XML")

    try:
        text = file_bytes.decode("utf-8", errors="replace")
        raw = pd.read_csv(io.StringIO(text), header=None, dtype=object)
        header_row = _detect_header_row(raw)
        raw_hdrs = raw.iloc[header_row].tolist()
        valid_idx = [
            i for i, v in enumerate(raw_hdrs)
            if pd.notna(v) and str(v).strip()
            and not _UNNAMED_PAT.match(str(v).strip())
        ]
        headers = [str(raw_hdrs[i]).strip() for i in valid_idx]
        data = raw.iloc[header_row + 1:, valid_idx].reset_index(drop=True)
        columns = _build_columns(headers, data, valid_idx)
        sheets_out.append({
            "sheet_name": filename or "Sheet1",
            "header_row": header_row,
            "data_start_row": header_row + 1,
            "columns": columns,
            "row_strategy": "policy",
        })
        return {"sheets": sheets_out}
    except Exception:
        log.debug("Not a CSV template")

    try:
        df = pd.read_xml(io.BytesIO(file_bytes))
        headers = [str(c) for c in df.columns]
        data = df.reset_index(drop=True)
        columns = _build_columns(headers, data)
        sheets_out.append({
            "sheet_name": filename or "Sheet1",
            "header_row": 0,
            "data_start_row": 1,
            "columns": columns,
            "row_strategy": "policy",
        })
        return {"sheets": sheets_out}
    except Exception:
        log.debug("Not an XML template")

    return {"sheets": []}


# ---- Data-dictionary / spec sheets -----------------------------------------
# Some BDX workbooks ship a spec sheet per data sheet (e.g. "POL" describing the
# columns in "POL Data") with a Field_Name + Description layout. The Description
# documents each column's MEANING and, for coded fields, its ALLOWED VALUES
# ("1 - Annual / 2 - Semi-Annual / …"). We parse these so the mapping can pick a
# field by its meaning (not just a lexical name match) and value_in_set rules can
# validate against the documented codes.

def _dict_norm(s) -> str:
    """Tight normalization (drop ALL non-alphanumerics) for header matching."""
    return re.sub(r"[^a-z0-9]+", "", str(s).lower())


def _join_norm(s) -> str:
    """Space-collapsed normalization for joining a spec Field_Name to a data
    column name — mirrors output_schema._norm_field."""
    return re.sub(r"[^a-z0-9]+", " ", str(s).lower()).strip()


# "<CODE> - <label>" — capture BOTH the code AND its label, because a BDX cell
# may hold either form (InstallmentPlan "1" OR "Annual"; PolicyType "RB" OR
# "Renewal Business"). The label is kept only when it's a plausible cell VALUE
# (short), not a long descriptive sentence (e.g. "Cancel by carrier - loss
# experience").
_ENUM_LINE = re.compile(r"^\s*([A-Za-z0-9]{1,8})\s*[-–—]\s+(\S.*?)\s*$")
_X_OR_Y = re.compile(r"^\s*([A-Za-z0-9]{1,4})\s+or\s+([A-Za-z0-9]{1,4})\s*$", re.I)
_MAX_LABEL_LEN = 30


def _parse_allowed_values(desc: str) -> list[str]:
    """Extract the documented allowed VALUES a BDX cell could hold from a free-text
    Description — BOTH the code AND its short label. Conservative: only returns a
    list when the text clearly enumerates options."""
    if not desc:
        return []
    text = str(desc)
    # 1) "<CODE> - <label>" lines (the most common enum form).
    pairs: list[tuple[str, str]] = []
    for ln in text.splitlines():
        m = _ENUM_LINE.match(ln.strip())
        if m:
            pairs.append((m.group(1), m.group(2)))
    if len(pairs) >= 2:
        seen: set[str] = set()
        out: list[str] = []
        for code, label in pairs:
            for v in (code, label):
                # include the code always; include the label only if it looks like
                # a value (short), not a sentence-length description
                if v is label and len(label) > _MAX_LABEL_LEN:
                    continue
                if v and v not in seen:
                    seen.add(v)
                    out.append(v)
        return out
    # 2) "Y or N" style binary.
    m = _X_OR_Y.match(text)
    if m:
        return [m.group(1), m.group(2)]
    # 3) colon-introduced comma list: "Enter GL Segment: A, B, C, D".
    m = re.search(r":\s*(.+)$", text.replace("\n", " "))
    if m and m.group(1).count(",") >= 2:
        items = [t.strip() for t in m.group(1).split(",")
                 if t.strip() and len(t.strip()) <= 60]
        if len(items) >= 2:
            return items
    return []


def _parse_spec_sheet(raw: "pd.DataFrame") -> dict | None:
    """If `raw` is a data-dictionary sheet (has Field_Name + Description columns),
    return {join_norm(field) -> {field_name, description, allowed_values,
    field_format, required}}. Otherwise None."""
    header_row = _detect_header_row(raw)
    hdrs = [str(v).strip() if pd.notna(v) else "" for v in raw.iloc[header_row].tolist()]
    norm = [_dict_norm(h) for h in hdrs]

    def find(cands: set[str]):
        for i, n in enumerate(norm):
            if n in cands:
                return i
        return None

    fn_i = find({"fieldname", "fieldnames"})
    de_i = find({"description", "desc", "descriptions"})
    if fn_i is None or de_i is None:
        return None
    fmt_i = find({"format", "type", "datatype"})
    req_i = find({"required", "req"})
    if req_i is None:    # Aurenity header: "(O)ptional, (C)onditional, (R)equired"
        for i, h in enumerate(hdrs):
            hl = h.lower()
            if "required" in hl and "conditional" in hl:
                req_i = i
                break

    def cell(row, i):
        if i is None or i >= len(row):
            return None
        v = row[i]
        return str(v).strip() if pd.notna(v) and str(v).strip() else None

    spec: dict = {}
    for r in range(header_row + 1, len(raw)):
        row = raw.iloc[r].tolist()
        fname = cell(row, fn_i)
        if not fname or _dict_norm(fname) in {"primarykeyvalues", "notes"}:
            continue
        key = _join_norm(fname)
        if not key:
            continue
        desc = cell(row, de_i)
        spec[key] = {
            "field_name": fname,
            "description": desc,
            "allowed_values": _parse_allowed_values(desc or ""),
            "field_format": cell(row, fmt_i),
            "required": cell(row, req_i),
        }
    return spec or None


def spec_sheet_names(file_bytes: bytes) -> set[str]:
    """Names of the data-dictionary / spec sheets in a workbook (sheets whose
    header carries Field_Name + Description columns). These document the real
    data columns rather than being data themselves, so they are not mapping
    targets. Returns an empty set for non-Excel payloads."""
    names: set[str] = set()
    try:
        xl = pd.ExcelFile(io.BytesIO(file_bytes))
        for sheet_name in xl.sheet_names:
            raw = pd.read_excel(xl, sheet_name=sheet_name, header=None, dtype=object)
            if not raw.empty and _parse_spec_sheet(raw):
                names.add(str(sheet_name))
    except Exception:
        log.debug("spec_sheet_names: not an Excel workbook")
    return names


def _enrich_columns_from_specs(sheets_out: list[dict], specs: dict[str, dict]) -> None:
    """Attach description / allowed_values / type / required from the dictionary
    sheets onto the matching data columns (joined by normalized field name)."""
    merged: dict[str, dict] = {}
    for spec in specs.values():
        for k, info in spec.items():
            merged.setdefault(k, info)
    for sheet in sheets_out:
        for col in sheet.get("columns", []):
            info = merged.get(_join_norm(col.get("column_name") or ""))
            if not info:
                continue
            if info.get("description"):
                col["description"] = info["description"]
            if info.get("allowed_values"):
                col["allowed_values"] = list(info["allowed_values"])
            if info.get("field_format"):
                col["field_format"] = info["field_format"]
            if info.get("required"):
                col["required"] = info["required"]


def _build_columns(
    headers: list[str],
    data: "pd.DataFrame",
    col_indices: list[int] | None = None,
) -> list[dict]:
    """Build column descriptors for the parsed template.

    `col_indices` carries each kept column's ORIGINAL spreadsheet position.
    Blank / NaN / Unnamed columns are dropped from the mapping, but the
    surviving columns keep their real index so `column_index` still addresses
    the correct physical cell when the workbook is regenerated at export time.
    Falls back to sequential indices when omitted (e.g. flat XML input).
    """
    columns = []
    for i, name in enumerate(headers):
        col_samples: list[str] = []
        if i < data.shape[1]:
            col_samples = [
                str(v)
                for v in data.iloc[:, i].dropna().head(MAX_SAMPLES).tolist()
            ]
        columns.append({
            "column_index": col_indices[i] if col_indices is not None else i,
            "column_name": name,
            "samples": col_samples,
            "canonical_field": None,
            "confidence": 0.0,
            "candidates": [],
            "transform": None,
            "static_value": None,
            # Filled from a data-dictionary / spec sheet when present.
            "description": None,
            "allowed_values": [],
            "field_format": None,
            "required": None,
        })
    return columns


_UNNAMED_PAT = re.compile(r"^Unnamed:\s*\d+$")


def _detect_header_row(raw: "pd.DataFrame", max_scan: int = 20) -> int:
    """Two-pass: prefer all-text rows; fall back to highest text count."""
    best_idx, best_score = 0, -1
    fall_idx, fall_score = 0, -1
    for i, row in raw.head(max_scan).iterrows():
        non_empty = [v for v in row if pd.notna(v) and str(v).strip()]
        text_count = sum(
            1 for v in non_empty
            if not str(v).strip().lstrip("-").replace(".", "", 1).isnumeric()
        )
        num_count = len(non_empty) - text_count
        if text_count > fall_score:
            fall_score, fall_idx = text_count, int(i)
        if num_count == 0 and text_count > best_score:
            best_score, best_idx = text_count, int(i)
    return best_idx if best_score > 0 else fall_idx


# ---- Heuristic pre-pass -----------------------------------------------------

_YYYYMMDD = re.compile(r"^\d{8}$")
_YYYY_MM  = re.compile(r"^\d{6}$")
_ISO_DATE = re.compile(r"^\d{4}-\d{2}-\d{2}")
_US_STATE = re.compile(r"^[A-Z]{2}$")
# Matches: comma-grouped (1,000,000.00), plain decimal (1000000.00) AND plain
# integer (1000000) — the last alternative is essential, otherwise un-grouped
# integer amounts read as "text" and numeric-gated heuristics never fire.
_CURRENCY = re.compile(r"^-?\d{1,3}(?:[,]\d{3})*(?:\.\d+)?$|^-?\d+(?:\.\d+)?$")


def _sample_kind(samples: list[str]) -> str:
    cleaned = [s.strip() for s in samples if s and str(s).strip()]
    if not cleaned:
        return "empty"
    if all(_YYYYMMDD.match(s) for s in cleaned):
        return "date_yyyymmdd"
    if all(_YYYY_MM.match(s) for s in cleaned):
        return "yyyymm"
    if all(_ISO_DATE.match(s) for s in cleaned):
        return "iso_date"
    if all(_US_STATE.match(s) for s in cleaned):
        return "state_code"
    if all(_CURRENCY.match(s.replace("$", "")) for s in cleaned):
        return "numeric"
    if all(s.upper() in ("Y", "N", "YES", "NO", "TRUE", "FALSE") for s in cleaned):
        return "boolean"
    return "text"


# Each rule: (regex on lowercased column name, required sample kind or None,
#             canonical_field).
# ORDER MATTERS — first match wins. Put specific patterns before generic ones.
_HEURISTIC_RULES: list[tuple[re.Pattern, str | None, str]] = [
    # NOTE: targets must be REAL DATA_MODEL keys — _heuristic_pin silently
    # ignores any key not in _EXPORTABLE_FIELDS. Order matters: first match
    # wins, so specific/prefixed patterns precede generic ones.
    # --- Policy dates (prefixed) ---------------------------------------------
    (re.compile(r"policy.?effective|pol.?eff.?date|policy.?eff\b|policy.?inception"),
     None, "policy_effective_date"),
    (re.compile(r"policy.?expir|pol.?exp.?date|policy.?exp\b|policy.?end"),
     None, "policy_expiration_date"),
    # --- Program dates — MUST come before generic "program" rule below -------
    (re.compile(r"program.?effective|program.?start|program.?from|program.?eff|program.?inception"),
     None, "program_valid_from"),
    (re.compile(r"program.?expir|program.?end|program.?to|program.?exp"),
     None, "program_valid_until"),
    # --- Contract dates -------------------------------------------------------
    (re.compile(r"contract.?effective|contract.?from|contract.?start|contract.?inception"),
     None, "contract_inception_date"),
    (re.compile(r"contract.?expir|contract.?to|contract.?end"),
     None, "contract_expiry_date"),
    # --- Identifiers ----------------------------------------------------------
    (re.compile(r"\bumr\b|unique.?market.?ref"),
     None, "policy_umr"),
    (re.compile(r"certificate.?(no|number|ref)\b|cert.?no\b|\bcert\b"),
     None, "policy_certificate_reference"),
    (re.compile(r"claim.?number|claim.?no\b|clm.?no\b"),
     None, "claim_number"),
    (re.compile(r"policy.?(number|ref|reference|no|id)\b|pol.?no\b|polno\b"),
     None, "policy_number"),
    # --- Insured name (specific — avoid 'Insured State/City/Zip' columns) ----
    (re.compile(r"insured.?name|named.?insured|name.?of.?insured|policyholder|insured.?legal"),
     None, "policyholder_legal_name"),
    # --- Premium / financial (before the carrier-name rule, so 'Net Premium
    #     to Carrier' is treated as a premium, not a carrier name) ------------
    (re.compile(r"net.?prem"),
     "numeric", "premium_transaction_net_premium_amount"),
    (re.compile(r"\btria\b|terror"),
     "numeric", "premium_transaction_terrorism_premium_amount"),
    # Specific premium phrasings only — NOT a bare "premium" (which would
    # wrongly swallow "Premium Tax", "Premium Fee", etc.).
    (re.compile(r"gross.?written.?prem|gwp\b|written.?prem|gross.?prem|annual.?prem|total.?prem"),
     "numeric", "premium_transaction_total_gross_written_premium_amount"),
    (re.compile(r"commission.?amount|comm.?amount|\bcommission\b"),
     "numeric", "premium_transaction_commission_amount"),
    # --- Carrier / issuing-company NAME (specific — never bare 'carrier') ----
    (re.compile(r"issuing.?(company|carrier|entity)|writing.?company|carrier.?(entity|name|company)|\bcarrier\b.*\b(name|entity|company)\b"),
     None, "ingested_party_legal_name"),
    # --- Accounting period ----------------------------------------------------
    (re.compile(r"accounting.?(yr|period|yrmo)"),
     "yyyymm", "premium_transaction_accounting_date"),
    # --- Transaction ----------------------------------------------------------
    (re.compile(r"transaction.?type|trans.?type|tran.?type"),
     None, "premium_transaction_type"),
    (re.compile(r"transaction.?date|trans.?date|tran.?date"),
     None, "premium_transaction_effective_date"),
    # --- Limits ---------------------------------------------------------------
    (re.compile(r"occurrence.?limit|per.?occurrence|each.?occurrence|occ.?limit"),
     "numeric", "coverage_occurrence_limit"),
    (re.compile(r"aggregate.?limit|general.?aggregate|agg.?limit"),
     "numeric", "coverage_aggregate_limit"),
    # --- Coverage / class / line of business ---------------------------------
    (re.compile(r"line.?of.?business|\blob\b|class.?of.?business|risk.?class|coverage.?type|cvg.?type|\bclass\b"),
     None, "coverage_code"),
    # --- Location -------------------------------------------------------------
    (re.compile(r"state.?code|risk.?state|ins.?state|domicile.?state|\bstate\b"),
     None, "risk_location_subdivision"),
    (re.compile(r"zip.?code|postal.?code|zip\b"),
     None, "risk_location_postal_code"),
    (re.compile(r"\bcity\b"),
     None, "risk_location_city"),
    (re.compile(r"\bcountry\b"),
     None, "risk_location_country"),
    # --- Generic (unprefixed) policy dates — LAST among dates so prefixed
    #     program/contract/coverage rules above win first --------------------
    (re.compile(r"\beffective.?date\b|\binception.?date\b|\beffective\b"),
     None, "policy_effective_date"),
    (re.compile(r"\bexpir|\bexpiry\b|\bexpiration\b"),
     None, "policy_expiration_date"),
    # --- Program name — LAST so date patterns above take priority ------------
    (re.compile(r"program.?name|program\b"),
     None, "program_name"),
]


def _heuristic_pin(column_name: str, samples: list[str]) -> str | None:
    """Return a canonical key if the column is mechanically obvious, else None."""
    h = column_name.lower()
    kind = _sample_kind([str(s) for s in samples])
    for pattern, required_kind, canonical in _HEURISTIC_RULES:
        if required_kind and kind != required_kind:
            continue
        if pattern.search(h):
            if canonical in _EXPORTABLE_FIELDS:
                return canonical
    return None


# ---- LLM field catalog ------------------------------------------------------

def _fields_for_strategy(row_strategy: str) -> list[dict[str, str]]:
    """Return only the canonical fields relevant to this sheet's row_strategy.

    Sending a filtered subset (~50-150 fields) instead of all 600+ fields:
      - Reduces prompt tokens significantly
      - Forces the model to focus on the right entity's fields
      - Prevents cross-entity confusion (program_name matching a date column
        just because both are in the same giant list)
    """
    allowed_tables = _STRATEGY_TABLES.get(row_strategy, set()) | _ALWAYS_INCLUDED_TABLES
    out = []
    for k, v in DATA_MODEL.items():
        if v.get("source") == "system":
            continue
        if v.get("table") not in allowed_tables:
            continue
        out.append({
            "key": k,
            "table": v.get("table", ""),
            "column": v.get("column", ""),
            "type": v.get("type", ""),
            "desc": (v.get("description") or "")[:120],
        })
    return out


def _fields_as_text(fields: list[dict[str, str]]) -> str:
    """Format fields as compact plain-text lines instead of JSON.

    Format: key  |  table.column  (type)  —  description
    This is ~40% smaller than JSON and easier for the model to scan.
    """
    lines = []
    for f in fields:
        table_col = f"{f['table']}.{f['column']}" if f["column"] else f["table"]
        type_hint = f"  ({f['type']})" if f["type"] else ""
        desc = f"  — {f['desc']}" if f["desc"] else ""
        lines.append(f"  {f['key']}  |  {table_col}{type_hint}{desc}")
    return "\n".join(lines)


def _build_sheet_prompt(
    sheet_name: str,
    row_strategy: str,
    columns: list[dict],
    fields: list[dict[str, str]],
    pinned: dict[int, str],
) -> str:
    """Build the LLM prompt for one chunk of columns on one sheet.

    Design decisions that directly improve accuracy:
    1. Thinking disabled — caller sets thinking_budget=0, so the model spends
       zero tokens reasoning and puts everything into JSON output.
    2. Field list is pre-filtered to the tables relevant for this row_strategy
       (~50-150 entries instead of 600+), so the model doesn't confuse fields
       from other entities.
    3. Fields are plain-text lines, not JSON — saves ~40% of prompt tokens.
    4. Explicit disambiguation examples for common mistake pairs.
    5. Pinned columns are declared upfront so the model can skip reasoning
       about them and focus budget on harder columns.
    6. Sample values are prominently formatted so the model uses them.
    """
    # Build column block as plain readable text
    col_lines = []
    for c in columns:
        samples_str = ", ".join(repr(s) for s in (c.get("samples") or [])[:MAX_SAMPLES])
        pin_note = f"  ← PINNED={pinned[c['column_index']]}" if c["column_index"] in pinned else ""
        col_lines.append(
            f"  [{c['column_index']}] \"{c['column_name']}\""
            f"  samples=[{samples_str}]{pin_note}"
        )

    pins_section = ""
    if pinned:
        pin_lines = [f"  column_index {idx} → {canon}" for idx, canon in sorted(pinned.items())]
        pins_section = (
            "\nPINNED (heuristic, must be top candidate with s=1.0):\n"
            + "\n".join(pin_lines) + "\n"
        )

    disambiguation = """
CRITICAL DISAMBIGUATION — common mistakes to avoid:
- "Policy Effective Date" / "Inception Date"
    → policy_effective_date  NOT a program or contract field
- "Policy Expiry" / "Policy Expiration"
    → policy_expiration_date  NOT policy_effective_date
- "Contract Effective" → contract_inception_date
- "Program Name" / "Program" (text values like "GL 2024")
    → program_name  (only when samples are text names, not dates)
- "Commission" with numeric samples → premium_transaction_commission_amount
- "Insured Name" / "Named Insured" / "Insured / Named Insured"
    → policyholder_legal_name   (the POLICYHOLDER's name)
- "Issuing Company" / "Carrier" / "Carrier Entity" / "Writing Company"
    → ingested_party_legal_name  (an organisation read off the file;
      never policyholder_legal_name)
- "Class of Business" / "Risk Class" / "Line of Business" → coverage_code
- "Annual Premium" / "Written Premium" / "Gross Premium"
    → premium_transaction_total_gross_written_premium_amount
- "Net Premium" / "Net Premium to Carrier"
    → premium_transaction_net_premium_amount
- "Per Occurrence Limit" / "Occurrence Limit" → coverage_occurrence_limit
- "General Aggregate" / "Aggregate Limit" → coverage_aggregate_limit
- "Commission Rate" with % or decimal samples
    → premium_transaction_commission_percent
- A date column → look for a field whose name contains the same date concept
  (effective, expiry, inception, cancellation, etc.), NOT a name/text field
"""

    fields_text = _fields_as_text(fields)

    return (
        "You map OUTPUT Excel template columns to canonical insurance data model fields.\n"
        "Direction: output_column → which canonical field POPULATES it during export.\n\n"
        f"SHEET: \"{sheet_name}\"  |  ROW STRATEGY: {row_strategy}\n"
        f"(Each row represents one {row_strategy}. "
        f"Prefer fields from table matching '{row_strategy}' when ambiguous.)\n"
        + disambiguation
        + pins_section
        + f"\nFor EACH column below, return exactly {TOP_N_CANDIDATES} candidates "
        f"ranked by fit.\n"
        "Use BOTH the column name AND sample values. Never invent keys.\n"
        "If no field fits well, return the closest matches with low scores (s<0.30).\n\n"
        "Confidence: >=0.85 clear match | 0.65-0.85 strong | 0.45-0.65 likely "
        "| 0.20-0.45 weak | <0.20 no fit\n\n"
        "OUTPUT FORMAT — strict compact JSON, no prose, no markdown:\n"
        '{"sheet_name":"...","row_strategy":"policy","columns":['
        '{"column_index":0,"candidates":[{"c":"policy_number","s":0.95},'
        '{"c":"external_policy_number","s":0.40}]}'
        "]}\n\n"
        f"CANONICAL FIELDS for row_strategy={row_strategy}:\n"
        f"{fields_text}\n\n"
        f"COLUMNS TO MAP:\n"
        + "\n".join(col_lines)
        + "\n"
    )


# ---- JSON recovery ----------------------------------------------------------

def _lenient_json_loads(text: str, label: str = "") -> dict | None:
    """Multi-strategy JSON recovery for truncated / fence-wrapped responses."""
    if not text:
        return None
    text = text.strip()

    # 1. Strip fences and try direct parse.
    cleaned = re.sub(r"^```(?:json)?\s*", "", text).rstrip("`").strip()
    for candidate in (cleaned, text):
        try:
            return json.loads(candidate)
        except json.JSONDecodeError:
            pass

    # 2. Locate first '{' and try from there.
    start = text.find("{")
    if start < 0:
        return None
    text = text[start:]
    try:
        return json.loads(text)
    except json.JSONDecodeError:
        pass

    # 3. Find the longest balanced brace block (respects strings + escapes).
    depth = 0
    in_str = False
    escape = False
    last_balanced_end = -1
    for i, ch in enumerate(text):
        if escape:
            escape = False
            continue
        if ch == "\\":
            escape = True
            continue
        if ch == '"':
            in_str = not in_str
            continue
        if in_str:
            continue
        if ch == "{":
            depth += 1
        elif ch == "}":
            depth -= 1
            if depth == 0:
                last_balanced_end = i
    if last_balanced_end > 0:
        try:
            return json.loads(text[: last_balanced_end + 1])
        except json.JSONDecodeError:
            pass

    # 4. Tail-trim: cut after last closed column entry (']'), balance braces.
    #    Recovers partial responses where Gemini stopped mid-stream.
    cut = text.rfind("]")
    if cut > 0:
        trimmed = text[: cut + 1].rstrip().rstrip(",")
        depth = 0
        in_str = False
        escape = False
        for ch in trimmed:
            if escape:
                escape = False
                continue
            if ch == "\\":
                escape = True
                continue
            if ch == '"':
                in_str = not in_str
                continue
            if in_str:
                continue
            if ch == "{":
                depth += 1
            elif ch == "}":
                depth -= 1
        if depth > 0:
            trimmed += "}" * depth
        try:
            return json.loads(trimmed)
        except json.JSONDecodeError:
            pass

    # 5. Give up — save raw for debugging.
    try:
        tf = tempfile.NamedTemporaryFile(
            delete=False, prefix=f"exporter_{label}_", suffix=".txt", dir="/tmp"
        )
        tf.write(text.encode("utf-8", errors="ignore"))
        tf.close()
        log.error(
            "JSON recovery exhausted for '%s'; saved to %s. First 300: %r",
            label, tf.name, text[:300],
        )
    except Exception:
        log.exception("Could not save raw Gemini response for '%s'", label)
    return None


# ---- Candidate validation and application -----------------------------------

def _is_acceptable_canonical(cf: str) -> bool:
    return cf in _EXPORTABLE_FIELDS


def _parse_candidates_response(raw: dict, valid_col_indices: set[int]) -> list[dict]:
    """Extract and validate the columns list from a parsed LLM response."""
    out: list[dict] = []
    for cm in raw.get("columns") or []:
        idx = cm.get("column_index")
        if idx not in valid_col_indices:
            continue
        raw_cands = cm.get("candidates") or []
        cleaned: list[dict] = []
        seen: set[str] = set()
        for item in (raw_cands if isinstance(raw_cands, list) else []):
            if not isinstance(item, dict):
                continue
            cf = item.get("c") or item.get("canonical")
            if not isinstance(cf, str) or not _is_acceptable_canonical(cf) or cf in seen:
                continue
            seen.add(cf)
            try:
                conf = float(
                    item.get("s") if "s" in item else item.get("confidence", 0.0) or 0.0
                )
            except (TypeError, ValueError):
                conf = 0.0
            cleaned.append({
                "canonical": cf,
                "confidence": max(0.0, min(1.0, conf)),
                "reason": (item.get("reason") or "")[:140],
            })
            if len(cleaned) >= TOP_N_CANDIDATES:
                break
        if cleaned:
            cleaned.sort(key=lambda c: c["confidence"], reverse=True)
            out.append({"column_index": idx, "candidates": cleaned})
    return out


def _apply_parsed_mapping(parsed: dict, sheet: dict, chunk_columns: list[dict]) -> None:
    """Write validated LLM candidates back into `sheet` columns."""
    valid_indices = {c["column_index"] for c in chunk_columns}
    col_map = {c["column_index"]: c for c in sheet["columns"]}
    for result in _parse_candidates_response(parsed, valid_indices):
        idx = result["column_index"]
        if idx not in col_map:
            continue
        cands = result["candidates"]
        col_map[idx]["candidates"] = cands
        top = cands[0]
        col_map[idx]["canonical_field"] = top["canonical"]
        col_map[idx]["confidence"] = top["confidence"]


def _seed_heuristic_pins(sheet: dict) -> dict[int, str]:
    """Run heuristic pre-pass, pre-seed candidates, return pins dict."""
    pins: dict[int, str] = {}
    for c in sheet["columns"]:
        canonical = _heuristic_pin(c["column_name"], c.get("samples") or [])
        if not canonical:
            continue
        idx = c["column_index"]
        pins[idx] = canonical
        pinned_cand = {"canonical": canonical, "confidence": 1.0, "reason": "heuristic"}
        existing = [x for x in (c.get("candidates") or []) if x.get("canonical") != canonical]
        c["candidates"] = ([pinned_cand] + existing)[:TOP_N_CANDIDATES]
        c["canonical_field"] = canonical
        c["confidence"] = 1.0
    if pins:
        log.info("  Heuristic-pinned %d col(s) on '%s': %s",
                 len(pins), sheet["sheet_name"],
                 {v: k for k, v in pins.items()})
    return pins


# ---- Gemini call ------------------------------------------------------------

def _call_gemini_for_chunk(
    client,
    sheet: dict,
    chunk_columns: list[dict],
    chunk_index: int,
    pinned: dict[int, str],
) -> dict | None:
    """One Gemini call for a chunk of columns. Thinking disabled."""
    sheet_name = sheet["sheet_name"]
    row_strategy = sheet.get("row_strategy") or "policy"
    label = f"{sheet_name}_chunk{chunk_index}"

    # Field list filtered to this sheet's row_strategy — smaller and more focused.
    fields = _fields_for_strategy(row_strategy)
    prompt = _build_sheet_prompt(sheet_name, row_strategy, chunk_columns, fields, pinned)

    log.info("    Gemini ← '%s' chunk %d (%d cols, %d fields, %d pinned)",
             sheet_name, chunk_index, len(chunk_columns), len(fields),
             sum(1 for c in chunk_columns if c["column_index"] in pinned))

    t0 = time.time()

    # Route through the shared AI gateway (gemini_service) so retry/backoff and
    # the global concurrency + rate limiter live in ONE place. We keep exporter's
    # own client (600 s timeout) and its lenient JSON recovery below.
    kwargs = {
        "model": "gemini-2.5-flash",
        "contents": prompt,
        "config": {
            "response_mime_type": "application/json",
            "max_output_tokens": _LLM_MAX_OUTPUT_TOKENS,
            # KEY FIX: disable extended thinking so all output tokens go to the
            # JSON response instead of internal reasoning. Without this, ~15k
            # tokens are consumed "thinking" and the model runs out of budget.
            "thinking_config": {"thinking_budget": 0},
        },
    }
    try:
        resp = _gateway_invoke(kwargs, label=label, gen_client=client)
    except Exception as e:
        log.error("Gemini call failed (after gateway retries) for '%s': %s", label, e)
        return None

    finish_reason = "?"
    raw_cands_obj = getattr(resp, "candidates", None) or []
    if raw_cands_obj:
        finish_reason = getattr(raw_cands_obj[0], "finish_reason", "?")
    usage = getattr(resp, "usage_metadata", None)
    elapsed = time.time() - t0
    text = (resp.text or "").strip()
    log.info("    Gemini → '%s' chunk %d in %.2fs (%d chars, finish=%s, "
             "output_tokens=%s, think_tokens=%s)",
             sheet_name, chunk_index, elapsed, len(text), finish_reason,
             getattr(usage, "candidates_token_count", "?"),
             getattr(usage, "thoughts_token_count", "?"))

    if finish_reason not in ("?",) and str(finish_reason) == "FinishReason.MAX_TOKENS":
        log.warning("    '%s' chunk %d hit MAX_TOKENS — "
                    "reduce _CHUNK_SIZE if this recurs.", sheet_name, chunk_index)

    parsed = _lenient_json_loads(text, label)
    if parsed is None:
        log.error("Could not recover JSON for '%s' chunk %d (finish=%s)",
                  sheet_name, chunk_index, finish_reason)
    return parsed


def _map_sheet(client, sheet: dict) -> None:
    """Map all columns on one sheet. Mutates `sheet` in-place."""
    sheet_name = sheet["sheet_name"]
    all_cols = sheet["columns"]
    log.info("  Mapping sheet '%s' (%d columns)…", sheet_name, len(all_cols))

    pinned = _seed_heuristic_pins(sheet)
    chunks = [all_cols[i: i + _CHUNK_SIZE] for i in range(0, len(all_cols), _CHUNK_SIZE)]
    row_strategy_set = False

    for chunk_idx, chunk in enumerate(chunks):
        chunk_pins = {
            c["column_index"]: pinned[c["column_index"]]
            for c in chunk if c["column_index"] in pinned
        }
        parsed = _call_gemini_for_chunk(client, sheet, chunk, chunk_idx, chunk_pins)
        if parsed is None:
            continue

        if not row_strategy_set:
            rs = parsed.get("row_strategy")
            if rs in ROW_STRATEGY_TABLES:
                sheet["row_strategy"] = rs
                row_strategy_set = True

        _apply_parsed_mapping(parsed, sheet, chunk)

    # Re-enforce heuristic pins as top candidate (defensive against any
    # partial-parse that might have overwritten them).
    col_map = {c["column_index"]: c for c in all_cols}
    for idx, canonical in pinned.items():
        c = col_map[idx]
        existing = [x for x in (c.get("candidates") or []) if x.get("canonical") != canonical]
        c["candidates"] = (
            [{"canonical": canonical, "confidence": 1.0, "reason": "heuristic"}]
            + existing
        )[:TOP_N_CANDIDATES]
        c["canonical_field"] = canonical
        c["confidence"] = 1.0

    mapped = sum(1 for c in all_cols if c.get("canonical_field"))
    log.info("  Sheet '%s': %d/%d columns mapped (%.0f%%)",
             sheet_name, mapped, len(all_cols),
             100 * mapped / len(all_cols) if all_cols else 0)


# --- Sheet role classification (LLM, header-only) ---------------------------
# A BDX workbook usually mixes the actual transactional schedules (whose rows are
# what validation rules check) with REFERENCE / LOOKUP tabs that merely enumerate
# or map the values used on the data tabs (e.g. a small nickname→canonical-name
# table, an approved list, a code table). Rules must NEVER be generated against a
# reference tab. We let the model decide which tabs are reference data purely from
# the sheet names + column headers — deliberately WITHOUT hardcoding that any
# particular column implies "data" or "reference", so it generalises to any
# workbook and works even when a tab has no rows yet.

_SHEET_ROLE_SCHEMA = {
    "type": "object",
    "properties": {
        "sheets": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "sheet_name": {"type": "string"},
                    "role": {"type": "string", "enum": ["data", "reference"]},
                    "reason": {"type": "string"},
                },
                "required": ["sheet_name", "role"],
            },
        }
    },
    "required": ["sheets"],
}


def is_reference_sheet(sheet: dict) -> bool:
    """True when a template sheet was classified as a reference / lookup tab and is
    therefore NOT a rule target — no rules are generated for, displayed on, or
    compiled against its columns. See classify_sheet_roles. Central so every
    rule-field path (generation, extraction vocabulary, the setup picker) agrees."""
    return (sheet.get("sheet_role") == "reference"
            or sheet.get("rule_generatable") is False)


def _build_sheet_role_prompt(structure: dict) -> str:
    blocks = []
    for sh in structure.get("sheets", []):
        headers = [c.get("column_name") for c in sh.get("columns", [])
                   if c.get("column_name")]
        blocks.append(f'Sheet "{sh.get("sheet_name", "")}":\n  columns: {headers}')
    sheets_block = "\n\n".join(blocks)
    return f"""You are analysing the tabs of a single insurance bordereaux (BDX) workbook.

Some tabs hold the workbook's actual reportable DATA — the policies / transactions
/ cessions that downstream validation rules are run against. Other tabs are
REFERENCE (lookup) tabs: they do not hold the reportable transactions themselves,
they only list or map values that the data tabs draw on (for example a table that
maps a short name to its full/canonical form, an approved list, or a code table).
Reference tabs must be excluded from rule generation.

For EACH tab below, decide whether it is the workbook's own transactional DATA or a
REFERENCE / lookup tab. Reason from the sheet names, the column headers, and how the
tabs relate to one another. Do not assume any particular column name implies one
role or the other — judge each workbook on its own.

Tabs:
{sheets_block}

Return JSON of the form
  {{"sheets":[{{"sheet_name": "...", "role": "data" | "reference", "reason": "..."}}]}}
Classify every tab exactly once, echoing its exact sheet_name."""


def classify_sheet_roles(client, structure: dict) -> dict:
    """LLM-classify each sheet as 'data' or 'reference' from headers only, then
    annotate the structure in place: sets sheet['sheet_role'] +
    sheet['rule_generatable'] (reference tabs are NOT rule-generatable) and
    sheet['sheet_role_reason']. Non-fatal: on any error a sheet defaults to
    'data' (rule-generatable) so we never silently drop rules."""
    sheets = structure.get("sheets", [])
    if not sheets:
        return structure

    prompt = _build_sheet_role_prompt(structure)
    parsed = None
    for attempt in range(1, 3):
        try:
            resp = client.models.generate_content(
                model="gemini-2.5-flash",
                contents=prompt,
                config={
                    "response_mime_type": "application/json",
                    "response_schema": _SHEET_ROLE_SCHEMA,
                    "thinking_config": {"thinking_budget": 0},
                },
            )
            parsed = _lenient_json_loads((resp.text or "").strip(), "sheet-roles")
            if parsed is not None:
                break
        except Exception as e:
            log.warning("Sheet-role classification attempt %d/2 failed: %s", attempt, e)

    roles = {
        item.get("sheet_name"): (item.get("role"), item.get("reason"))
        for item in (parsed or {}).get("sheets", [])
        if item.get("sheet_name")
    }
    for sh in sheets:
        role, reason = roles.get(sh["sheet_name"], (None, None))
        is_ref = str(role).strip().lower() == "reference"
        sh["sheet_role"] = "reference" if is_ref else "data"
        sh["rule_generatable"] = not is_ref
        if reason:
            sh["sheet_role_reason"] = reason
        if is_ref:
            log.info("  Sheet '%s' → REFERENCE (excluded from rule generation): %s",
                     sh["sheet_name"], reason)
    return structure


def propose_template_mapping(structure: dict[str, Any]) -> dict[str, Any]:
    """Run LLM (per-sheet, in parallel), mutating `structure` in-place."""
    api_key = os.getenv("GEMINI_API_KEY")
    if not api_key:
        log.warning("No GEMINI_API_KEY — returning template with empty mappings.")
        return structure

    from google import genai
    from google.genai import types as genai_types

    client = genai.Client(
        api_key=api_key,
        http_options=genai_types.HttpOptions(timeout=600_000),  # 600 s
    )

    # Classify tabs (data vs reference/lookup) from headers so reference tabs are
    # kept out of rule generation. Header-only + non-fatal — see classify_sheet_roles.
    try:
        classify_sheet_roles(client, structure)
    except Exception:
        log.exception("Sheet-role classification failed; treating all tabs as data.")

    sheets = structure["sheets"]
    log.info(
        "Mapping %d sheet(s) via Gemini "
        "(≤%d workers | chunk=%d | top-%d | thinking=OFF)…",
        len(sheets), _LLM_MAX_WORKERS, _CHUNK_SIZE, TOP_N_CANDIDATES,
    )
    t_total = time.time()

    with ThreadPoolExecutor(max_workers=_LLM_MAX_WORKERS) as pool:
        futures = {
            pool.submit(_map_sheet, client, sheet): sheet["sheet_name"]
            for sheet in sheets
        }
        for future in as_completed(futures):
            sheet_name = futures[future]
            try:
                future.result()
                log.info("  ✓ '%s' done.", sheet_name)
            except Exception:
                log.exception("  ✗ '%s' raised an exception.", sheet_name)

    total_cols = sum(len(sh["columns"]) for sh in sheets)
    total_mapped = sum(
        sum(1 for c in sh["columns"] if c.get("canonical_field"))
        for sh in sheets
    )
    log.info(
        "All sheets mapped in %.2fs — %d/%d columns mapped overall (%.0f%%).",
        time.time() - t_total, total_mapped, total_cols,
        100 * total_mapped / total_cols if total_cols else 0,
    )
    return structure


# ---- Value resolution + workbook generation --------------------------------

def _resolve_value(
    policy: dict,
    canonical_field: str | None,
    row_table: str | None,
    row_item: dict | None,
) -> Any:
    if not canonical_field:
        return None

    from extras import is_extras_key, parse_xf_key

    if is_extras_key(canonical_field):
        entity, key = parse_xf_key(canonical_field)
        if entity == "policy":
            return ((policy.get("policy") or {}).get("extras") or {}).get(key)
        if row_table == entity and row_item is not None:
            return (row_item.get("extras") or {}).get(key)
        arr = policy.get(entity)
        if isinstance(arr, list) and arr:
            return (arr[0].get("extras") or {}).get(key)
        if isinstance(arr, dict):
            return (arr.get("extras") or {}).get(key)
        return None

    meta = DATA_MODEL.get(canonical_field)
    if not meta:
        return None
    table, column = meta.get("table"), meta.get("column")
    if not table or not column:
        return None

    if row_table and table == row_table and row_item is not None:
        return row_item.get(column)
    if table == "policy":
        return (policy.get("policy") or {}).get(column)
    if table in SCALAR_TABLES:
        return (policy.get(table) or {}).get(column)
    if table in COLLECTION_TABLES:
        arr = policy.get(table) or []
        if arr:
            return arr[0].get(column)
    return None


def _iter_rows(policy: dict, row_strategy: str):
    if row_strategy == "policy" or row_strategy not in ROW_STRATEGY_TABLES:
        yield (None, None)
        return
    table = ROW_STRATEGY_TABLES[row_strategy]
    items = policy.get(table) or []
    if not items:
        return
    for item in items:
        yield (table, item)


def _apply_transform(value: Any, transform: str | None) -> Any:
    if value is None:
        return None
    if not transform:
        return value
    t = transform.lower()
    try:
        if t in ("upper", "uppercase"):
            return str(value).upper()
        if t in ("lower", "lowercase"):
            return str(value).lower()
        if t in ("date", "yyyy-mm-dd"):
            return str(value)[:10]
        if t in ("number2dp", "money"):
            return round(float(value), 2)
        if t == "int":
            return int(float(value))
        if t == "str":
            return str(value)
    except Exception:
        return value
    return value


def _copy_cell_style(src_cell, dst_cell) -> None:
    if src_cell is None or src_cell is dst_cell:
        return
    try:
        dst_cell.font = copy(src_cell.font)
        dst_cell.fill = copy(src_cell.fill)
        dst_cell.border = copy(src_cell.border)
        dst_cell.alignment = copy(src_cell.alignment)
        dst_cell.number_format = src_cell.number_format
        dst_cell.protection = copy(src_cell.protection)
    except Exception as e:
        log.debug("copy style failed: %s", e)


def _cell_has_visible_style(cell) -> bool:
    if cell is None:
        return False
    fill = getattr(cell, "fill", None)
    if fill is not None:
        ft = getattr(fill, "fill_type", None) or getattr(fill, "patternType", None)
        if ft and ft != "none":
            return True
    return bool(getattr(cell, "has_style", False))


def _find_style_source(ws, header_row_1based: int, data_start_1based: int, col_idx: int):
    scan_end = min((ws.max_row or data_start_1based) + 1, data_start_1based + 10)
    for r in range(data_start_1based, scan_end):
        cell = ws.cell(row=r, column=col_idx)
        if _cell_has_visible_style(cell):
            return cell
    return ws.cell(row=header_row_1based, column=col_idx)


def _generate_with_template(
    structure: dict[str, Any],
    policies: list[dict],
    template_bytes: bytes,
) -> bytes:
    wb = load_workbook(io.BytesIO(template_bytes))
    existing_sheets = {ws.title: ws for ws in wb.worksheets}

    for sh in structure["sheets"]:
        cols = sorted(sh["columns"], key=lambda c: c["column_index"])
        sheet_name = sh.get("sheet_name") or ""
        if sheet_name in existing_sheets:
            ws = existing_sheets[sheet_name]
        else:
            ws = wb.create_sheet(title=sheet_name[:31] or "Sheet")
            for c in cols:
                ws.cell(row=1, column=c["column_index"] + 1,
                        value=c.get("column_name") or "")

        header_row_1b = (sh.get("header_row") or 0) + 1
        data_start = (sh.get("data_start_row") or 1) + 1

        row_merge_spans: list[tuple[int, int]] = []
        for mr in list(ws.merged_cells.ranges):
            if (mr.min_row == data_start and mr.max_row == data_start
                    and mr.max_col > mr.min_col):
                row_merge_spans.append((mr.min_col, mr.max_col))
                ws.unmerge_cells(str(mr))

        # Dissolve any OTHER merged ranges sitting inside the data region (e.g. a
        # stray "A8:F8" left over in the uploaded sample). The data writer can
        # only set the top-left anchor of a merge and SKIPS the rest as
        # MergedCells — which silently blanks those columns for that row. The
        # per-row pattern detected above is re-applied later; everything else in
        # the data area must be free cells so every value can be written.
        for mr in list(ws.merged_cells.ranges):
            if mr.min_row >= data_start:
                ws.unmerge_cells(str(mr))

        style_template: dict[int, Any] = {}
        for c in cols:
            col_idx = c["column_index"] + 1
            style_template[col_idx] = _find_style_source(
                ws, header_row_1b, data_start, col_idx,
            )
        sample_row_height = ws.row_dimensions[data_start].height

        max_row = ws.max_row or data_start
        for r in range(data_start, max_row + 1):
            for c in cols:
                cell = ws.cell(row=r, column=c["column_index"] + 1)
                if isinstance(cell, MergedCell):
                    continue
                cell.value = None

        out_row = data_start
        row_strategy = sh.get("row_strategy") or "policy"
        for policy in policies:
            for row_table, row_item in _iter_rows(policy, row_strategy):
                for c in cols:
                    col_idx = c["column_index"] + 1
                    if c.get("static_value") is not None:
                        val = c["static_value"]
                    else:
                        val = _resolve_value(
                            policy, c.get("canonical_field"),
                            row_table, row_item,
                        )
                        val = _apply_transform(val, c.get("transform"))
                    cell = ws.cell(row=out_row, column=col_idx)
                    if isinstance(cell, MergedCell):
                        continue
                    if val is not None:
                        cell.value = val
                    _copy_cell_style(style_template.get(col_idx), cell)
                if sample_row_height is not None:
                    ws.row_dimensions[out_row].height = sample_row_height
                for (min_col, max_col) in row_merge_spans:
                    ws.merge_cells(
                        start_row=out_row, end_row=out_row,
                        start_column=min_col, end_column=max_col,
                    )
                out_row += 1

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


def _normalize_datetime(data):
    """Recursively strip tzinfo so openpyxl can write datetime values."""
    if isinstance(data, dict):
        for key, value in data.items():
            if isinstance(value, datetime) and value.tzinfo is not None:
                data[key] = value.replace(tzinfo=None)
            elif isinstance(value, (dict, list)):
                _normalize_datetime(value)
    elif isinstance(data, list):
        for i, value in enumerate(data):
            if isinstance(value, datetime) and value.tzinfo is not None:
                data[i] = value.replace(tzinfo=None)
            elif isinstance(value, (dict, list)):
                _normalize_datetime(value)
    return data


def build_output_records(
    structure: dict[str, Any],
    policies: list[dict],
) -> list[dict]:
    """Resolve the exact post-transform values that WILL be written to the
    output, per sheet, without rendering any xlsx.

    Mirrors `generate_workbook`'s per-cell resolution (same `_iter_rows`,
    `_resolve_value`, `_apply_transform`, `static_value`) so validation can run
    on the real typed values BEFORE the workbook is created — avoiding a lossy
    "render xlsx then read it back" round-trip.

    Returns:
        [ {"sheet": <name>, "records": [ {column_name: value, ...}, ... ]}, ... ]
    """
    policies = _normalize_datetime(policies)
    out: list[dict] = []
    for sh in structure["sheets"]:
        cols = sorted(sh["columns"], key=lambda c: c["column_index"])
        records: list[dict] = []
        row_strategy = sh.get("row_strategy") or "policy"
        for policy in policies:
            for row_table, row_item in _iter_rows(policy, row_strategy):
                rec: dict[str, Any] = {}
                for c in cols:
                    name = c.get("column_name") or ""
                    if not name:
                        continue
                    if c.get("static_value") is not None:
                        val = c["static_value"]
                    else:
                        val = _resolve_value(
                            policy, c.get("canonical_field"),
                            row_table, row_item,
                        )
                        val = _apply_transform(val, c.get("transform"))
                    rec[name] = val
                records.append(rec)
        out.append({"sheet": sh.get("sheet_name", ""), "records": records})
    return out


def generate_workbook(
    structure: dict[str, Any],
    policies: list[dict],
    template_bytes: bytes | None = None,
) -> bytes:
    """Build an xlsx according to `structure` + assembled policies."""
    policies = _normalize_datetime(policies)

    if template_bytes:
        try:
            return _generate_with_template(structure, policies, template_bytes)
        except Exception as e:
            log.warning(
                "Style-preserving generation failed (%s); falling back to plain workbook.", e
            )

    wb = Workbook()
    wb.remove(wb.active)
    for sh in structure["sheets"]:
        ws = wb.create_sheet(title=(sh["sheet_name"] or "Sheet")[:31])
        cols = sorted(sh["columns"], key=lambda c: c["column_index"])
        for c in cols:
            ws.cell(row=1, column=c["column_index"] + 1,
                    value=c.get("column_name") or "")
        row_idx = 2
        row_strategy = sh.get("row_strategy") or "policy"
        for policy in policies:
            for row_table, row_item in _iter_rows(policy, row_strategy):
                for c in cols:
                    if c.get("static_value") is not None:
                        val = c["static_value"]
                    else:
                        val = _resolve_value(
                            policy, c.get("canonical_field"),
                            row_table, row_item,
                        )
                        val = _apply_transform(val, c.get("transform"))
                    if val is not None:
                        ws.cell(row=row_idx, column=c["column_index"] + 1, value=val)
                row_idx += 1

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


# Excel "Light Red Fill" (the same colour Excel uses for failed conditional
# formatting) so flagged cells read as errors at a glance.
_INVALID_FILL_RGB = "FFC7CE"


def highlight_exceptions(
    xlsx_bytes: bytes,
    structure: dict[str, Any],
    exceptions: list[dict],
) -> bytes:
    """Flag every cell that failed validation with a light-red fill + a comment.

    Used when the user chooses "Generate anyway" despite open validation
    exceptions: the produced workbook is re-opened and each offending cell is
    painted light red with a comment explaining why, so the problems are visible
    in the downloaded file.

    Cell positions mirror `_generate_with_template`'s row math exactly: per sheet,
    the i-th output record (1-based — which is what the validator stores in each
    exception's `row`/`__rowid`) is written at Excel row `data_start_row + 1 +
    (i - 1)`, and each output column lives at `column_index + 1`. Exceptions that
    don't resolve to a real cell (unknown sheet/column, missing row, merged-cell
    anchor) are skipped rather than guessed.

    Pure, defensive post-processing: it NEVER raises. On any failure it returns
    the input bytes unchanged so adding highlights can never break the download.
    """
    if not xlsx_bytes or not exceptions:
        return xlsx_bytes
    try:
        from openpyxl.styles import PatternFill
        from openpyxl.comments import Comment

        # Match sheet/column names tolerantly: collapse internal whitespace, trim,
        # lowercase. The validator's `sheet`/`field` are SQL string literals the
        # LLM copied from the schema, so a stray space or case difference vs the
        # template column name would otherwise silently skip the cell.
        def _norm(s) -> str:
            return re.sub(r"\s+", " ", str(s)).strip().lower()

        wb = load_workbook(io.BytesIO(xlsx_bytes))
        # normalized worksheet title -> real title (so wb[...] uses the real one)
        ws_by_norm = {_norm(t): t for t in wb.sheetnames}

        # normalized sheet_name -> (data_start_1based, {normalized_col: col_idx})
        sheet_meta: dict[str, tuple[int, dict[str, int]]] = {}
        for sh in structure.get("sheets") or []:
            name = sh.get("sheet_name") or ""
            data_start = (sh.get("data_start_row") or 1) + 1
            colmap: dict[str, int] = {}
            for c in sh.get("columns") or []:
                cn = c.get("column_name")
                if cn:
                    colmap[_norm(cn)] = c["column_index"] + 1
            sheet_meta[_norm(name)] = (data_start, colmap)

        red = PatternFill(start_color=_INVALID_FILL_RGB,
                          end_color=_INVALID_FILL_RGB, fill_type="solid")

        # Accumulate messages per cell so multiple violations on the same cell
        # merge into one comment (openpyxl allows only one comment per cell).
        notes: dict[tuple[str, int, int], list[str]] = {}

        # Diagnostics: why each exception was/ wasn't mapped to a cell.
        skip = {"no_sheet": 0, "no_worksheet": 0, "no_field": 0,
                "col_not_in_template": 0, "bad_row": 0, "merged_cell": 0}
        missed_fields: set[str] = set()

        for exc in exceptions:
            sheet = exc.get("sheet")
            field = exc.get("column") or exc.get("field")
            nsheet = _norm(sheet) if sheet else ""
            if not nsheet or nsheet not in sheet_meta:
                skip["no_sheet"] += 1
                continue
            ws_title = ws_by_norm.get(nsheet)
            if not ws_title:
                skip["no_worksheet"] += 1
                continue
            if not field:
                skip["no_field"] += 1
                continue
            data_start, colmap = sheet_meta[nsheet]
            col_idx = colmap.get(_norm(field))
            if not col_idx:
                skip["col_not_in_template"] += 1
                missed_fields.add(f"{sheet!r}:{field!r}")
                continue
            try:
                rowid = int(exc.get("row"))
            except (TypeError, ValueError):
                skip["bad_row"] += 1
                continue
            if rowid < 1:
                skip["bad_row"] += 1
                continue
            excel_row = data_start + (rowid - 1)
            cell = wb[ws_title].cell(row=excel_row, column=col_idx)
            if isinstance(cell, MergedCell):
                skip["merged_cell"] += 1
                continue
            cell.fill = red
            msg = (exc.get("message") or exc.get("reason")
                   or exc.get("rule_name") or "Validation exception")
            notes.setdefault((ws_title, excel_row, col_idx), []).append(str(msg))

        for (ws_title, excel_row, col_idx), msgs in notes.items():
            seen, lines = set(), []
            for m in msgs:
                if m not in seen:
                    seen.add(m)
                    lines.append(m)
            text = "\n".join(
                (f"• {m}" for m in lines) if len(lines) > 1 else lines
            )
            comment = Comment(text, "Kavachio validation")
            comment.width = 320
            comment.height = 22 + 14 * min(len(lines), 8)
            wb[ws_title].cell(row=excel_row, column=col_idx).comment = comment

        # One-line diagnostic so it's clear in the server log why cells did/didn't
        # get highlighted (sits alongside the existing "[DuckDB validation]" line).
        print(f"[highlight] exceptions={len(exceptions)} "
              f"highlighted_cells={len(notes)} skipped={skip}")
        if missed_fields:
            avail = {n: sorted(cm.keys()) for n, (_, cm) in sheet_meta.items()}
            print(f"[highlight] field/sheet not found in template (first 10): "
                  f"{sorted(missed_fields)[:10]} | template columns: {avail}")

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf.getvalue()
    except Exception as e:  # never let highlighting break the download
        log.warning("highlight_exceptions failed (%s); returning unhighlighted file.", e)
        return xlsx_bytes