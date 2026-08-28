"""Render direct-lane projected rows into an xlsx output file.

The existing exporter.generate_workbook resolves cells from canonical `policies`
via `canonical_field`; the direct lane has already produced finished output rows
keyed by output column name. This renderer writes those rows straight into the
output template, reusing exporter's style helpers so template fonts/borders are
preserved.

Also exposes `to_validation_blocks()` so the projected rows can be fed to the
existing DuckDB contract-rule validator (duckdb_validation.run_validation), which
expects the same shape exporter.build_output_records produces.
"""
from __future__ import annotations

import io
import re
from datetime import datetime
from typing import Any

_ISO_DATE_RE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
_ISO_DT_RE = re.compile(r"^\d{4}-\d{2}-\d{2}[T ]\d{2}:\d{2}")

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import MergedCell


def to_validation_blocks(
    structure: dict, output_rows_by_sheet: dict[str, list[dict]]
) -> list[dict]:
    """Shape projected rows like exporter.build_output_records:
    [{"sheet": name, "records": [{col: val}, ...]}, ...]."""
    blocks: list[dict] = []
    for sh in structure.get("sheets", []):
        name = sh.get("sheet_name", "")
        blocks.append({"sheet": name, "records": output_rows_by_sheet.get(name, [])})
    return blocks


def _coerce_cell(val: Any) -> Any:
    if isinstance(val, datetime) and val.tzinfo is not None:
        return val.replace(tzinfo=None)
    # ISO date/datetime strings (from the landing JSON) → real date objects so the
    # output cell is a genuine date (template date format applies) rather than text.
    if isinstance(val, str):
        s = val.strip()
        if _ISO_DATE_RE.match(s):
            try:
                return datetime.strptime(s, "%Y-%m-%d").date()
            except ValueError:
                return val
        if _ISO_DT_RE.match(s):
            try:
                return datetime.fromisoformat(s.replace(" ", "T")).replace(tzinfo=None)
            except ValueError:
                return val
    return val


def render_output(
    structure: dict,
    output_rows_by_sheet: dict[str, list[dict]],
    template_bytes: bytes | None = None,
) -> bytes:
    """Write the projected output rows into an xlsx. Uses the template (style
    preserving) when provided, else a plain workbook."""
    if template_bytes:
        try:
            return _render_with_template(structure, output_rows_by_sheet, template_bytes)
        except Exception:  # noqa: BLE001 — never fail the delivery on styling
            pass

    wb = Workbook()
    wb.remove(wb.active)
    for sh in structure.get("sheets", []):
        ws = wb.create_sheet(title=(sh.get("sheet_name") or "Sheet")[:31])
        cols = sorted(sh.get("columns", []), key=lambda c: c.get("column_index", 0))
        for c in cols:
            ws.cell(row=1, column=c.get("column_index", 0) + 1,
                    value=c.get("column_name") or "")
        rows = output_rows_by_sheet.get(sh.get("sheet_name", ""), [])
        for r, row in enumerate(rows, start=2):
            for c in cols:
                name = c.get("column_name")
                if not name:
                    continue
                val = _coerce_cell(row.get(name))
                if val is not None:
                    ws.cell(row=r, column=c.get("column_index", 0) + 1, value=val)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


def _render_with_template(
    structure: dict,
    output_rows_by_sheet: dict[str, list[dict]],
    template_bytes: bytes,
) -> bytes:
    from exporter import _copy_cell_style, _find_style_source

    wb = load_workbook(io.BytesIO(template_bytes))
    existing = {ws.title: ws for ws in wb.worksheets}

    for sh in structure.get("sheets", []):
        cols = sorted(sh.get("columns", []), key=lambda c: c.get("column_index", 0))
        sheet_name = sh.get("sheet_name") or ""
        if sheet_name in existing:
            ws = existing[sheet_name]
        else:
            ws = wb.create_sheet(title=sheet_name[:31] or "Sheet")
            for c in cols:
                ws.cell(row=1, column=c.get("column_index", 0) + 1,
                        value=c.get("column_name") or "")

        header_row_1b = (sh.get("header_row") or 0) + 1
        data_start = (sh.get("data_start_row") or 1) + 1

        # Free up merged ranges in the data area so every cell can be written.
        for mr in list(ws.merged_cells.ranges):
            if mr.min_row >= data_start:
                ws.unmerge_cells(str(mr))

        style_template = {
            c.get("column_index", 0) + 1: _find_style_source(
                ws, header_row_1b, data_start, c.get("column_index", 0) + 1)
            for c in cols
        }
        sample_row_height = ws.row_dimensions[data_start].height

        # Clear stale data rows.
        for r in range(data_start, (ws.max_row or data_start) + 1):
            for c in cols:
                cell = ws.cell(row=r, column=c.get("column_index", 0) + 1)
                if not isinstance(cell, MergedCell):
                    cell.value = None

        rows = output_rows_by_sheet.get(sheet_name, [])
        out_row = data_start
        for row in rows:
            for c in cols:
                col_idx = c.get("column_index", 0) + 1
                cell = ws.cell(row=out_row, column=col_idx)
                if isinstance(cell, MergedCell):
                    continue
                name = c.get("column_name")
                val = _coerce_cell(row.get(name)) if name else None
                if val is not None:
                    cell.value = val
                _copy_cell_style(style_template.get(col_idx), cell)
            if sample_row_height is not None:
                ws.row_dimensions[out_row].height = sample_row_height
            out_row += 1

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()
