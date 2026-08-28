"""
Builds the Kavachio Microservices Architecture documentation as a Word (.docx) file.
Written in simple English. Covers services, DuckDB placement, YAML files, and CI/CD.
"""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ---------- Theme colors ----------
BLUE = RGBColor(0x1F, 0x4E, 0x79)
LIGHT_BLUE = RGBColor(0x2E, 0x74, 0xB5)
GREY = RGBColor(0x59, 0x59, 0x59)
GREEN = RGBColor(0x2E, 0x7D, 0x32)
CODE_BG = "F2F2F2"

doc = Document()

# ---------- Base styles ----------
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.15


def shade_cell(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), color_hex)
    tcPr.append(shd)


def set_cell_text(cell, text, bold=False, color=None, size=10, white=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(2)
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.bold = bold
    if white:
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    elif color:
        run.font.color.rgb = color


def add_heading(text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    run.font.bold = True
    if level == 1:
        run.font.size = Pt(17)
        run.font.color.rgb = BLUE
    elif level == 2:
        run.font.size = Pt(14)
        run.font.color.rgb = LIGHT_BLUE
    else:
        run.font.size = Pt(12)
        run.font.color.rgb = GREY
    return p


def add_para(text, bold=False, italic=False, color=None, size=11):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.bold = bold
    run.font.italic = italic
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    return p


def add_bullet(text, level=0, bold_lead=None):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.25 + 0.25 * level)
    p.paragraph_format.space_after = Pt(3)
    if bold_lead:
        r = p.add_run(bold_lead)
        r.font.bold = True
        r.font.size = Pt(11)
        p.add_run(text).font.size = Pt(11)
    else:
        p.add_run(text).font.size = Pt(11)
    return p


def add_number(text, bold_lead=None):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(3)
    if bold_lead:
        r = p.add_run(bold_lead)
        r.font.bold = True
    p.add_run(text)
    return p


def add_code(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.15)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), CODE_BG)
    pPr.append(shd)
    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
    return p


def add_table(headers, rows, col_widths=None, header_color=None):
    header_color = header_color or "1F4E79"
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        set_cell_text(hdr[i], h, bold=True, white=True, size=10)
        shade_cell(hdr[i], header_color)
    for r_idx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, val in enumerate(row):
            set_cell_text(cells[i], val, size=9.5)
            if r_idx % 2 == 1:
                shade_cell(cells[i], "F7F9FB")
    if col_widths:
        for i, w in enumerate(col_widths):
            for cell in table.columns[i].cells:
                cell.width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_divider():
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "1F4E79")
    pbdr.append(bottom)
    pPr.append(pbdr)


# ============================================================
# TITLE PAGE
# ============================================================
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_before = Pt(120)
r = title.add_run("KAVACHIO")
r.font.size = Pt(40)
r.font.bold = True
r.font.color.rgb = BLUE

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run("Microservices Architecture")
r.font.size = Pt(26)
r.font.color.rgb = LIGHT_BLUE

sub2 = doc.add_paragraph()
sub2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub2.add_run("Complete Design & Migration Documentation")
r.font.size = Pt(15)
r.font.color.rgb = GREY

for _ in range(3):
    doc.add_paragraph()

meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = meta.add_run("Includes: Service Breakdown  •  DuckDB Placement  •  YAML Files  •  CI/CD Pipeline")
r.font.size = Pt(12)
r.font.italic = True
r.font.color.rgb = GREY

info = doc.add_paragraph()
info.alignment = WD_ALIGN_PARAGRAPH.CENTER
info.paragraph_format.space_before = Pt(40)
r = info.add_run("Version 1.0    |    Prepared for the Kavachio Engineering Team")
r.font.size = Pt(11)
r.font.color.rgb = GREY

doc.add_page_break()

# ============================================================
# TABLE OF CONTENTS
# ============================================================
add_heading("Table of Contents", 1)
toc_items = [
    "1. Introduction",
    "2. Why We Are Moving to Microservices",
    "3. Current System (Before)",
    "4. Target Architecture (After)",
    "5. The Seven Microservices - Full Detail",
    "6. Where DuckDB Will Stay",
    "7. Shared Database Strategy",
    "8. Shared File Storage",
    "9. API Gateway",
    "10. All YAML Files We Need to Create",
    "11. CI/CD Pipeline",
    "12. Step-by-Step Migration Plan",
    "13. Risks and How We Handle Them",
    "14. Summary",
]
for t in toc_items:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(t)
    r.font.size = Pt(12)
    if t[0].isdigit() and t.split(".")[0].isdigit():
        r.font.bold = True
        r.font.color.rgb = BLUE
doc.add_page_break()

# ============================================================
# 1. INTRODUCTION
# ============================================================
add_heading("1. Introduction", 1)
add_para(
    "This document explains how we will change the Kavachio application from one big "
    "program (a monolith) into many small, separate programs (microservices). It is "
    "written in simple English so that both technical and non-technical readers can "
    "follow it."
)
add_para("The document covers four things you asked about:")
add_bullet("Every microservice and what lives inside each one.", bold_lead="What services: ")
add_bullet("Exactly where DuckDB will run.", bold_lead="DuckDB: ")
add_bullet("All the YAML files we need to write.", bold_lead="YAML files: ")
add_bullet("How automatic build and deploy (CI/CD) fits into the plan.", bold_lead="CI/CD: ")
add_para(
    "Important decisions already agreed: (1) The old Node.js rule-engine is dropped - "
    "the code stays in the repo but is not used. Validation will be pure Python. "
    "(2) In the first phase we keep ONE shared database that every microservice can use."
)

# ============================================================
# 2. WHY MICROSERVICES
# ============================================================
add_heading("2. Why We Are Moving to Microservices", 1)
add_para("Today everything runs as one large FastAPI program. This causes some problems:")
add_bullet("If one part crashes, the whole application can go down.")
add_bullet("We must deploy everything together, even for a tiny change.")
add_bullet("Heavy jobs (like file uploads or LLM mapping) slow down light jobs (like login).")
add_bullet("We cannot scale only the busy parts - we have to scale everything.")
add_para("Splitting into microservices gives us these benefits:")
add_bullet("Each service can be deployed on its own.", bold_lead="Independent deploys: ")
add_bullet("We can run more copies of only the busy service.", bold_lead="Independent scaling: ")
add_bullet("If one service fails, the others keep working.", bold_lead="Fault isolation: ")
add_bullet("Small teams can own one service each.", bold_lead="Clear ownership: ")

# ============================================================
# 3. CURRENT SYSTEM
# ============================================================
add_heading("3. Current System (Before)", 1)
add_para("The system today has these main parts:")
add_table(
    ["Layer", "Technology", "Port", "Notes"],
    [
        ["Frontend", "React + Vite", "5173", "Single web application (SPA)"],
        ["Backend API", "FastAPI (one big app)", "8000", "About 6,400 lines in 4 route files"],
        ["Database", "PostgreSQL (single)", "5432", "Multi-tenant. Ops tables + 53 warehouse tables"],
        ["Validation DB", "DuckDB", "in-memory", "Created per request, inside the backend"],
        ["Old JS engine", "Node/Express", "4000", "DROPPED - code stays but is not used"],
    ],
    col_widths=[1.4, 1.9, 0.8, 2.6],
)
add_para(
    "All backend logic lives in one FastAPI process. There are no background workers "
    "and no message queues. Everything runs step by step in the same program."
)

# ============================================================
# 4. TARGET ARCHITECTURE
# ============================================================
add_heading("4. Target Architecture (After)", 1)
add_para(
    "After the migration we will have seven small backend services, one frontend, one "
    "API gateway in front, one shared database, and one shared file store. The picture "
    "below shows how they connect."
)
add_code(
    "                          +-------------+\n"
    "                          |  Frontend   |   React/Vite (:5173)\n"
    "                          +------+------+\n"
    "                                 | HTTPS (one address)\n"
    "                          +------v------+\n"
    "                          | API Gateway |   nginx / Kong (:8080)\n"
    "                          |  + JWT check|\n"
    "                          +------+------+\n"
    "      +--------+--------+--------+--------+--------+--------+--------+\n"
    "      v        v        v        v        v        v        v\n"
    "   [ auth ] [tenant] [mapper] [ingest] [valid.] [export] [contract]\n"
    "   :8001    :8007    :8002    :8003    :8004    :8005    :8006\n"
    "      |        |        |        |        |        |        |\n"
    "      +--------+--------+---+----+--------+--------+--------+\n"
    "                            v\n"
    "               +-------------------------+\n"
    "               |   Shared PostgreSQL DB  |  (:5432)  all tables\n"
    "               +-------------------------+\n"
    "                            |\n"
    "               +-------------------------+\n"
    "               |   Shared File Store     |  MinIO / S3 (:9000)\n"
    "               |   (Excel / PDF files)   |\n"
    "               +-------------------------+"
)
add_para("Key points about this design:", bold=True)
add_bullet("The frontend talks ONLY to the API gateway, never directly to a service.")
add_bullet("The gateway checks the login token (JWT) and sends the request to the right service.")
add_bullet("In Phase 1 all services share ONE PostgreSQL database.")
add_bullet("Files (Excel, PDF) are kept in one shared file store, not inside a single service.")

# ============================================================
# 5. THE SEVEN MICROSERVICES
# ============================================================
add_heading("5. The Seven Microservices - Full Detail", 1)
add_para(
    "Below is every service, one by one. For each we list: what it does, which code "
    "files move into it, which database tables it owns (is allowed to write), which "
    "tables it only reads, and its main API paths."
)

services = [
    {
        "name": "5.1  auth-service",
        "port": ":8001",
        "purpose": "Handles login, logout, token refresh, password reset, and user identity. "
                   "Every other service trusts the tokens this service creates.",
        "files": "auth_deps.py, auth_tokens.py, auth_utils.py, settings.py, email_utils.py",
        "owns": "users",
        "reads": "(none)",
        "endpoints": "/auth/login, /auth/refresh, /auth/logout, /auth/reset, /auth/jwks",
        "why": "Login is used by everyone. Keeping it separate means we can add MFA or "
               "OAuth later without touching business logic.",
    },
    {
        "name": "5.2  tenant-admin-service",
        "port": ":8007",
        "purpose": "Manages master data: tenants (customers), parties (companies), and "
                   "programs. These change slowly and are low volume.",
        "files": "tenant / party / program routes taken from app_routes.py",
        "owns": "tenants, parties, programs",
        "reads": "users",
        "endpoints": "/tenants, /parties, /programs (create, read, update, delete)",
        "why": "Admin data changes rarely and does not need the same scaling as uploads.",
    },
    {
        "name": "5.3  mapper-service",
        "port": ":8002",
        "purpose": "Takes the column headers of an uploaded Excel file and works out which "
                   "canonical (standard) field each column means. Uses Google Gemini plus "
                   "sentence embeddings, with a cosine-similarity fallback.",
        "files": "mapper.py, fingerprint.py, data_model.py, canonical.py (schema read)",
        "owns": "mappers, fingerprints",
        "reads": "canonical schema definitions",
        "endpoints": "/mapper/generate, /mapper/{id}, /mapper (list), /bdx/sheets, "
                     "/bdx/preview, /data-model",
        "why": "The LLM work is heavy and slow. Keeping it separate lets us scale it on "
               "its own and warm up the model without affecting other services.",
    },
    {
        "name": "5.4  ingestion-service",
        "port": ":8003",
        "purpose": "The main data-loading engine. Reads real Excel files, applies the "
                   "approved mapping spec (glom), and writes rows into the 53-table "
                   "canonical warehouse using SCD2 history tracking.",
        "files": "ingester.py, direct_lane.py, direct_mapper.py, direct_render.py, "
                 "assembler.py, scd2_sql.py",
        "owns": "uploads, uploads_policy, canonical warehouse tables (writes)",
        "reads": "mappers (approved specs)",
        "endpoints": "/bdx/upload, /direct/upload, /direct/*, /uploads, /uploads/{id}",
        "why": "This is the heaviest input/output path. Isolating it lets us add more "
               "copies during busy upload times without scaling everything else.",
    },
    {
        "name": "5.5  validation-service",
        "port": ":8004",
        "purpose": "Checks the loaded data against business rules. Loads the needed "
                   "warehouse rows into an in-memory DuckDB database, runs the rules, "
                   "and returns any exceptions. Now 100% Python (no Node).",
        "files": "duckdb_validation.py, validation_routes.py, rule definitions",
        "owns": "exception_decisions",
        "reads": "canonical warehouse tables (read only)",
        "endpoints": "/api/validate, /api/validate/exceptions/decide",
        "why": "DuckDB uses a lot of memory per run and has its own lifecycle. Keeping it "
               "separate protects the other services' memory.",
    },
    {
        "name": "5.6  export-service",
        "port": ":8005",
        "purpose": "Builds the final output Excel workbooks from approved export templates "
                   "and warehouse data. Good candidate for async (background) jobs.",
        "files": "exporter.py, assembler.py (shared read)",
        "owns": "export_templates",
        "reads": "canonical warehouse tables (read only)",
        "endpoints": "/export/template/*, /export/generate",
        "why": "Building a workbook can take several seconds. A separate service can do "
               "this in the background without blocking the main API.",
    },
    {
        "name": "5.7  contract-service",
        "port": ":8006",
        "purpose": "Handles contract uploads: reads PDF and Word files, extracts rules "
                   "using LLM, compiles and stores them. A self-contained domain of its own.",
        "files": "contract_upload_services/ (15+ modules: extraction, rule_compiler, "
                 "gemini_service, versioning, etc.)",
        "owns": "contracts, contract rule tables",
        "reads": "parties, programs",
        "endpoints": "/contracts/*, contract rule endpoints",
        "why": "Contract handling is a distinct area with its own document parsing and "
               "rule logic. It changes independently from the data pipeline.",
    },
]

for s in services:
    add_heading(s["name"] + "   " + s["port"], 2)
    add_para(s["purpose"])
    add_table(
        ["Property", "Detail"],
        [
            ["Code files that move here", s["files"]],
            ["Tables it OWNS (can write)", s["owns"]],
            ["Tables it only READS", s["reads"]],
            ["Main API paths", s["endpoints"]],
            ["Why it is separate", s["why"]],
        ],
        col_widths=[1.9, 4.8],
    )

add_heading("Quick Reference: All Services at a Glance", 2)
add_table(
    ["#", "Service", "Port", "Owns (writes)", "Main Job"],
    [
        ["1", "auth-service", "8001", "users", "Login & tokens"],
        ["2", "tenant-admin-service", "8007", "tenants, parties, programs", "Master data"],
        ["3", "mapper-service", "8002", "mappers, fingerprints", "LLM column mapping"],
        ["4", "ingestion-service", "8003", "uploads, warehouse", "Load data"],
        ["5", "validation-service", "8004", "exception_decisions", "Rule checks (DuckDB)"],
        ["6", "export-service", "8005", "export_templates", "Build output files"],
        ["7", "contract-service", "8006", "contracts, rules", "Contract extraction"],
    ],
    col_widths=[0.4, 1.9, 0.7, 2.1, 1.8],
)

# ============================================================
# 6. WHERE DUCKDB WILL STAY
# ============================================================
add_heading("6. Where DuckDB Will Stay", 1)
add_para(
    "DuckDB will live ONLY inside the validation-service. Nothing else uses it. "
    "This is an important and clear decision.", bold=True
)
add_para("How it works:")
add_number("A validation request arrives at validation-service with an upload ID.")
add_number("The service reads the needed rows from the shared PostgreSQL warehouse.")
add_number("It creates a fresh in-memory DuckDB database just for that request.")
add_number("It loads the rows into DuckDB and runs the validation rules as fast SQL.")
add_number("It returns the exceptions, then throws the DuckDB database away.")
add_para("Key facts about DuckDB placement:", bold=True)
add_bullet("DuckDB is in-memory by default. It is created and destroyed per request.")
add_bullet("It is NOT a shared database. Each validation-service copy has its own DuckDB.")
add_bullet(
    "Because DuckDB keeps state in memory, we should be careful when running many "
    "copies of validation-service. Each copy needs enough RAM for its own DuckDB."
)
add_bullet(
    "Optional debugging: DuckDB can write a file to /tmp (KAVACHIO_DUCKDB_PERSIST=True). "
    "This is only for looking at problems, not for production."
)
add_para(
    "In short: PostgreSQL is the permanent shared database for all services. DuckDB is a "
    "temporary in-memory helper that belongs to the validation-service alone.",
    italic=True, color=GREEN
)

# ============================================================
# 7. SHARED DATABASE STRATEGY
# ============================================================
add_heading("7. Shared Database Strategy", 1)
add_para(
    "As agreed, in Phase 1 there is ONE PostgreSQL database and every microservice "
    "connects to it. We do NOT split the database yet. This keeps the move simple and "
    "easy to undo if needed."
)
add_para("To keep order even with a shared database, we use a simple rule:", bold=True)
add_bullet(
    "Each service gets its own database login (for example auth_svc, mapper_svc). "
    "That login can WRITE only to the tables the service owns and READ the shared tables."
)
add_bullet(
    "This gives table-level safety inside one shared database - the safety of ownership "
    "without the cost of splitting."
)
add_table(
    ["Service", "Can WRITE", "Can READ"],
    [
        ["auth-service", "users", "-"],
        ["tenant-admin-service", "tenants, parties, programs", "users"],
        ["mapper-service", "mappers, fingerprints", "canonical schema"],
        ["ingestion-service", "uploads, warehouse tables", "mappers"],
        ["validation-service", "exception_decisions", "warehouse tables"],
        ["export-service", "export_templates", "warehouse tables"],
        ["contract-service", "contracts, contract rules", "parties, programs"],
    ],
    col_widths=[1.9, 2.6, 2.2],
)
add_para(
    "Later (Phase 3, optional): if one part needs its own scaling, we can move the "
    "canonical warehouse into a separate PostgreSQL instance. This is not required now."
)

# ============================================================
# 8. SHARED FILE STORAGE
# ============================================================
add_heading("8. Shared File Storage", 1)
add_para(
    "Today files are handled inside the single program's memory. Once services are "
    "separate, they cannot pass files in memory anymore. So we add a shared file store."
)
add_bullet("We will use MinIO for local and test (it works like Amazon S3).")
add_bullet("In the cloud we use Amazon S3 (or the cloud provider's object storage).")
add_para("How it works in simple steps:")
add_number("The user uploads an Excel or PDF through the gateway.")
add_number("The ingestion-service (or contract-service) saves the file in MinIO and gets a key (a short id).")
add_number("Only the key is passed around between services, not the whole file.")
add_number("Any service that needs the file reads it from MinIO using the key.")
add_para(
    "This keeps services small and stateless, and means a restart never loses an "
    "uploaded file.", italic=True
)

# ============================================================
# 9. API GATEWAY
# ============================================================
add_heading("9. API Gateway", 1)
add_para(
    "The API gateway is the single front door. The frontend sends every request to the "
    "gateway, and the gateway forwards it to the correct service."
)
add_para("The gateway does these jobs:")
add_bullet("Checks the login token (JWT) once, at the door.")
add_bullet("Adds helpful headers like X-User-Id and X-Tenant-Id for the services.")
add_bullet("Routes each URL path to the right service (see table).")
add_bullet("Can also do rate limiting and logging in one place.")
add_table(
    ["URL Path", "Goes To Service"],
    [
        ["/auth/*", "auth-service"],
        ["/tenants/*, /parties/*, /programs/*", "tenant-admin-service"],
        ["/mapper/*, /bdx/sheets, /bdx/preview, /data-model", "mapper-service"],
        ["/bdx/upload, /direct/*, /uploads/*", "ingestion-service"],
        ["/api/validate*", "validation-service"],
        ["/export/*", "export-service"],
        ["/contracts/*", "contract-service"],
    ],
    col_widths=[3.9, 2.8],
)

# ============================================================
# 10. YAML FILES
# ============================================================
add_heading("10. All YAML Files We Need to Create", 1)
add_para(
    "YAML files are simple text files that describe how to build, run, and deploy the "
    "services. Below is every YAML file we plan to create, grouped by purpose."
)

add_heading("10.1  Local Development - Docker Compose", 2)
add_para("One file to run everything on a developer's laptop with a single command.")
add_table(
    ["File", "What it does"],
    [
        ["infra/docker-compose.yml",
         "Starts all 7 services + PostgreSQL + MinIO + gateway together"],
        ["infra/docker-compose.override.yml",
         "Extra settings for local dev only (live code reload, debug ports)"],
        ["infra/.env (used by compose)",
         "Holds shared values: DB URL, JWT secret, Gemini API key"],
    ],
    col_widths=[2.6, 4.1],
)
add_para("Small example of what the compose file looks like:", italic=True, size=10)
add_code(
    "services:\n"
    "  postgres:\n"
    "    image: postgres:16\n"
    "    environment:\n"
    "      POSTGRES_DB: kavachio\n"
    "    ports: [\"5432:5432\"]\n"
    "  minio:\n"
    "    image: minio/minio\n"
    "    command: server /data\n"
    "    ports: [\"9000:9000\"]\n"
    "  auth-service:\n"
    "    build: ./services/auth\n"
    "    ports: [\"8001:8001\"]\n"
    "    depends_on: [postgres]\n"
    "  gateway:\n"
    "    image: nginx:alpine\n"
    "    volumes: [\"./infra/nginx.conf:/etc/nginx/nginx.conf\"]\n"
    "    ports: [\"8080:80\"]"
)

add_heading("10.2  Kubernetes - One Set Per Service", 2)
add_para(
    "For running in the cloud we use Kubernetes. Each service gets a small folder of "
    "YAML files. We create the same set of files for all seven services."
)
add_table(
    ["File (per service)", "What it does"],
    [
        ["deployment.yaml", "Tells Kubernetes how many copies to run and which image to use"],
        ["service.yaml", "Gives the service a stable internal address"],
        ["configmap.yaml", "Non-secret settings (ports, log level, feature flags)"],
        ["secret.yaml", "Secret values (DB password, JWT secret, Gemini key)"],
        ["hpa.yaml", "Auto-scaling rules (add copies when busy)"],
        ["ingress.yaml", "Only for the gateway - the public entry point"],
    ],
    col_widths=[2.3, 4.4],
)
add_para("Shared cluster-level YAML files (created once):")
add_bullet("namespace.yaml - a named space to hold all Kavachio objects.")
add_bullet("postgres-statefulset.yaml - the shared database (or a managed cloud DB instead).")
add_bullet("minio-statefulset.yaml - the shared file store (or cloud S3 instead).")
add_bullet("network-policy.yaml - controls which service can talk to which.")

add_heading("10.3  Helm Chart (Optional but Recommended)", 2)
add_para(
    "Instead of copying the same Kubernetes YAML seven times, we can use one Helm chart "
    "with values files. This keeps things tidy."
)
add_table(
    ["File", "What it does"],
    [
        ["charts/kavachio/Chart.yaml", "Describes the chart"],
        ["charts/kavachio/values.yaml", "Default settings for all services"],
        ["charts/kavachio/values-dev.yaml", "Overrides for the dev environment"],
        ["charts/kavachio/values-prod.yaml", "Overrides for production"],
        ["charts/kavachio/templates/*.yaml", "The reusable service templates"],
    ],
    col_widths=[2.9, 3.8],
)

add_heading("10.4  CI/CD Pipeline Files", 2)
add_para("These YAML files run the automatic build and deploy. Details are in Section 11.")
add_table(
    ["File", "What it does"],
    [
        [".github/workflows/ci.yaml", "Runs on every push: lint, test, build images"],
        [".github/workflows/cd-dev.yaml", "Deploys to the dev environment automatically"],
        [".github/workflows/cd-prod.yaml", "Deploys to production after approval"],
    ],
    col_widths=[2.9, 3.8],
)

# ============================================================
# 11. CI/CD PIPELINE
# ============================================================
add_heading("11. CI/CD Pipeline", 1)
add_para(
    "CI/CD means Continuous Integration and Continuous Deployment. In simple words: when "
    "a developer pushes code, the computer automatically tests it, builds it, and puts it "
    "into the running system - with little or no manual work."
)

add_heading("11.1  The Two Halves", 2)
add_bullet(
    "Every time code is pushed, run checks and build the service images. This catches "
    "mistakes early.", bold_lead="CI (Continuous Integration): ")
add_bullet(
    "After the checks pass, deploy the new images to the dev or production cluster.",
    bold_lead="CD (Continuous Deployment): ")

add_heading("11.2  The CI Steps (on every push)", 2)
add_number("Check out the code.", bold_lead="Checkout: ")
add_number("Run code style checks (lint) and type checks.", bold_lead="Lint: ")
add_number("Run automated tests for the changed service.", bold_lead="Test: ")
add_number("Build a Docker image for the service.", bold_lead="Build: ")
add_number("Push the image to the container registry with a version tag.", bold_lead="Push: ")
add_number("Scan the image for security problems.", bold_lead="Scan: ")

add_heading("11.3  The CD Steps (after CI passes)", 2)
add_number("Update the Kubernetes files with the new image version.")
add_number("Deploy to the DEV cluster automatically.")
add_number("Run quick smoke tests to confirm the service is healthy.")
add_number("Wait for a human to approve the production release.")
add_number("Deploy to PRODUCTION using a safe rolling update.")
add_number("If something breaks, roll back to the previous version automatically.")

add_heading("11.4  Smart Build - Only What Changed", 2)
add_para(
    "Because this is a monorepo (all services in one code repository), we build only the "
    "service whose code actually changed. If a developer edits the mapper-service, we do "
    "not rebuild the other six. This saves a lot of time."
)

add_heading("11.5  Example CI Workflow (GitHub Actions)", 2)
add_para("A short example of what the ci.yaml file looks like:", italic=True, size=10)
add_code(
    "name: CI\n"
    "on:\n"
    "  push:\n"
    "    branches: [ development, main ]\n"
    "jobs:\n"
    "  build-and-test:\n"
    "    runs-on: ubuntu-latest\n"
    "    strategy:\n"
    "      matrix:\n"
    "        service: [auth, mapper, ingestion, validation,\n"
    "                  export, contract, tenant-admin]\n"
    "    steps:\n"
    "      - uses: actions/checkout@v4\n"
    "      - name: Set up Python\n"
    "        uses: actions/setup-python@v5\n"
    "        with: { python-version: '3.11' }\n"
    "      - name: Install deps\n"
    "        run: pip install -r services/${{ matrix.service }}/requirements.txt\n"
    "      - name: Run tests\n"
    "        run: pytest services/${{ matrix.service }}/tests\n"
    "      - name: Build image\n"
    "        run: docker build -t kavachio/${{ matrix.service }}:${{ github.sha }} \\\n"
    "             services/${{ matrix.service }}\n"
    "      - name: Push image\n"
    "        run: docker push kavachio/${{ matrix.service }}:${{ github.sha }}"
)

add_heading("11.6  Environments", 2)
add_table(
    ["Environment", "Purpose", "Deploy Trigger"],
    [
        ["Local", "Developer laptop with docker-compose", "Manual (docker compose up)"],
        ["Dev", "Shared test cluster", "Automatic on push to development"],
        ["Production", "Live system for real users", "Manual approval, then automatic"],
    ],
    col_widths=[1.6, 2.9, 2.2],
)

# ============================================================
# 12. MIGRATION PLAN
# ============================================================
add_heading("12. Step-by-Step Migration Plan", 1)
add_para(
    "We move in careful phases. Each phase gives a working system, so we never take a "
    "big risk all at once."
)

add_heading("Phase 0 - Reorganize the Code (about 1 week)", 3)
add_bullet("Reshape the repo into a folder per service, plus a shared folder for common code.")
add_bullet("No logic changes yet. Just moving files into their new homes.")
add_code(
    "kavachio/\n"
    "  gateway/\n"
    "  services/\n"
    "    auth/  tenant-admin/  mapper/  ingestion/\n"
    "    validation/  export/  contract/\n"
    "  shared/          # db.py, canonical.py, data_model.py, auth helpers\n"
    "  frontend/\n"
    "  infra/           # docker-compose.yml, nginx.conf\n"
    "  charts/          # helm chart\n"
    "  .github/workflows/   # ci.yaml, cd-dev.yaml, cd-prod.yaml"
)

add_heading("Phase 1 - Containers + Shared DB + Gateway (2 to 3 weeks)", 3)
add_bullet("Write a Dockerfile for each service.")
add_bullet("Write docker-compose.yml with all 7 services + PostgreSQL + MinIO + gateway.")
add_bullet("One shared PostgreSQL, exactly as agreed.")
add_bullet("Add health-check endpoints to every service.")
add_bullet("Result: real microservice deployment with NO business-logic rewrite.")

add_heading("Phase 2 - Split the Routes Into Real Services (4 to 6 weeks)", 3)
add_para("Pull the routes out of the old monolith one service at a time. Safest first:")
add_table(
    ["Order", "Service", "Risk", "Reason"],
    [
        ["1", "auth-service", "Low", "Cleanest boundary, stateless tokens"],
        ["2", "tenant-admin-service", "Low", "Simple low-volume CRUD"],
        ["3", "mapper-service", "Medium", "Clear input and output"],
        ["4", "validation-service", "Medium", "Self-contained with DuckDB"],
        ["5", "export-service", "Low", "Read-only and can be async"],
        ["6", "contract-service", "Medium", "Self-contained domain"],
        ["7", "ingestion-service", "High", "Heaviest, touches the warehouse - do last"],
    ],
    col_widths=[0.7, 2.0, 0.9, 3.1],
)

add_heading("Phase 3 - Hardening (ongoing)", 3)
add_bullet("Add background job queue (Redis + ARQ or Celery) for ingestion and export.")
add_bullet("Add central logging and tracing so we can follow a request across services.")
add_bullet("Add per-service metrics and dashboards.")
add_bullet("Optional: give the warehouse its own database only if it needs separate scaling.")

# ============================================================
# 13. RISKS
# ============================================================
add_heading("13. Risks and How We Handle Them", 1)
add_table(
    ["Risk", "How We Handle It"],
    [
        ["Ingestion writes to many tables at once",
         "Keep shared DB in Phase 1 so we can still use one database transaction"],
        ["LLM mapping is slow",
         "Make mapper async - return quickly and let the client poll for the result"],
        ["DuckDB uses a lot of memory",
         "Give each validation-service copy enough RAM; be careful adding many copies"],
        ["Files can no longer pass in memory",
         "Use shared MinIO/S3 store; pass a file key, not the whole file"],
        ["JWT secret must be shared safely",
         "auth-service owns the secret and publishes a JWKS endpoint for others"],
        ["Harder to follow a request across services",
         "Add tracing (OpenTelemetry) and a request id passed through the gateway"],
        ["Deploy mistakes",
         "CI/CD with smoke tests and automatic rollback to the last good version"],
    ],
    col_widths=[2.8, 3.9],
)

# ============================================================
# 14. SUMMARY
# ============================================================
add_heading("14. Summary", 1)
add_para("The plan in one paragraph:", bold=True)
add_para(
    "We split the one big FastAPI program into seven small Python services: auth, "
    "tenant-admin, mapper, ingestion, validation, export, and contract. A gateway sits "
    "in front and sends each request to the right service. In Phase 1 all services share "
    "ONE PostgreSQL database, and files live in a shared MinIO/S3 store. DuckDB stays "
    "only inside the validation-service, created fresh in memory for each check. We write "
    "YAML files for local Docker, for Kubernetes, and for automatic build-and-deploy "
    "(CI/CD). The old Node.js engine is dropped. We move in safe phases, starting with "
    "containers and the gateway, then splitting routes one service at a time, with the "
    "heavy ingestion service moved last."
)
add_divider()
add_para("Recommended first move:", bold=True, color=GREEN)
add_bullet("Do Phase 0 and Phase 1 together: reorganize the code, containerize, add the "
           "shared database and gateway. About 3 to 4 weeks, with no business-logic rewrite.")
add_bullet("Then split routes one by one, easiest first, ingestion last.")

add_para("")
end = doc.add_paragraph()
end.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = end.add_run("--- End of Document ---")
r.font.color.rgb = GREY
r.font.italic = True

out = "/Users/at-mac11/Documents/Dinesh/POC/Kawachu/Code/Git/kavachio/docs/Kavachio_Microservices_Architecture.docx"
doc.save(out)
print("SAVED:", out)
