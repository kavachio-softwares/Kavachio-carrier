"""
version_diff_sheet.py
─────────────────────
The Previous Submission Reconciliation (version_diff.py) as people read it:
a "Previous submission" sheet appended to the downloaded BDX Excel, beside
the "Results" sheet (plan 23.4), and the changes as CSV.

The sheet reads like a file comparison, but rows are paired on the business
key, never by position. It is written for someone opening it for the first
time, in plain sentences:
  - which version this file is, which version it was compared with, and how
    a row is recognised in both;
  - the counts as coloured boxes: All, then the task's four outcomes — New,
    Changed, Disappeared, Restated — and what each colour in the table means
    (only the colours the table uses);
  - a totals warning only when an amount's change is not explained by the
    rows;
  - the bordereau itself: every row of this file in its own order, with its
    outcome; what happened to the row is a note on its Outcome cell; a
    changed value is highlighted and carries the accepted value in a note; a
    new row is tinted; a missing row is put back where it was in the
    accepted version, struck through.

The CSV and the API keep one line per changed value.

Same look as the Results sheet (export_results): navy headers, light borders,
never the two fills the in-site viewers read as exception marks
(exporter._INVALID_FILL_RGB / _WARNING_FILL_RGB). Nothing is frozen: a
frozen block taller than the window hid the rows. Never raises: on any
failure the workbook comes back unchanged.

No database and no web imports, so it can be tested on its own.
"""
from __future__ import annotations

import csv
import io
import os
import re
import zipfile
from collections import defaultdict
from datetime import date, datetime, timezone
from typing import Optional

import version_diff as vd

SHEET_NAME = "Previous submission"
TITLE = "Previous submission reconciliation"
# The workbook is stamped "<LAYOUT>|<state>" when the sheet is built. Bump
# LAYOUT whenever the sheet's layout or wording changes: "Download BDX"
# rebuilds a stored sheet with an older stamp, or one that was not compared
# (version_diff_store.refresh_download), and serves the rest as stored.
LAYOUT = "3"
_STAMP = "Kavachio previous submission"
# Lines listed on the sheet unless PREVIOUS_SUBMISSION_MAX_LINES says
# otherwise (10,000 ≈ 550 KB stored with the output); the CSV and the API
# always have every row.
_DEFAULT_MAX_LINES = 10_000


def max_lines() -> int:
    try:
        return max(1, int(os.getenv("PREVIOUS_SUBMISSION_MAX_LINES", "").strip()
                          or _DEFAULT_MAX_LINES))
    except ValueError:
        return _DEFAULT_MAX_LINES

_FONT = "Calibri"
_NAVY = "1F3A5F"
_LINE = "D0D7DE"
_ZEBRA = "F6F8FA"
_MUTED = "5F6368"
_TINT = {                          # outcome → (soft fill, text)
    vd.NEW: ("F3E8FD", "6A1B9A"),
    vd.CHANGED: ("E8F0FE", "1A56DB"),
    vd.DISAPPEARED: ("FEF7E0", "8A5A00"),
    vd.RESTATED: ("FDECEA", "8C1D18"),
    vd.UNCHANGED: ("E6F4EA", "1E7B34"),
}
# The boxes and the colour guide: the task's four outcomes, in its order.
SHOWN = (vd.NEW, vd.CHANGED, vd.DISAPPEARED, vd.RESTATED)
# What each outcome means, for someone reading the sheet for the first time
# (the note on its box).
MEANING = {
    vd.NEW: "Is in this file, was not in the accepted version "
            "(a new endorsement, a new policy or another new row).",
    vd.CHANGED: "Same row (same policy, transaction type and effective date), "
                "but some details other than amounts are different.",
    vd.DISAPPEARED: "Was in the accepted version, is not in this file.",
    vd.RESTATED: "Same row (same policy, transaction type and effective date), "
                 "but an amount is different.",
}
# Who and which row first, then the value (was → now) and why; the rest of
# the key after, for filtering.
LINE_COLUMNS = ("Outcome", "Policy", "Insured", "Row in this file", "Row in accepted version",
                "Field", "Accepted version", "This file", "Difference", "What happened",
                "Transaction type", "Effective date")
# A changed cell: a stronger shade of its row's outcome colour.
_CELL_TINT = {vd.CHANGED: "C6DAFC", vd.RESTATED: "F8D0CB"}
# The bordereau view: these columns, then the columns the setup reads. What
# happened to a row is a note on its Outcome cell, not a column: a changed
# value already shows itself (highlighted, the old value in its note).
VIEW_FIXED = ("Outcome", "Row in this file", "Row in accepted version")
# Column widths follow what is in them, within these bounds (characters);
# left-aligned text is indented, which takes about _INDENT characters.
_MIN_WIDTH, _MAX_WIDTH, _INDENT = 10, 40, 2
LEGEND_TITLE = "How to read the table below"
# One sentence per colour the table can use, for someone with no background:
# {v} is the version this file was compared with ("version 1").
LEGEND = {
    vd.NEW: "A new row: it was not in {v}.",
    vd.CHANGED: "The same row as in {v}, but some details (not amounts) are different. The "
                "changed values are highlighted: point at one to see what it was in {v}.",
    vd.DISAPPEARED: "This row was in {v} but is not in this file. It is shown crossed out, "
                    "where it used to be.",
    vd.RESTATED: "The same row as in {v}, but an amount is different. The changed amounts are "
                 "highlighted: point at one to see what it was in {v}.",
}
LEGEND_NOTE = "Point at the word in the Outcome column to read what happened to that row."


def _lines_needed(text, width: float, bold: bool = False) -> int:
    """Lines a wrapped cell needs at this column width — word by word, on
    the safe side (neither Excel nor LibreOffice grows rows openpyxl wrote)."""
    if text is None or text == "":
        return 1
    cap = max(4, int(width * (1.0 if bold else 1.2)) - 2)
    total = 0
    for para in str(text).split("\n"):
        n, cur = 1, 0
        for word in para.split():
            size = len(word)
            if cur and cur + 1 + size <= cap:
                cur += 1 + size
                continue
            if cur:
                n += 1
            n += (size - 1) // cap
            cur = size % cap or cap
        total += n
    return total


def _note(text: str):
    """A cell note big enough to show `text` without scrolling (~45
    characters a line at this width)."""
    from openpyxl.comments import Comment
    lines = sum(max(1, -(-len(p) // 45)) for p in str(text).split("\n"))
    return Comment(text, "Kavachio", height=24 + 16 * lines, width=320)


def _free_name(existing: list[str]) -> str:
    taken = {n.lower() for n in existing}
    name, n = SHEET_NAME, 2
    while name.lower() in taken:
        name, n = f"{SHEET_NAME} {n}", n + 1
    return name


def _num(v) -> str:
    return "" if v is None else f"{float(v):,.2f}"


def _signed(v) -> str:
    v = float(v or 0)
    return ("+" if v > 0 else "−" if v < 0 else "") + _num(abs(v))


# ── what goes on the sheet ──────────────────────────────────────────────────

def money_fields(result: dict) -> dict:
    """{money column: data-model field or None} over every compared sheet."""
    out: dict = {}
    for s in result.get("sheets") or []:
        out.update(s.get("money_fields") or {c: None for c in s.get("money_columns") or []})
    return out


def _amounts(cols: list[str], fields: dict) -> list[str]:
    """The amount columns among `cols` — every one the setup maps to a money
    field, rates (commission %) left out: a rate is never added up."""
    return [c for c in cols if not vd._is_rate(c, fields.get(c))]


def value_lines(row: dict, fields: Optional[dict] = None) -> list[tuple]:
    """(field, accepted value, this file's value, difference) — one per value
    that differs. A changed or restated row lists every value; a new or
    disappeared row its amounts."""
    fields = fields or {}
    o, money = row.get("outcome"), row.get("money") or {}
    out: list[tuple] = []
    if o in (vd.NEW, vd.DISAPPEARED):
        for c in _amounts(list(money), fields):
            m = money[c]
            out.append((c, m.get("was"), m.get("now"), m.get("change")))
    else:
        for c in sorted(money, key=lambda c: vd._is_rate(c, fields.get(c))):
            m = money[c]
            out.append((c, m.get("was"), m.get("now"), m.get("change")))
        for c, ch in (row.get("changes") or {}).items():
            was, now = ch.get("was"), ch.get("now")
            out.append((c, "(blank)" if was in (None, "") else was,
                        "(blank)" if now in (None, "") else now, None))
    return out or [(None, None, None, None)]


def _when(v) -> Optional[date]:
    try:
        return date.fromisoformat(str(v)[:10]) if v else None
    except ValueError:
        return None


def line_values(row: dict, line: tuple) -> list:
    """One sheet / CSV line (LINE_COLUMNS) for one value of one row."""
    outcome = vd.OUTCOME_WORDS.get(row.get("outcome"), row.get("outcome"))
    field, was, now, diff = line
    eff = row.get("effective_date")
    return [outcome, row.get("policy"), row.get("insured"), row.get("row"),
            row.get("accepted_row"), field, was, now, diff, row.get("detail") or "",
            row.get("type"), _when(eff) or eff]


def sorted_rows(result: dict) -> list[dict]:
    """New, Changed, Disappeared, Restated, then Unchanged (the task's order)
    — and by row within each."""
    order = {o: i for i, o in enumerate(vd.OUTCOMES)}
    return sorted(result.get("rows") or [], key=lambda r: (
        order.get(r.get("outcome"), 99), str(r.get("sheet") or ""),
        r.get("row") or 0, r.get("accepted_row") or 0))


def _received(info: dict) -> Optional[datetime]:
    try:
        return datetime.fromisoformat(info["received_at"]) if info.get("received_at") else None
    except (TypeError, ValueError):
        return None


def _matched_on(result: dict) -> Optional[str]:
    """'policy number, transaction type and transaction effective date' — the
    key parts the compared sheets used, as words in a sentence."""
    used = []
    for s in result.get("sheets") or []:
        for p in s.get("parts_used") or []:
            if p not in used:
                used.append(p)
    words = [vd.rk.PART_WORDS[p].lower() for p in used]
    if not words:
        return None
    return words[0] if len(words) == 1 else f"{', '.join(words[:-1])} and {words[-1]}"


def _version_word(result: dict) -> str:
    """'version 1' — the version this file was compared with."""
    acc = (result.get("accepted") or {}).get("version")
    return f"version {acc}" if acc else "the last version sent to the carrier"


def about_for(result: dict, checked=None) -> dict:
    """The lines under the title, as plain sentences: which version this file
    is, which one it was compared with, how a row is recognised in both, and
    when. `checked` is shown in the same time zone as the versions' own
    times."""
    this, acc = result.get("this") or {}, result.get("accepted") or {}
    tz = next((d.tzinfo for d in (_received(this), _received(acc)) if d and d.tzinfo), None)
    if isinstance(checked, datetime) and tz is not None:
        if checked.tzinfo is None:
            checked = checked.replace(tzinfo=timezone.utc)
        checked = checked.astimezone(tz)

    def received(info: dict) -> str:
        when = _received(info)
        return f"received {when:%d %b %Y}" if when else ""

    mine = (f"This file is version {this['version']} of this submission"
            + (f", {received(this)}" if received(this) else "") + "."
            if this.get("version") else None)
    theirs = (f"It is compared with version {acc['version']}, the last version that was sent "
              f"to the carrier" + (f" ({received(acc)})" if received(acc) else "") + "."
              if acc.get("version") else None)
    key = _matched_on(result)
    return {
        "This file": mine,
        "Compared with": theirs,
        "Rows matched on": (f"A row counts as the same row in both versions when its {key} "
                            f"are the same." if key else None),
        "Checked": checked,
    }


def _about_lines(about: Optional[dict]) -> list[str]:
    """One sentence per line; the check time as 'Checked on 16 Oct 2026 at
    18:13 (UTC+05:30).'"""
    lines = []
    for v in (about or {}).values():
        if v is None or (isinstance(v, str) and not v.strip()):
            continue
        if isinstance(v, datetime):
            off = v.utcoffset() if v.tzinfo else None
            z = v.strftime("%z")
            zone = ("" if off is None else "UTC" if not off else f"UTC{z[:3]}:{z[3:]}")
            v = f"Checked on {v:%d %b %Y} at {v:%H:%M}" + (f" ({zone})" if zone else "") + "."
        elif isinstance(v, date):
            v = f"Checked on {v:%d %b %Y}."
        lines.append(" ".join(str(v).split()))
    return lines


def in_short(result: dict) -> list[str]:
    """What the sheet says above the table besides the counts. A file not
    compared: why. A compared file: nothing, unless an amount changed by
    more than the rows explain — then one warning per such amount (the
    counts are in the boxes, the changes in the table)."""
    s = result.get("summary") or {}
    if result.get("state") != "compared":
        return [s.get("plain") or "Not compared."]
    name = _version_word(result)
    fields = money_fields(result)
    money = {c: m for c, m in (s.get("money") or {}).items()
             if m.get("accepted") or m.get("this")}
    out = []
    for c in _amounts(list(money), fields):
        m = money[c]
        gap = m.get("control_gap") or 0
        if gap:
            was, now = m.get("accepted") or 0, m.get("this") or 0
            out.append(f"Please check the totals: {c} is {_num(now)} in this file and "
                       f"{_num(was)} in {name}, and {_signed(gap)} of that difference is not "
                       f"explained by the rows below.")
    return out


def _shown(v) -> str:
    """A cell as it is shown: amounts with grouping and two decimals."""
    if v is None:
        return ""
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return _num(v)
    if isinstance(v, date):
        return v.strftime("%d %b %Y")
    return str(v)


def column_widths(tables: list[tuple[list[list], set]]) -> list[float]:
    """A width per column from what the tables put in it: the longest value,
    or the longest word where the text wraps (every heading, and the columns
    a table names as wrapping), plus the indent left-aligned text gets.
    `tables` = [(rows, {column indexes that wrap})], heading row first."""
    n = max((len(row) for t, _w in tables for row in t), default=0)
    out = []
    for j in range(n):
        need = _MIN_WIDTH
        for t, wraps in tables:
            for i, row in enumerate(t):
                if j >= len(row):
                    continue
                v = row[j]
                text = _shown(v)
                words = max((len(w) for w in text.split()), default=0)
                number = isinstance(v, (int, float)) and not isinstance(v, bool)
                if i == 0:              # a bold heading: whole words, at most ~2 lines
                    size = max(round(words * 1.15) + 1, -(-len(text) // 2))
                else:
                    size = words if j in wraps else len(text)
                    if not number:      # LibreOffice draws text ~20% wider than this counts
                        size = round(size * 1.2)
                need = max(need, size + 2 + (0 if number else _INDENT))
        out.append(min(need, _MAX_WIDTH))
    return out


# ── the bordereau view ──────────────────────────────────────────────────────

def _was_now(item: dict) -> str:
    """'Gross Premium: 9,800.00 → 10,800.00; Insured: A → B'."""
    bits = []
    for c, m in (item.get("money") or {}).items():
        if m.get("was") is not None and m.get("now") is not None:
            bits.append(f"{c}: {_shown(m['was'])} → {_shown(m['now'])}")
    for c, ch in (item.get("changes") or {}).items():
        bits.append(f"{c}: {_shown(ch.get('was')) or '(blank)'} → "
                    f"{_shown(ch.get('now')) or '(blank)'}")
    return "; ".join(bits)


def _entry(item: Optional[dict], values: dict, side: str, row: Optional[int],
           accepted_row: Optional[int]) -> dict:
    """One line of the view: a row of this file, or one only the accepted
    version had (side 'accepted'). `what` is the note on its Outcome cell:
    what happened, then — for a changed row — which values changed, since a
    highlighted cell can be far to the right."""
    if item is None:
        return {"outcome": vd.OUTCOME_WORDS[vd.UNCHANGED], "tint": vd.UNCHANGED, "what": "",
                "row": row, "accepted_row": accepted_row, "values": values, "was": {},
                "side": side}
    was = {c: m["was"] for c, m in (item.get("money") or {}).items()
           if m.get("was") is not None and m.get("now") is not None}
    was.update({c: ch.get("was") for c, ch in (item.get("changes") or {}).items()})
    what = item.get("detail") or ""
    if was:
        what = f"{what}\nWhat changed: {_was_now(item)}".strip()
    o = item.get("outcome")
    return {"outcome": vd.OUTCOME_WORDS.get(o, o), "tint": o,
            "what": what, "row": item.get("row"), "accepted_row": item.get("accepted_row"),
            "values": values, "was": was, "side": side}


def view_tables(result: dict) -> list[dict]:
    """Per compared sheet: {sheet, columns, entries}. Every row of this file
    in its own order (repeated heading lines left out); a row only the
    accepted version had is put back right after the row that came before it
    there, as a file comparison shows a deleted line."""
    items = result.get("rows") or []
    out = []
    for sheet, v in (result.get("view") or {}).items():
        mine = [x for x in items if x.get("sheet") == sheet]
        by_row = {x["row"]: x for x in mine if x.get("row")}
        pairs = list(v.get("matched") or [])
        this_to_acc = {c: b for c, b in pairs}
        partner = {b: c for c, b in pairs}             # accepted row → this row, 0-based
        for x in mine:
            if x.get("row") and x.get("accepted_row"):
                partner[x["accepted_row"] - 1] = x["row"] - 1
        after: dict = defaultdict(list)
        for x in sorted((x for x in mine if not x.get("row") and x.get("accepted_row")),
                        key=lambda x: x["accepted_row"]):
            before = [p for p in partner if p < x["accepted_row"] - 1]
            after[partner[max(before)] if before else -1].append(x)
        cols = list(v.get("columns") or [])
        dated = set(v.get("date_columns") or [])
        acc_col = v.get("accepted_column") or {}

        def shown(c, value):
            return (_when(value) or value) if c in dated and isinstance(value, str) else value

        acc_rows = v.get("accepted_rows") or []

        def gone(x):
            i = x["accepted_row"] - 1
            row = acc_rows[i] if 0 <= i < len(acc_rows) else {}
            return _entry(x, {c: shown(c, row.get(acc_col.get(c) or c)) for c in cols},
                          "accepted", None, x["accepted_row"])

        entries = [gone(x) for x in after[-1]]
        for i, row in enumerate(v.get("this_rows") or []):
            x = by_row.get(i + 1)
            if x is None and i not in this_to_acc:
                continue
            acc = this_to_acc.get(i)
            entries.append(_entry(x, {c: shown(c, row.get(c)) for c in cols}, "this", i + 1,
                                  acc + 1 if acc is not None else None))
            entries.extend(gone(y) for y in after.get(i, []))
        out.append({"sheet": sheet, "columns": cols, "entries": entries})
    return out


def count_boxes(result: dict) -> list[tuple[str, int, Optional[str]]]:
    """(label, rows, outcome) for the boxes at the top: All rows, then the
    task's four outcomes — new, changed, disappeared, restated (the rest of
    All is unchanged)."""
    o = (result.get("summary") or {}).get("outcomes") or {}
    return ([("All", sum(o.get(k, 0) for k in vd.OUTCOMES), "all")]
            + [(vd.OUTCOME_WORDS[k], o.get(k, 0), k) for k in SHOWN])


# ── the sheet ───────────────────────────────────────────────────────────────

def _styled(ws, result: dict, about: Optional[dict]) -> None:
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    thin = Side(style="thin", color=_LINE)
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_box = Border(left=thin, right=thin, top=thin, bottom=Side(style="medium", color=_NAVY))
    acc_no = (result.get("accepted") or {}).get("version")
    acc_word = f"Version {acc_no}" if acc_no else "Accepted version"
    cap = max_lines()
    fields = money_fields(result)

    # Everything the tables will hold, first: the column widths follow it.
    tables = view_tables(result) if result.get("state") == "compared" else []
    fixed = list(VIEW_FIXED)
    if acc_no:
        fixed[2] = f"Row in version {acc_no}"

    def cells(e: dict, columns: list) -> list:
        return [e["outcome"], e["row"], e["accepted_row"],
                *[e["values"].get(c) for c in columns]]

    width_rows = []
    left = cap
    for t in tables:
        t["shown"] = t["entries"][:max(0, left)]
        left -= len(t["shown"])
        width_rows.append(([fixed + t["columns"]] + [cells(e, t["columns"]) for e in t["shown"]],
                           set()))
    boxes = count_boxes(result)
    width_rows.append(([[label for label, _n, _t in boxes]], set()))
    widths = column_widths(width_rows) if width_rows else []

    # One object per look, shared by every cell that has it: building new
    # style objects per cell made a 16,000-line sheet take over 20 seconds.
    cache: dict = {}

    def font(**kw):
        key = ("font", *sorted(kw.items()))
        if key not in cache:
            cache[key] = Font(name=_FONT, **{"size": 11, **kw})
        return cache[key]

    def fill(rgb):
        if ("fill", rgb) not in cache:
            cache[("fill", rgb)] = PatternFill("solid", fgColor=rgb)
        return cache[("fill", rgb)]

    def align(number: bool, wrap: bool, centre: bool = False):
        key = ("align", number, wrap, centre)
        if key not in cache:
            cache[key] = Alignment(
                horizontal="center" if centre else "right" if number else "left",
                vertical="center" if centre else "top", indent=0 if number or centre else 1,
                wrap_text=wrap and not number, shrink_to_fit=number and not centre)
        return cache[key]

    def put(r, j, v, *, money=False, bg=None, border=True, wrap=True, **f):
        c = ws.cell(row=r, column=j, value=v)
        number = isinstance(v, (int, float)) and not isinstance(v, bool)
        if money and number:
            c.number_format = "#,##0.00"
        elif isinstance(v, date):
            c.number_format = "DD MMM YYYY"
        c.alignment = align(number, wrap)
        c.font = font(**f)
        if border:
            c.border = box
        if bg:
            c.fill = fill(bg)
        return c

    def fit(r, values, bold=False):
        """Row height for the wrapped text cells of row r."""
        n = max([1] + [_lines_needed(v, widths[j - 1] - _INDENT, bold)
                       for j, v in enumerate(values, start=1)
                       if isinstance(v, str) and j <= len(widths)])
        if n > 1:
            ws.row_dimensions[r].height = 15 * n + 2

    def section(r, text, *notes):
        ws.cell(row=r, column=1, value=text).font = font(size=13, bold=True, color=_NAVY)
        for note in notes:
            r += 1
            ws.cell(row=r, column=1, value=note).font = font(size=10, color=_MUTED)
        return r + 1

    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = _NAVY
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w

    ws["A1"] = TITLE
    ws["A1"].font = font(size=16, bold=True, color=_NAVY)
    r = 2
    for line in _about_lines(about):
        ws.cell(row=r, column=1, value=line).font = font(size=10, color=_MUTED)
        r += 1
    r += 1

    if result.get("state") != "compared":
        # Nothing was compared: say why for each sheet, and stop there.
        r = section(r, "Why this file was not compared")
        for line in in_short(result):
            ws.cell(row=r, column=1, value=line).font = font(bold=True)
            r += 1
        sheets = [x for x in result.get("sheets") or [] if x.get("state_text")]
        if sheets:
            r = section(r + 1, "Sheets")
            for x in sheets:
                ws.cell(row=r, column=1, value=f"{x.get('sheet')}: {x['state_text']}").font = font()
                r += 1
        return

    # The counts, as boxes: a label over a number, coloured by outcome.
    for j, (label, n, tint) in enumerate(boxes, start=1):
        bg, ink = (_NAVY, "FFFFFF") if tint == "all" else _TINT[tint]
        for rr, v, f in ((r, label, {"bold": True, "size": 10}),
                         (r + 1, n, {"bold": True, "size": 16})):
            c = ws.cell(row=rr, column=j, value=v)
            c.fill, c.font, c.border = fill(bg), font(color=ink, **f), box
            c.alignment = align(False, True, centre=True)
            if rr == r and tint != "all":
                c.comment = Comment(MEANING.get(tint, ""), "Kavachio")
    ws.row_dimensions[r + 1].height = 26
    r += 3

    # What the colours mean: one line per outcome the table actually shows,
    # its Outcome word in that outcome's colours, then a plain sentence.
    shown = [e for t in tables for e in t["shown"]]
    present = {e["tint"] for e in shown}
    if present & set(SHOWN):
        ws.cell(row=r, column=1, value=LEGEND_TITLE).font = font(bold=True, color=_NAVY)
        r += 1
        for tint in SHOWN:
            if tint not in present:
                continue
            obg, oink = _TINT[tint]
            c = ws.cell(row=r, column=1, value=vd.OUTCOME_WORDS[tint])
            c.fill, c.font, c.border = fill(obg), font(color=oink, bold=True), box
            c.alignment = align(False, False)
            ws.cell(row=r, column=2,
                    value=LEGEND[tint].format(v=_version_word(result))).font = font()
            r += 1
        if any(e["what"] for e in shown):
            ws.cell(row=r, column=1, value=LEGEND_NOTE).font = font(size=10, color=_MUTED)
            r += 1
        r += 1
    warnings = in_short(result)
    for line in warnings:
        ws.cell(row=r, column=1, value=line).font = font(
            bold=True, color=_TINT[vd.RESTATED][1])
        r += 1
    if warnings:
        r += 1

    # The bordereau: every row, coloured by what happened to it.
    first_table = None
    for t in tables:
        if len(tables) > 1:
            r = section(r, f"Sheet: {t['sheet']}")
        labels = fixed + t["columns"]
        for j, h in enumerate(labels, start=1):
            c = put(r, j, h, bold=True, color="FFFFFF", bg=_NAVY)
            c.border = head_box
            c.alignment = Alignment(horizontal="left", vertical="center", indent=1,
                                    wrap_text=True)
        fit(r, labels, bold=True)
        top = r
        for e in t["shown"]:
            r += 1
            vals = cells(e, t["columns"])
            tint = e["tint"]
            gone = e["side"] == "accepted"
            row_bg = (_TINT[vd.DISAPPEARED][0] if gone and tint == vd.DISAPPEARED
                      else _TINT[vd.NEW][0] if tint == vd.NEW
                      else None)
            for j, v in enumerate(vals, start=1):
                col = labels[j - 1]
                changed = j > len(fixed) and col in e["was"]
                c = put(r, j, v, money=col in fields,
                        bg=_CELL_TINT.get(tint) if changed else row_bg,
                        bold=changed, strike=gone or None,
                        color=_MUTED if gone else None)
                if changed:
                    c.comment = Comment(f"{acc_word}: {_shown(e['was'][col]) or '(blank)'}",
                                        "Kavachio")
            obg, oink = _TINT.get(tint, _TINT[vd.UNCHANGED])
            c = ws.cell(row=r, column=1)
            c.fill, c.font = fill(obg), font(color=oink, bold=True)
            if e["what"]:
                c.comment = _note(e["what"])
            fit(r, vals)
        if first_table is None and t["shown"]:
            first_table = f"A{top}:{get_column_letter(len(labels))}{r}"
        r += 2
    if first_table:
        ws.auto_filter.ref = first_table
    total = sum(len(t["entries"]) for t in tables)
    if total > cap:
        ws.cell(row=r, column=1, value=(
            f"Only the first {cap:,} rows are shown here; the CSV download lists every "
            f"change.")).font = font(size=10, color=_MUTED)

    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def _plain_sheet(ws, result: dict, about: Optional[dict]) -> None:
    ws.append([TITLE])
    for line in _about_lines(about):
        ws.append([line])
    ws.append([])
    if result.get("state") == "compared":
        boxes = count_boxes(result)
        ws.append([label for label, _n, _t in boxes])
        ws.append([n for _l, n, _t in boxes])
        ws.append([])
    for line in in_short(result):
        ws.append([line])
    left = max_lines()
    for t in view_tables(result) if result.get("state") == "compared" else []:
        ws.append([])
        ws.append(list(VIEW_FIXED) + t["columns"])
        for e in t["entries"][:max(0, left)]:
            ws.append([e["outcome"], e["row"], e["accepted_row"],
                       *[e["values"].get(c) for c in t["columns"]]])
            if e["what"]:
                ws.cell(row=ws.max_row, column=1).comment = _note(e["what"])
        left -= len(t["entries"])


def _is_ours(ws) -> bool:
    """A sheet this module made — not a template's own sheet that happens to
    be called "Previous submission"."""
    return ws.title.startswith(SHEET_NAME) and ws["A1"].value == TITLE


def has_sheet(xlsx_bytes: bytes) -> bool:
    """Whether a saved workbook carries a sheet this module made — read from
    the zip, without loading the workbook."""
    try:
        z = zipfile.ZipFile(io.BytesIO(xlsx_bytes))
        names = z.read("xl/workbook.xml").decode("utf-8", "ignore")
        if f'name="{SHEET_NAME}' not in names:
            return False
        texts = (z.read(n).decode("utf-8", "ignore") for n in z.namelist()
                 if n == "xl/sharedStrings.xml" or n.startswith("xl/worksheets/sheet"))
        return any(TITLE in t for t in texts)
    except Exception:  # noqa: BLE001 — not a workbook: nothing of ours in it
        return False


def stamp_of(xlsx_bytes: bytes) -> Optional[str]:
    """The stamp the sheet was built with ('2|compared'), or None (built
    before stamps, or not ours) — read from the zip."""
    try:
        xml = zipfile.ZipFile(io.BytesIO(xlsx_bytes)).read("docProps/custom.xml").decode("utf-8")
    except Exception:  # noqa: BLE001 — no custom properties at all
        return None
    m = re.search(r'<property\b[^>]*\bname="' + re.escape(_STAMP)
                  + r'"[^>]*>\s*<[^>]*>([^<]*)<', xml)
    return m.group(1) if m else None


def is_current(xlsx_bytes: bytes) -> bool:
    """Built with today's layout, and compared: nothing to rebuild."""
    return stamp_of(xlsx_bytes) == f"{LAYOUT}|compared"


def add_sheet(xlsx_bytes: bytes, result: dict, about: Optional[dict] = None,
              replace: bool = False) -> bytes:
    """Append the "Previous submission" sheet to a rendered workbook. With
    `replace`, it takes the place of the one already there (an output
    downloaded after it was made), and every formula keeps the result saved
    with it. Never raises: the plain sheet if the styled one fails, else the
    input back."""
    for build in (_styled, _plain_sheet):
        try:
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(xlsx_bytes))
            if replace:
                for ws in [w for w in wb.worksheets if _is_ours(w)]:
                    wb.remove(ws)
            build(wb.create_sheet(_free_name(wb.sheetnames)), result, about)
            from openpyxl.packaging.custom import StringProperty
            props = wb.custom_doc_props
            if _STAMP in props.names:
                del props[_STAMP]
            props.append(StringProperty(name=_STAMP, value=f"{LAYOUT}|{result.get('state')}"))
            out = io.BytesIO()
            wb.save(out)
            return _keep_formula_results(xlsx_bytes, out.getvalue()) if replace else out.getvalue()
        except Exception:  # noqa: BLE001 — the reconciliation never costs the file
            continue
    return xlsx_bytes


def _sheet_paths(z: zipfile.ZipFile) -> dict[str, str]:
    """{sheet name: 'xl/worksheets/sheetN.xml'} from workbook.xml and its rels
    (attribute order inside an element is not fixed)."""
    wbxml = z.read("xl/workbook.xml").decode("utf-8")
    rels = z.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    target = {}
    for rel in re.findall(r"<Relationship\b[^>]*>", rels):
        rid, tgt = re.search(r'\bId="([^"]+)"', rel), re.search(r'\bTarget="([^"]+)"', rel)
        if rid and tgt:
            t = tgt.group(1)
            target[rid.group(1)] = t.lstrip("/") if t.startswith("/") else "xl/" + t
    out = {}
    for m in re.finditer(r"<sheet\b[^>]*>", wbxml):
        name, rid = re.search(r'\bname="([^"]+)"', m.group(0)), re.search(r'\br:id="([^"]+)"', m.group(0))
        if name and rid and rid.group(1) in target:
            out[name.group(1)] = target[rid.group(1)]
    return out


# A formula with its saved result, and one whose result openpyxl dropped.
_WITH_RESULT = re.compile(r'<c\b[^>]*?\br="([A-Z]+\d+)"[^>]*>(?:<f\b[^>]*>[^<]*</f>)<v>([^<]+)</v></c>')
_NO_RESULT = re.compile(r'(<c\b[^>]*?\br="([A-Z]+\d+)"[^>]*><f\b[^>]*>[^<]*</f>)<v\s*(?:/>|></v>)(</c>)')


def _keep_formula_results(before: bytes, after: bytes) -> bytes:
    """openpyxl keeps a formula but drops its saved result. The renderer
    writes the totals' results into the file after its last openpyxl save
    (direct_render.inject_summary_formulas), so a re-saved output would show
    empty totals in any viewer that reads saved results. Put each result back,
    by sheet name and cell."""
    zb, za = zipfile.ZipFile(io.BytesIO(before)), zipfile.ZipFile(io.BytesIO(after))
    pb, pa = _sheet_paths(zb), _sheet_paths(za)
    patched = {}
    for name, path in pb.items():
        if name not in pa:
            continue
        saved = {m.group(1): m.group(2) for m in _WITH_RESULT.finditer(zb.read(path).decode("utf-8"))}
        if not saved:
            continue
        xml = za.read(pa[name]).decode("utf-8")
        patched[pa[name]] = _NO_RESULT.sub(
            lambda m: (f"{m.group(1)}<v>{saved[m.group(2)]}</v>{m.group(3)}"
                       if m.group(2) in saved else m.group(0)), xml).encode("utf-8")
    if not patched:
        return after
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in za.infolist():
            zout.writestr(item, patched.get(item.filename) or za.read(item.filename))
    return buf.getvalue()


def to_csv(result: dict) -> str:
    """The row-by-row list as CSV: one line per value that differs (every
    line, no cap), after a line with the plain-words summary."""
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([TITLE, (result.get("summary") or {}).get("plain") or ""])
    w.writerow(list(LINE_COLUMNS))
    fields = money_fields(result)
    for row in sorted_rows(result):
        for line in value_lines(row, fields):
            vals = line_values(row, line)
            w.writerow(["" if v is None else v.isoformat() if isinstance(v, date) else v
                        for v in vals])
    return buf.getvalue()
