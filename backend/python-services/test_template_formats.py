"""The output BDX written into a template: a heading row formatted as dates.

The Risksmith workbook has 88 of its 93 headings formatted as dates and no
styling on its data rows, so each data cell took its style from the heading —
date format included. 2,000 then read as 22/06/1905, and 10,000,000 (past
Excel's last date) became #VALUE! once the file was opened again to paint
exceptions.

Run:  pytest test_template_formats.py      (no network, no database)
"""
from __future__ import annotations

import io
import os
from datetime import date, datetime

# exporter (where the style copying lives) builds an AI client when imported;
# nothing here calls it.
os.environ.setdefault("GEMINI_API_KEY", "test-not-used")

from openpyxl import Workbook, load_workbook
from openpyxl.styles.numbers import is_date_format

import direct_render

HEADINGS = ["Policy", "Effective Date", "Gross Premium", "Attachment Point"]


def _template() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "ITD (H)"
    for i, h in enumerate(HEADINGS, start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.number_format = "mm-dd-yy"          # as in the Risksmith workbook
    for r in range(2, 5):                      # sample rows, no styling of their own
        for i in range(1, len(HEADINGS) + 1):
            ws.cell(row=r, column=i, value=f"old {r}")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


STRUCTURE = {"sheets": [{
    "sheet_name": "ITD (H)", "header_row": 0, "data_start_row": 1,
    "columns": [{"column_index": i, "column_name": h} for i, h in enumerate(HEADINGS)],
}]}

ROW = {"Policy": "23-XSP-0698", "Effective Date": "2023-12-08",
       "Gross Premium": 2000, "Attachment Point": 10_000_000}


def _reopened(xlsx: bytes):
    # Opened and saved again, as the exception painting does after rendering.
    wb = load_workbook(io.BytesIO(xlsx))
    buf = io.BytesIO()
    wb.save(buf)
    return load_workbook(io.BytesIO(buf.getvalue()))["ITD (H)"]


def test_a_headings_date_format_never_reaches_a_number():
    out = direct_render._render_with_template(STRUCTURE, {"ITD (H)": [ROW]}, _template())
    ws = _reopened(out)
    premium, attachment = ws["C2"], ws["D2"]
    assert premium.value == 2000 and not is_date_format(premium.number_format)
    # past Excel's last date: kept as a number, not turned into #VALUE!
    assert attachment.value == 10_000_000 and attachment.data_type == "n"
    assert ws["A2"].value == "23-XSP-0698"


def test_a_real_date_still_reads_as_a_date():
    out = direct_render._render_with_template(STRUCTURE, {"ITD (H)": [ROW]}, _template())
    ws = _reopened(out)
    eff = ws["B2"]
    assert isinstance(eff.value, (date, datetime)) and is_date_format(eff.number_format)
    assert eff.value.date() == date(2023, 12, 8)
