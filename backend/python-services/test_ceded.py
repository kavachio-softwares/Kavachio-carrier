"""
Reinsurance (ceded) bordereaux — built from the inward bordereaux a carrier has
processed, per treaty and period, split by reinsurer, sent, and answered.
"""
import io
import json
import os
os.environ["MAIL_ALLOWED_RECIPIENTS"] = "nobody@example.invalid"
os.environ.setdefault("APP_BASE_URL", "http://localhost:5173")

import openpyxl
import pytest
from fastapi.testclient import TestClient

import ceded_engine as ce
import main
import treaty_email
import treaty_routes
from auth_tokens import mint_access_token
from db import (AppUser, CededBordereau, ExportTemplate, LandingRecord,
                OutputExport, Party, PartyContact, Program, SessionLocal, Tenant)
from intake_models import FileArrival

client = TestClient(main.app)
PERIOD = "2027-01"
COLUMNS = [
    ("Policy", "policy_number"), ("Sum Insured", "policy_sum_insured_amount"),
    ("Gross Premium", "premium_transaction_gross_premium_amount"),
    ("Claim", "claim_number"), ("Incurred", "claim_total_incurred_amount"),
]


def _xlsx(rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "BDX"
    ws.append([h for h, _f in COLUMNS])
    for r in rows:
        ws.append(list(r))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


STRUCTURE = {"sheets": [{"sheet_name": "BDX", "header_row": 0, "data_start_row": 1,
                         "columns": [{"column_index": i, "column_name": h,
                                      "canonical_field": f}
                                     for i, (h, f) in enumerate(COLUMNS)]}]}


@pytest.fixture(scope="module")
def w():
    sfx = os.urandom(4).hex()
    with SessionLocal() as s:
        car = Tenant(tenant_name=f"ceded-{sfx}", legal_name="Insurisk Specialty")
        s.add(car); s.commit()
        ids = {}
        for key, name in (("swiss", "Swiss Re"), ("munich", "Munich Re")):
            pt = Party(tenant_id=car.id, party_type="reinsurer", legal_name=name,
                       reference=f"cd-{key}-{sfx}", is_active=True)
            s.add(pt); s.commit()
            ids[key] = pt.id
            s.add(PartyContact(party_id=pt.id, tenant_id=car.id,
                               full_name=f"{name} UW", email=f"{key}-{sfx}@x.invalid"))
        br = Party(tenant_id=car.id, party_type="broker", legal_name="CRC",
                   reference=f"cd-b-{sfx}", is_active=True)
        s.add(br); s.commit()
        pr = Program(tenant_id=car.id, name="Spectrum Property", is_app_managed=True)
        s.add(pr); s.commit()
        tmpl = ExportTemplate(tenant_id=car.id, name=f"tmpl-{sfx}",
                              structure=STRUCTURE, is_active=1)
        s.add(tmpl); s.commit()
        f = (s.query(AppUser).filter(AppUser.role == "kavachio_admin")
             .order_by(AppUser.id).first())
        dana = AppUser(tenant_id=car.id, email=f"d-{sfx}@x.test", full_name="Dana",
                       role="carrier_admin", invited_by_user_id=f.id)
        s.add(dana); s.commit()
        car.owner_user_id = dana.id
        s.commit()
        return {"t": car.id, "b": br.id, "p": pr.id, "tmpl": tmpl.id, **ids,
                "ch": {"Authorization":
                       f"Bearer {mint_access_token(dana.id, car.id, 'carrier_admin')}"}}


def _inward(w, rows, period=PERIOD, landing=None):
    """A processed inward bordereau: its rendered file, the landing that points
    at it, and the arrival that gives it its period."""
    with SessionLocal() as s:
        oe = OutputExport(tenant_id=w["t"], filename="bdx.xlsx", blob=_xlsx(rows),
                          template_id=w["tmpl"], program_id=w["p"],
                          broker_party_id=w["b"], status="clean")
        s.add(oe); s.commit()
        if landing is None:
            lr = LandingRecord(tenant_id=w["t"], data={}, output_export_id=oe.id)
            s.add(lr); s.commit()
            landing = lr.id
        else:
            s.get(LandingRecord, landing).output_export_id = oe.id
            s.commit()
        s.add(FileArrival(tenant_id=w["t"], filename="bdx.xlsx", outcome="accepted",
                          run_export_id=oe.id, run_landing_id=landing,
                          reporting_period=period))
        s.commit()
        return oe.id, landing


def _treaty(w, structure, name, *, setup=True, **setup_over):
    r = client.post("/contracts", headers=w["ch"], json={
        "program_id": w["p"], "contract_type": "insurer_reinsurer", "name": name,
        "inception_dt": "2027-01-01", "expiry_dt": "2027-12-31",
        "year_of_account": "2027", "notice_period_days": 90, "create_as": "draft",
        "treaty": structure})
    assert r.status_code == 200, r.text
    cid = r.json()["id"]
    if setup:
        _setup(w, cid, **setup_over)
    return cid


def _setup(w, cid, activate=True, **over):
    d = client.get(f"/reinsurance/setups/{cid}", headers=w["ch"]).json()
    body = {**d["defaults"], **over}
    r = client.put(f"/reinsurance/setups/{cid}", headers=w["ch"], json=body)
    assert r.status_code == 200, r.text
    if activate:
        r = client.post(f"/reinsurance/setups/{cid}/activate", headers=w["ch"])
        assert r.status_code == 200, r.text
    return r.json()


def _qs(w, pct=40):
    return {"treaty_type": "quota_share",
            "terms": {"currency": "USD", "cession_pct": pct, "ceding_commission_pct": 30},
            "layers": [{"name": "QS", "shares": [
                {"reinsurer_party_id": w["swiss"], "share_pct": 60, "is_lead": True},
                {"reinsurer_party_id": w["munich"], "share_pct": 40}]}]}


# ── the engine ──────────────────────────────────────────────────────────────
def test_quota_share_cedes_premium_and_commission():
    st = {"treaty_type": "quota_share",
          "terms": {"cession_pct": 40, "ceding_commission_pct": 30},
          "layers": [{"name": "QS", "shares": [
              {"reinsurer_party_id": 1, "share_pct": 60},
              {"reinsurer_party_id": 2, "share_pct": 40}]}]}
    rows = [{"policy_number": "P1", "premium_transaction_gross_premium_amount": 1000}]
    out = ce.compute(st, rows)
    ln = out["lines"][0]
    assert (ln["ceded_premium"], ln["ceding_commission"], ln["net_premium"]) == (400, 120, 280)
    split = ce.by_reinsurer(st, out["lines"])
    assert split[1]["totals"]["ceded_premium"] == 240
    assert split[2]["totals"]["net_premium"] == 112


def test_surplus_cedes_above_the_line_by_sum_insured():
    st = {"treaty_type": "surplus", "terms": {"retention_line": 1_000_000, "lines": 4},
          "layers": [{"name": "S", "shares": [{"reinsurer_party_id": 1, "share_pct": 100}]}]}
    rows = [
        {"policy_number": "BIG", "policy_sum_insured_amount": 4_000_000,
         "premium_transaction_gross_premium_amount": 8000},      # 75% ceded
        {"policy_number": "SMALL", "policy_sum_insured_amount": 500_000,
         "premium_transaction_gross_premium_amount": 900},       # within line
        {"policy_number": "NOSI", "premium_transaction_gross_premium_amount": 10},
    ]
    out = ce.compute(st, rows)
    assert [ln["policy"] for ln in out["lines"]] == ["BIG"]
    assert out["lines"][0]["ceded_pct"] == 75 and out["lines"][0]["ceded_premium"] == 6000
    assert {s["policy"] for s in out["skipped"]} == {"SMALL", "NOSI"}


def test_xol_recovers_claims_above_each_layer():
    st = {"treaty_type": "xol", "terms": {"cover_basis": "per_risk"},
          "layers": [
              {"name": "L1", "terms": {"retention": 100, "limit": 400},
               "shares": [{"reinsurer_party_id": 1, "share_pct": 50}]},
              {"name": "L2", "terms": {"retention": 500, "limit": 1000},
               "shares": [{"reinsurer_party_id": 1, "share_pct": 100}]}]}
    rows = [{"claim_number": "C1", "claim_total_incurred_amount": 800}]
    out = ce.compute(st, rows)
    rec = {ln["layer"]: ln["recovery"] for ln in out["lines"]}
    assert rec == {"L1": 400, "L2": 300}
    assert ce.by_reinsurer(st, out["lines"])[1]["totals"]["recovery"] == 500
    # A claim reaching two layers is one claim, not two.
    assert out["totals"].get("claims_paid") is None


def test_stop_loss_recovers_above_the_loss_ratio():
    st = {"treaty_type": "stop_loss", "terms": {"measure": "loss_ratio"},
          "layers": [{"name": "SL", "terms": {"attachment_pct": 70, "limit_pct": 30},
                      "shares": [{"reinsurer_party_id": 1, "share_pct": 100}]}]}
    rows = [{"premium_transaction_gross_premium_amount": 1000,
             "claim_total_incurred_amount": 900}]          # 90% loss ratio
    out = ce.compute(st, rows)
    assert out["lines"][0]["recovery"] == 200               # 20 points of 1000


def test_facultative_takes_only_its_policy():
    st = {"treaty_type": "facultative",
          "terms": {"risk_reference": "POL-9", "share_ceded_pct": 50},
          "layers": [{"name": "F", "shares": [{"reinsurer_party_id": 1, "share_pct": 100}]}]}
    rows = [{"policy_number": "POL-9", "premium_transaction_gross_premium_amount": 200},
            {"policy_number": "POL-1", "premium_transaction_gross_premium_amount": 999}]
    out = ce.compute(st, rows)
    assert [ln["policy"] for ln in out["lines"]] == ["POL-9"]
    assert out["lines"][0]["ceded_premium"] == 100


def test_a_processed_file_is_read_back_by_data_model_field():
    rows, notes = ce.read_export_rows(STRUCTURE, _xlsx([("P1", 1000, 50, None, None)]),
                                      "bdx.xlsx")
    assert notes == []
    assert rows == [{"policy_number": "P1", "policy_sum_insured_amount": 1000,
                     "premium_transaction_gross_premium_amount": 50,
                     # …and every column by its own heading, for exact copies.
                     "_col": {"Policy": "P1", "Sum Insured": 1000, "Gross Premium": 50},
                     # …and which column gave each field, for a value fixed on Exceptions.
                     "_from": {"policy_number": "Policy", "policy_sum_insured_amount": "Sum Insured",
                               "premium_transaction_gross_premium_amount": "Gross Premium"}}]


def test_period_shapes():
    assert ce.period_bounds("2027-02")[1].day == 28
    assert ce.period_bounds("2027-Q2")[0].month == 4
    assert ce.period_bounds("Week 3") is None


# ── the flow ────────────────────────────────────────────────────────────────
@pytest.fixture
def mail(monkeypatch):
    sent = []

    def fake(to, subject, html, text, attachments=None, cc=None):
        sent.append({"to": to, "subject": subject, "attachments": attachments or [],
                     "cc": cc or []})
        return True
    monkeypatch.setattr(treaty_email, "safe_send", fake)
    monkeypatch.setattr(treaty_routes, "_send_link_email",
                        lambda s, link, c, t, code, **kw: sent.append(
                            {"link": (link.id, t, code)}) or True)
    return sent


def test_prepared_automatically_sent_and_answered(w, mail):
    cid = _treaty(w, _qs(w), "QS 2027")
    _export, landing = _inward(w, [("P1", 500_000, 1000, None, None),
                                   ("P2", 200_000, 500, None, None)])

    lst = client.get("/reinsurance/bordereaux", headers=w["ch"]).json()
    mine = [i for i in lst["items"] if i["contract_id"] == cid]
    assert len(mine) == 1 and mine[0]["status"] == "draft" and mine[0]["auto_prepared"]
    cb_id = mine[0]["id"]

    d = client.get(f"/reinsurance/bordereaux/{cb_id}", headers=w["ch"]).json()
    assert d["totals"]["ceded_premium"] == 600                # 40% of 1500
    swiss = next(x for x in d["statements"] if x["reinsurer_party_id"] == w["swiss"])
    assert swiss["totals"]["ceded_premium"] == 360            # 60% of 600
    assert swiss["totals"]["net_premium"] == 252              # less 30% commission
    f = client.get(f"/reinsurance/bordereaux/{cb_id}/statements/{swiss['id']}/file",
                   headers=w["ch"])
    assert f.status_code == 200 and f.content[:2] == b"PK"

    r = client.post(f"/reinsurance/bordereaux/{cb_id}/send", headers=w["ch"], json={})
    assert r.status_code == 200, r.text
    assert r.json()["status"] == "sent"
    emails = [m for m in mail if "attachments" in m]
    assert len(emails) == 2 and all(m["attachments"][0][1][:2] == b"PK" for m in emails)

    # The reinsurer answers on its private link.
    client.post(f"/contracts/{cid}/reinsurer-access", headers=w["ch"],
                json={"contact_ids": [c["id"] for g in client.get(
                    f"/contracts/{cid}/reinsurer-access", headers=w["ch"]).json()["contacts"]
                    if g["party_id"] == w["swiss"] for c in g["contacts"]]})
    lid, token, code = next(m["link"] for m in reversed(mail) if "link" in m)
    sess = client.post(f"/treaty-view/{token}/verify", json={"code": code}).json()["session"]
    h = {"X-Treaty-Session": sess}
    view = client.get(f"/treaty-view/{token}", headers=h).json()
    assert [x["period"] for x in view["statements"]] == [PERIOD]
    sid = view["statements"][0]["id"]
    assert client.get(f"/treaty-view/{token}/statements/{sid}/file",
                      headers=h).content[:2] == b"PK"
    r = client.post(f"/treaty-view/{token}/statements/{sid}/query", headers=h,
                    json={"name": "Sam", "note": "P2 is not ours"})
    assert r.status_code == 200
    assert client.get(f"/reinsurance/bordereaux/{cb_id}",
                      headers=w["ch"]).json()["status"] == "queried"
    client.post(f"/treaty-view/{token}/statements/{sid}/agree", headers=h,
                json={"name": "Sam"})
    d = client.get(f"/reinsurance/bordereaux/{cb_id}", headers=w["ch"]).json()
    assert d["status"] == "sent" and d["agreed"] == 1          # Munich still to answer

    # A corrected inward file after sending makes the next version.
    _inward(w, [("P1", 500_000, 1000, None, None)], landing=landing)
    lst = client.get("/reinsurance/bordereaux", headers=w["ch"]).json()
    mine = sorted([i for i in lst["items"] if i["contract_id"] == cid],
                  key=lambda i: i["version_no"])
    assert [i["version_no"] for i in mine] == [1, 2]
    v2 = mine[1]
    assert v2["status"] == "draft" and v2["is_correction"]
    d2 = client.get(f"/reinsurance/bordereaux/{v2['id']}", headers=w["ch"]).json()
    sw2 = next(x for x in d2["statements"] if x["reinsurer_party_id"] == w["swiss"])
    assert sw2["previous_totals"]["ceded_premium"] == 360
    assert sw2["totals"]["ceded_premium"] == 240
    client.post(f"/reinsurance/bordereaux/{v2['id']}/send", headers=w["ch"], json={})
    with SessionLocal() as s:
        assert s.get(CededBordereau, cb_id).status == "superseded"

    dash = client.get("/reinsurance/dashboard", headers=w["ch"]).json()["ceded"]
    assert dash["sent"] >= 1


def test_a_period_outside_the_treaty_term_is_not_prepared(w, mail):
    cid = _treaty(w, _qs(w, 25), "QS term check")
    _inward(w, [("P9", 100, 100, None, None)], period="2026-06")
    lst = client.get("/reinsurance/bordereaux", headers=w["ch"]).json()
    assert not any(i["contract_id"] == cid and i["period"] == "2026-06"
                   for i in lst["items"])


def test_a_draft_can_be_prepared_by_hand_and_refreshed(w, mail):
    cid = _treaty(w, _qs(w, 10), "QS manual")
    periods = client.get(f"/reinsurance/bordereaux/periods/{cid}",
                         headers=w["ch"]).json()["periods"]
    assert PERIOD in periods
    r = client.post("/reinsurance/bordereaux", headers=w["ch"],
                    json={"contract_id": cid, "period": PERIOD})
    assert r.status_code == 200, r.text
    again = client.post("/reinsurance/bordereaux", headers=w["ch"],
                        json={"contract_id": cid, "period": PERIOD}).json()
    assert again["id"] == r.json()["id"] and again["version_no"] == 1



# ── the setup ───────────────────────────────────────────────────────────────
def test_nothing_is_prepared_until_the_treaty_is_set_up(w, mail):
    cid = _treaty(w, _qs(w, 15), "QS no setup", setup=False)
    _inward(w, [("Q1", 1, 100, None, None)])
    lst = client.get("/reinsurance/bordereaux", headers=w["ch"]).json()
    assert not any(i["contract_id"] == cid for i in lst["items"])
    assert any(t["contract_id"] == cid for t in lst["needs_setup"])
    r = client.post("/reinsurance/bordereaux", headers=w["ch"],
                    json={"contract_id": cid, "period": PERIOD})
    assert r.status_code == 409 and "set up" in r.json()["detail"]
    # A draft setup is not enough either.
    _setup(w, cid, activate=False)
    assert client.post("/reinsurance/bordereaux", headers=w["ch"],
                       json={"contract_id": cid, "period": PERIOD}).status_code == 409


def test_a_carrier_user_sends_the_setup_for_approval(w):
    cid = _treaty(w, _qs(w, 12), "QS approval", setup=False)
    with SessionLocal() as s:
        f = s.query(AppUser).filter(AppUser.role == "kavachio_admin").order_by(AppUser.id).first()
        u = AppUser(tenant_id=w["t"], email=f"u-{os.urandom(3).hex()}@x.test",
                    full_name="Uma", role="carrier_admin", invited_by_user_id=f.id)
        s.add(u); s.commit()
        uh = {"Authorization": f"Bearer {mint_access_token(u.id, w['t'], 'carrier_admin')}"}
    d = client.get(f"/reinsurance/setups/{cid}", headers=uh).json()
    client.put(f"/reinsurance/setups/{cid}", headers=uh, json=d["defaults"])
    r = client.post(f"/reinsurance/setups/{cid}/activate", headers=uh).json()
    assert r["setup"]["status"] == "pending_approval"
    r = client.post(f"/reinsurance/setups/{cid}/send-back", headers=w["ch"],
                    json={"note": "Quarterly, please"}).json()
    assert r["setup"]["status"] == "draft" and r["setup"]["note"] == "Quarterly, please"


def test_quarterly_setup_groups_the_months(w, mail):
    cid = _treaty(w, _qs(w, 20), "QS quarterly", frequency="quarterly")
    _inward(w, [("M2", 1, 1000, None, None)], period="2027-02")
    periods = client.get(f"/reinsurance/bordereaux/periods/{cid}",
                         headers=w["ch"]).json()["periods"]
    assert periods == ["2027-Q1"]
    d = client.post("/reinsurance/bordereaux", headers=w["ch"],
                    json={"contract_id": cid, "period": "2027-Q1"}).json()
    assert {s_["period"] for s_ in d["sources"]} >= {"2027-01", "2027-02"}


def test_setup_layout_recipients_and_csv(w, mail):
    with SessionLocal() as s:
        swiss_contacts = [c.id for c in s.query(PartyContact).filter(
            PartyContact.party_id == w["swiss"])]
        extra = PartyContact(party_id=w["swiss"], tenant_id=w["t"], full_name="Other",
                             email=f"other-{os.urandom(3).hex()}@x.invalid")
        s.add(extra); s.commit()
    cid = _treaty(w, _qs(w, 30), "QS csv", file_format="csv",
                  columns=[{"key": "policy", "heading": "Policy No"},
                           {"key": "ceded_premium", "heading": "Ceded"}],
                  recipients={str(w["swiss"]): swiss_contacts},
                  cc=["team@insurisk.invalid"])
    d = client.post("/reinsurance/bordereaux", headers=w["ch"],
                    json={"contract_id": cid, "period": PERIOD}).json()
    sw = next(x for x in d["statements"] if x["reinsurer_party_id"] == w["swiss"])
    f = client.get(f"/reinsurance/bordereaux/{d['id']}/statements/{sw['id']}/file",
                   headers=w["ch"]).content.decode("utf-8-sig")
    assert f.splitlines()[0] == "Policy No,Ceded,Your share %"
    mail.clear()
    client.post(f"/reinsurance/bordereaux/{d['id']}/send", headers=w["ch"], json={})
    swiss_mail = [m for m in mail if "attachments" in m and "swiss" in m["to"]]
    assert len(swiss_mail) == len(swiss_contacts)      # not the extra contact
    assert all(m["cc"] == ["team@insurisk.invalid"] for m in mail if "attachments" in m)


def test_automatic_sending(w, mail):
    cid = _treaty(w, _qs(w, 35), "QS auto", send_mode="auto")
    mail.clear()
    lst = client.get("/reinsurance/bordereaux", headers=w["ch"]).json()
    mine = [i for i in lst["items"] if i["contract_id"] == cid]
    assert mine and mine[0]["status"] == "sent"
    assert any("attachments" in m for m in mail)


def test_preview_saves_nothing(w):
    cid = _treaty(w, _qs(w, 45), "QS preview", setup=False)
    d = client.get(f"/reinsurance/setups/{cid}", headers=w["ch"]).json()
    r = client.post(f"/reinsurance/setups/{cid}/preview", headers=w["ch"],
                    json={**d["defaults"], "period": PERIOD})
    assert r.status_code == 200, r.text
    assert r.json()["period"] == PERIOD and r.json()["totals"]["ceded_premium"]
    assert client.get(f"/reinsurance/setups/{cid}", headers=w["ch"]).json()["setup"] is None


# ── the carrier's own layout ────────────────────────────────────────────────
def _no_ai(monkeypatch, pins=None):
    """No model call in tests: the headings the inward files share, the rules
    for reinsurance headings, and any pin given here decide the mapping."""
    import ceded_layout as cl

    def fake(structure):
        for sh in structure.get("sheets") or []:
            for col in sh.get("columns") or []:
                col["canonical_field"] = (pins or {}).get(col.get("column_name"))
    monkeypatch.setattr(cl, "map_structure", fake)


def test_carrier_uploads_its_own_excel_layout(w, mail, monkeypatch):
    from openpyxl.styles import Font
    _no_ai(monkeypatch)
    _inward(w, [("POL-L1", 1_000_000, 10_000, "CLM-9", 2_500)], period="2027-05")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "RI Statement"
    ws.append(["Policy", "Sum Insured", "Incurred", "Ceded Premium",
               "RI Commission", "Reinsurer share", "Office use"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FF0000")
    ws.append(["sample", 1, 2, 3, 4, 5, "x"])
    wb.create_sheet("Read me").append(["Fill one row per policy"])
    buf = io.BytesIO(); wb.save(buf)

    cid = _treaty(w, _qs(w, 50), "QS own layout", setup=False)
    r = client.post(f"/reinsurance/setups/{cid}/layout/upload", headers=w["ch"],
                    files={"file": ("ri.xlsx", buf.getvalue())})
    assert r.status_code == 200, r.text
    prop = r.json()
    keys = {c["heading"]: c["key"] for c in prop["columns"]}
    # A heading the brokers' bordereaux also have is copied from that exact
    # column; reinsurance figures are worked out.
    assert keys == {"Policy": "col:Policy", "Sum Insured": "col:Sum Insured",
                    "Incurred": "col:Incurred",
                    "Ceded Premium": "ceded_premium",
                    "RI Commission": "ceding_commission",
                    "Reinsurer share": "share_pct", "Office use": ""}
    assert prop["unfilled"] == ["Office use"]
    assert prop["layout"]["sheet_name"] == "RI Statement"

    d = client.get(f"/reinsurance/setups/{cid}", headers=w["ch"]).json()
    _setup(w, cid, **{**d["defaults"], "layout": prop["layout"],
                      "columns": prop["columns"], "layout_file": prop["layout_file"]})
    got = client.get(f"/reinsurance/setups/{cid}", headers=w["ch"]).json()["setup"]
    assert got["layout"]["source"] == "uploaded" and got["has_layout_file"]

    cb = client.post("/reinsurance/bordereaux", headers=w["ch"],
                     json={"contract_id": cid, "period": "2027-05"}).json()
    sw = next(x for x in cb["statements"] if x["reinsurer_party_id"] == w["swiss"])
    assert sw["filename"].endswith(".xlsx")
    data = client.get(f"/reinsurance/bordereaux/{cb['id']}/statements/{sw['id']}/file",
                      headers=w["ch"]).content
    out = openpyxl.load_workbook(io.BytesIO(data))
    assert out.sheetnames[:2] == ["RI Statement", "Read me"]
    assert "Reinsurance summary" in out.sheetnames
    rs = out["RI Statement"]
    assert [c.value for c in rs[1]] == ["Policy", "Sum Insured", "Incurred",
                                        "Ceded Premium", "RI Commission",
                                        "Reinsurer share", "Office use"]
    assert rs["A1"].font.color.rgb.endswith("FF0000")         # their styling kept
    rows = [[c.value for c in r] for r in rs.iter_rows(min_row=2)]
    assert "sample" not in [r[0] for r in rows]                # sample row cleared
    mine = next(r for r in rows if r[0] == "POL-L1")
    # 50% ceded, Swiss takes 60% of it: 10,000 × 0.5 × 0.6
    assert mine[1] == 1_000_000 and mine[2] == 2500 and mine[3] == 3000
    assert mine[4] == 900 and mine[5] == 60 and mine[6] is None


def test_lloyds_risk_sheet_layout(w, mail, monkeypatch):
    import reporting_standards as standards
    if not standards.discover():
        pytest.skip("no reporting standard bundled")
    _no_ai(monkeypatch, {"Certificate Ref": "policy_number",
                         "Sum Insured Amount": "policy_sum_insured_amount"})
    cid = _treaty(w, _qs(w, 25), "QS lloyds", setup=False)
    r = client.post(f"/reinsurance/setups/{cid}/layout/standard", headers=w["ch"],
                    json={"jurisdiction": "UK"})
    assert r.status_code == 200, r.text
    prop = r.json()
    cols = prop["columns"]
    by = {c["heading"]: c for c in cols}
    assert by["Certificate Ref"]["key"] == "policy" and by["Certificate Ref"]["keep"]
    assert by["Coverholder Name"]["required"] and by["Coverholder Name"]["keep"]
    assert by["Coverholder Name"]["key"] == "broker"
    assert by["Reporting Period (End Date)"]["key"] == "period_end"
    assert not by["Basement"]["keep"]                      # nothing to fill it with
    assert by["Commission %"]["key"] != "ceding_commission"  # the coverholder's, not ceded
    added = [c["key"] for c in cols if c.get("added")]
    assert {"ceded_premium", "ceding_commission", "net_premium", "share_pct"} <= set(added)

    d = client.get(f"/reinsurance/setups/{cid}", headers=w["ch"]).json()
    body = {**d["defaults"], "layout": prop["layout"], "layout_file": prop["layout_file"]}
    bad = [{**c, "keep": False} if c["heading"] == "Coverholder Name" else c for c in cols]
    r = client.put(f"/reinsurance/setups/{cid}", headers=w["ch"], json={**body, "columns": bad})
    assert r.status_code == 400 and "mandatory" in r.text
    _setup(w, cid, **{**body, "columns": cols})

    cb = client.post("/reinsurance/bordereaux", headers=w["ch"],
                     json={"contract_id": cid, "period": PERIOD}).json()
    sw = cb["statements"][0]
    data = client.get(f"/reinsurance/bordereaux/{cb['id']}/statements/{sw['id']}/file",
                      headers=w["ch"]).content
    ws = openpyxl.load_workbook(io.BytesIO(data))["UK"]
    heads = [c.value for c in ws[1]]
    assert heads == [c["heading"] for c in cols if c["keep"]]
    assert "Basement" not in heads and heads[-1] == "Your share %"
    first = dict(zip(heads, [c.value for c in ws[2]]))
    assert first["Certificate Ref"] and first["Ceded premium"] is not None
    assert first["Coverholder Name"] == "CRC"
    assert str(first["Reporting Period (End Date)"]).startswith("2027-01-31")


# ── the carrier's own input file ────────────────────────────────────────────
def _own_xlsx(rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Book"
    ws.append(["Pol Ref", "Client", "GWP", "TSI", "Notes"])
    for r in rows:
        ws.append(list(r))
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()


def test_a_setup_reads_the_carriers_own_file(w, mail, monkeypatch):
    _no_ai(monkeypatch, {"Pol Ref": "policy_number", "Client": "policyholder_legal_name",
                         "GWP": "premium_transaction_gross_premium_amount"})
    cid = _treaty(w, _qs(w, 20), "QS own input", setup=False)
    r = client.post(f"/reinsurance/setups/{cid}/input/upload", headers=w["ch"],
                    files={"file": ("book.xlsx", _own_xlsx([("S-1", "Acme", 100, 5, "x")]))})
    assert r.status_code == 200, r.text
    inp = r.json()
    by = {c["name"]: c["field"] for c in inp["columns"]}
    assert by["GWP"] == "premium_transaction_gross_premium_amount" and by["TSI"] is None
    # The user ties TSI themselves.
    for c in inp["columns"]:
        if c["name"] == "TSI":
            c["field"] = "policy_sum_insured_amount"
    d = client.get(f"/reinsurance/setups/{cid}", headers=w["ch"]).json()
    _setup(w, cid, **{**d["defaults"], "input_layout": inp})
    got = client.get(f"/reinsurance/setups/{cid}", headers=w["ch"]).json()
    assert got["setup"]["input_layout"]["source"] == "own"
    assert got["options"]["inputs"][0]["own"] is True
    assert "policy_sum_insured_amount" in {f["key"][6:] for f in got["options"]["fields"]}

    # Nothing is prepared from the brokers' bordereaux for this treaty.
    lst = client.get("/reinsurance/bordereaux", headers=w["ch"]).json()
    assert not [i for i in lst["items"] if i["contract_id"] == cid]
    t = next(x for x in lst["treaties"] if x["contract_id"] == cid)
    assert t["input_source"] == "own"

    r = client.post("/reinsurance/bordereaux/input", headers=w["ch"],
                    data={"contract_id": str(cid), "period": "2027-06"},
                    files={"file": ("june.xlsx", _own_xlsx([
                        ("S-1", "Acme", 1000, 50000, ""), ("S-2", "Beta", 3000, 90000, "")]))})
    assert r.status_code == 200, r.text
    cb = r.json()
    assert cb["period"] == "2027-06" and cb["totals"]["ceded_premium"] == 800
    assert cb["sources"][0]["filename"] == "june.xlsx"
    assert {ln["insured"] for ln in cb["lines"]} == {"Acme", "Beta"}

    # A corrected file for the same period refreshes the draft.
    r = client.post("/reinsurance/bordereaux/input", headers=w["ch"],
                    data={"contract_id": str(cid), "period": "2027-06"},
                    files={"file": ("june-v2.xlsx", _own_xlsx([("S-1", "Acme", 2000, 1, "")]))})
    assert r.json()["id"] == cb["id"] and r.json()["totals"]["ceded_premium"] == 400

    bad = client.post("/reinsurance/bordereaux/input", headers=w["ch"],
                      data={"contract_id": str(cid), "period": "June"},
                      files={"file": ("x.xlsx", _own_xlsx([("S", "A", 1, 1, "")]))})
    assert bad.status_code == 400


def test_own_file_mapping_by_column(w, mail, monkeypatch):
    """Mapping is column to column: one input column can fill two outputs, and
    a column no standard field fits is copied across as it is."""
    _no_ai(monkeypatch, {"GWP": "premium_transaction_gross_premium_amount"})
    cid = _treaty(w, _qs(w, 10), "QS own mapping", setup=False)
    inp = client.post(f"/reinsurance/setups/{cid}/input/upload", headers=w["ch"],
                      files={"file": ("b.xlsx", _own_xlsx([("P", "A", 1, 1, "n")]))}).json()
    cols = {c["name"]: c for c in inp["columns"]}
    # "Pol Ref" fills Policy AND, as an extra entry, the insured name; "Notes"
    # is copied across under a field of its own.
    cols["Pol Ref"]["field"] = "policy_number"
    extra = {**cols["Pol Ref"], "field": "policyholder_legal_name"}
    cols["Notes"]["field"] = "x_notes"
    inp["columns"] = list(cols.values()) + [extra]
    # A layout of the carrier's own (here a CSV of four columns), where any
    # input column can be copied across.
    lay = client.post(f"/reinsurance/setups/{cid}/layout/upload", headers=w["ch"],
                      files={"file": ("out.csv", b"Policy,Insured,Notes,Ceded\n")}).json()
    keys = {"Policy": "policy", "Insured": "insured", "Notes": "field:x_notes",
            "Ceded": "ceded_premium"}
    out_cols = [{**c, "key": keys[c["heading"]], "keep": True} for c in lay["columns"]]
    d = client.get(f"/reinsurance/setups/{cid}", headers=w["ch"]).json()
    _setup(w, cid, **{**d["defaults"], "input_layout": inp, "layout": lay["layout"],
                      "layout_file": lay["layout_file"], "columns": out_cols})
    r = client.post("/reinsurance/bordereaux/input", headers=w["ch"],
                    data={"contract_id": str(cid), "period": "2027-07"},
                    files={"file": ("jul.xlsx", _own_xlsx([("Z-9", "ignored", 500, 1, "hello")]))})
    assert r.status_code == 200, r.text
    cb = r.json()
    st = cb["statements"][0]
    f = client.get(f"/reinsurance/bordereaux/{cb['id']}/statements/{st['id']}/file",
                   headers=w["ch"]).content
    lines = f.decode("utf-8-sig").splitlines()
    row = dict(zip(lines[0].split(","), lines[1].split(",")))
    assert row["Policy"] == "Z-9" and row["Insured"] == "Z-9"
    assert row["Notes"] == "hello" and float(row["Ceded"]) == 30.0   # 500 × 10% × 60%


def test_check_saves_nothing(w, mail, monkeypatch):
    _no_ai(monkeypatch, {"Pol Ref": "policy_number",
                         "GWP": "premium_transaction_gross_premium_amount"})
    cid = _treaty(w, _qs(w, 50), "QS check", setup=False)
    inp = client.post(f"/reinsurance/setups/{cid}/input/upload", headers=w["ch"],
                      files={"file": ("b.xlsx", _own_xlsx([("P", "A", 1, 1, "")]))}).json()
    d = client.get(f"/reinsurance/setups/{cid}", headers=w["ch"]).json()
    _setup(w, cid, **{**d["defaults"], "input_layout": inp})
    r = client.post("/reinsurance/bordereaux/check", headers=w["ch"],
                    data={"contract_id": str(cid), "period": "2027-08"},
                    files={"file": ("aug.xlsx", _own_xlsx([("C-1", "A", 400, 1, "")]))})
    assert r.status_code == 200, r.text
    assert r.json()["totals"]["ceded_premium"] == 200 and r.json()["table"]["rows"]
    with SessionLocal() as s:
        from db import CededInput
        assert s.query(CededInput).filter(CededInput.contract_id == cid).count() == 0
        assert s.query(CededBordereau).filter(CededBordereau.contract_id == cid).count() == 0
    lst = client.get("/reinsurance/bordereaux", headers=w["ch"]).json()
    t = next(x for x in lst["treaties"] if x["contract_id"] == cid)
    assert t["input_file"] == "b.xlsx" and t["output"] == "kavachio"


def test_a_brokers_setup_takes_an_input_file_too(w, mail, monkeypatch):
    """Like a broker run: drop a file in for a period and the statement is
    worked out from it, its columns matched to the brokers' headings."""
    _no_ai(monkeypatch)
    cid = _treaty(w, _qs(w, 10), "QS file over brokers")        # reads the brokers
    wb = openpyxl.Workbook(); ws = wb.active
    ws.append(["Policy", "Sum Insured", "Gross Premium"])       # the brokers' headings
    ws.append(["F-1", 100000, 2000])
    buf = io.BytesIO(); wb.save(buf)
    r = client.post("/reinsurance/bordereaux/check", headers=w["ch"],
                    data={"contract_id": str(cid), "period": "2027-09"},
                    files={"file": ("sep.xlsx", buf.getvalue())})
    assert r.status_code == 200, r.text
    assert r.json()["totals"]["ceded_premium"] == 200
    r = client.post("/reinsurance/bordereaux/input", headers=w["ch"],
                    data={"contract_id": str(cid), "period": "2027-09"},
                    files={"file": ("sep.xlsx", buf.getvalue())})
    assert r.status_code == 200, r.text
    cb = r.json()
    assert cb["totals"]["ceded_premium"] == 200 and cb["sources"][0]["filename"] == "sep.xlsx"
    assert "2027-09" in client.get(f"/reinsurance/bordereaux/periods/{cid}",
                                   headers=w["ch"]).json()["periods"]


def test_send_attaches_each_reinsurers_bdx_file(w, mail):
    cid = _treaty(w, _qs(w, 15), "QS send file")
    d = client.post("/reinsurance/bordereaux", headers=w["ch"],
                    json={"contract_id": cid, "period": PERIOD}).json()
    plan = {p["reinsurer_party_id"]: p for p in d["send_plan"]}
    assert set(plan) == {w["swiss"], w["munich"]}
    assert all(p["attachment"].endswith(".xlsx") and p["attachment_bytes"] > 0
               and p["to"] for p in plan.values())
    mail.clear()
    client.post(f"/reinsurance/bordereaux/{d['id']}/send", headers=w["ch"], json={})
    sent = [m for m in mail if m.get("attachments")]
    assert {a[0] for m in sent for a in m["attachments"]} == {p["attachment"] for p in plan.values()}
    for m in sent:
        name, data, mime = m["attachments"][0]
        assert data[:2] == b"PK" and "spreadsheetml" in mime     # a real .xlsx


# ── the treaty's rules ──────────────────────────────────────────────────────
def test_the_treatys_terms_give_its_rules():
    import datetime as _d
    import treaty_rules as trr
    st = {"treaty_type": "quota_share",
          "terms": {"currency": "usd", "cession_pct": 40, "treaty_limit": 2_000_000}}
    rules = {r["id"]: r for r in trr.derive(st, _d.date(2027, 1, 1), _d.date(2027, 12, 31))}
    assert {"term", "currency", "required_premium", "limit_per_risk", "duplicate"} <= set(rules)
    assert rules["limit_per_risk"]["params"]["value"] == 5_000_000      # 2m ÷ 40%
    rows = [{"policy_number": "A", "policy_effective_date": "2026-06-01",
             "premium_transaction_gross_premium_amount": 10,
             "premium_transaction_original_currency": "GBP",
             "policy_sum_insured_amount": 9_000_000},
            {"policy_number": "A", "policy_effective_date": "2027-02-01"}]
    got = {c["id"]: c for c in trr.check(list(rules.values()), rows)}
    assert got["term"]["rows"] == 1 and got["currency"]["rows"] == 1
    assert got["required_premium"]["rows"] == 1 and got["limit_per_risk"]["severity"] == "warning"
    assert got["duplicate"]["rows"] == 1


def test_a_blocking_rule_stops_send_until_set_to_warning(w, mail):
    import treaty_rules as trr
    cid = _treaty(w, _qs(w, 12), "QS rules")
    # Every row of the period must be on policy "NOPE" — none is.
    d = client.get(f"/reinsurance/setups/{cid}", headers=w["ch"]).json()
    rule = {"id": "wording_1_policy_allowed", "label": "Only policy NOPE",
            "kind": "allowed", "subject": "policy", "params": {"values": ["NOPE"]},
            "severity": "critical", "enabled": True, "source": "wording", "clause": "x"}
    base = {k: v for k, v in d["setup"].items()
            if k in ("frequency", "due_days", "program_ids", "columns", "layout",
                     "input_layout", "file_format", "recipients", "cc", "send_mode")}
    _setup(w, cid, **{**base, "rules": d["setup"]["rules"] + [rule]})
    cb = client.post("/reinsurance/bordereaux", headers=w["ch"],
                     json={"contract_id": cid, "period": PERIOD}).json()
    assert cb["send_blocked_by"] == ["Only policy NOPE"] and cb["can_send"] is False
    assert client.post(f"/reinsurance/bordereaux/{cb['id']}/send", headers=w["ch"],
                       json={}).status_code == 409
    # Set to Warning: shown, and sent.
    _setup(w, cid, **{**base, "rules": d["setup"]["rules"] + [{**rule, "severity": "warning"}]})
    cb = client.post("/reinsurance/bordereaux", headers=w["ch"],
                     json={"contract_id": cid, "period": PERIOD}).json()
    assert cb["can_send"] is True and trr.critical(cb["checks"]) == []
    assert any(c["id"] == rule["id"] for c in cb["checks"])


def test_rules_are_read_from_the_treaty_wording(w, monkeypatch):
    import json as _j
    import treaty_rules as trr
    real = trr.extract
    answer = _j.dumps({"rules": [
        {"kind": "excluded", "subject": "territory", "values": ["Cuba", "Iran"],
         "label": "Sanctioned territories are excluded",
         "clause": "This treaty excludes risks in Cuba and Iran."},
        {"kind": "max", "subject": "sum_insured", "value": 25000000,
         "label": "At most 25m any one risk", "clause": "Maximum any one risk 25,000,000."},
        {"kind": "bogus", "subject": "territory", "values": ["x"]}]})
    monkeypatch.setattr(trr, "extract", lambda text, call=None: real(text, call=lambda p: answer))
    cid = _treaty(w, _qs(w, 11), "QS wording rules", setup=False)
    with SessionLocal() as s:
        from db import Contract
        c = s.get(Contract, cid)
        c.wording_sections = {"sections": [{"key": "excl", "title": "Exclusions",
                                            "body": "This treaty excludes risks in Cuba and Iran."}]}
        s.commit()
    r = client.post(f"/reinsurance/setups/{cid}/rules/extract", headers=w["ch"])
    assert r.status_code == 200, r.text
    got = r.json()
    assert got["found"] == 2 and got["read_from"]
    wording = [x for x in got["rules"] if x["source"] == "wording"]
    assert {x["kind"] for x in wording} == {"excluded", "max"}
    assert wording[0]["clause"]
    assert any(x["id"] == "term" for x in got["rules"])     # the terms rules stay


def test_a_draft_follows_its_setup_and_a_sent_one_does_not(w, mail):
    cid = _treaty(w, _qs(w, 13), "QS follows setup",
                  columns=[{"key": "policy", "heading": "Policy"},
                           {"key": "ceded_premium", "heading": "Ceded"}])
    cb = client.post("/reinsurance/bordereaux", headers=w["ch"],
                     json={"contract_id": cid, "period": PERIOD}).json()

    def headings(cb_id):
        d = client.get(f"/reinsurance/bordereaux/{cb_id}", headers=w["ch"]).json()
        st = d["statements"][0]
        f = client.get(f"/reinsurance/bordereaux/{cb_id}/statements/{st['id']}/file",
                       headers=w["ch"]).content
        return [c.value for c in openpyxl.load_workbook(io.BytesIO(f))["Lines"][1]], d

    assert headings(cb["id"])[0][:2] == ["Policy", "Ceded"]
    # The setup's layout changes: the draft is rebuilt to it, same id.
    d = client.get(f"/reinsurance/setups/{cid}", headers=w["ch"]).json()
    base = {k: v for k, v in d["setup"].items()
            if k in ("frequency", "due_days", "program_ids", "layout", "input_layout",
                     "file_format", "recipients", "cc", "send_mode", "rules")}
    _setup(w, cid, **{**base, "columns": [{"key": "policy", "heading": "Policy No"},
                                           {"key": "net_premium", "heading": "Net"}]})
    assert client.get(f"/reinsurance/bordereaux/{cb['id']}", headers=w["ch"]).json()["setup_changed"]
    client.get("/reinsurance/bordereaux", headers=w["ch"])          # prepares what is due
    heads, d2 = headings(cb["id"])
    assert heads[:2] == ["Policy No", "Net"] and d2["setup_changed"] is False
    # Sent: a later setup change does not re-issue it.
    client.post(f"/reinsurance/bordereaux/{cb['id']}/send", headers=w["ch"], json={})
    _setup(w, cid, **{**base, "columns": [{"key": "policy", "heading": "P"}]})
    lst = client.get("/reinsurance/bordereaux", headers=w["ch"]).json()
    mine = [i for i in lst["items"] if i["contract_id"] == cid and i["period"] == PERIOD]
    assert len(mine) == 1 and mine[0]["status"] == "sent"



def test_a_copied_column_is_that_exact_column(w, mail, monkeypatch):
    """Several input columns can carry one field (the AI tied every
    'Expiring…' column to the expiry date). A copied output column must read
    the input column it names — never whichever of them came first."""
    import ceded_layout as cl
    st = {"sheets": [{"sheet_name": "S", "header_row": 0, "data_start_row": 1, "columns": [
        {"column_index": 0, "column_name": "Expiring Rating Basis", "canonical_field": "policy_expiration_date"},
        {"column_index": 1, "column_name": "PolicyPeriodExpiryDate", "canonical_field": "policy_expiration_date"}]}]}
    rows = ce._rows_from_table([["Expiring Rating Basis", "PolicyPeriodExpiryDate"],
                                ["Flat", "2027-12-31"]], st["sheets"][0])
    assert rows[0]["policy_expiration_date"] == "Flat"          # first wins by field…
    line = {"src": {"_cols": rows[0]["_col"]}}
    assert cl.value(line, "col:PolicyPeriodExpiryDate") == "2027-12-31"   # …not by column
    _f, _h, names = cl.source_columns([st])
    layout, cols = cl.propose({"sheets": [{"sheet_name": "Out", "columns": [
        {"column_index": 0, "column_name": "PolicyPeriodExpiryDate"}]}]},
        source="uploaded", available=_f, by_heading=_h, treaty_type="quota_share", names=names)
    assert cols[0]["key"] == "col:PolicyPeriodExpiryDate"


def test_a_draft_setup_can_be_deleted_an_active_one_cannot(w):
    cid = _treaty(w, _qs(w, 9), "QS delete draft", activate=False)
    assert client.delete(f"/reinsurance/setups/{cid}", headers=w["ch"]).status_code == 200
    assert client.get(f"/reinsurance/setups/{cid}", headers=w["ch"]).json()["setup"] is None
    cid2 = _treaty(w, _qs(w, 8), "QS delete active")
    r = client.delete(f"/reinsurance/setups/{cid2}", headers=w["ch"])
    assert r.status_code == 409 and "active" in r.text


def test_exceptions_are_decided_row_by_row(w, mail):
    """The broker's exception resolution, for a reinsurance bordereau: each
    row a Critical rule catches is Approved (the value the rule recommends),
    Fixed (a value entered) or Dismissed (kept). Saving records them; Fix &
    Validate prepares the statement with them and checks every row again."""
    period = "2027-03"
    _inward(w, [("PX1", 500_000, 1000, None, None), ("PX2", 900_000, 2000, None, None),
                ("PX3", 100_000, 500, None, None), ("PX4", 700_000, 800, None, None)],
            period=period)
    cid = _treaty(w, _qs(w, 50), "QS exceptions")
    d = client.get(f"/reinsurance/setups/{cid}", headers=w["ch"]).json()
    rule = {"id": "wording_1_si_max", "label": "Sum insured at most 200,000",
            "kind": "max", "subject": "sum_insured", "params": {"value": 200_000},
            "severity": "critical", "enabled": True, "source": "wording", "clause": "x"}
    base = {k: v for k, v in d["setup"].items()
            if k in ("frequency", "due_days", "program_ids", "columns", "layout",
                     "input_layout", "file_format", "recipients", "cc", "send_mode")}
    _setup(w, cid, **{**base, "rules": d["setup"]["rules"] + [rule]})
    cb = client.post("/reinsurance/bordereaux", headers=w["ch"],
                     json={"contract_id": cid, "period": period}).json()
    chk = next(c for c in cb["checks"] if c["id"] == rule["id"])
    assert chk["recommendation"] == "200000"
    items = {it["policy"]: it for it in chk["items"]}
    assert set(items) == {"PX1", "PX2", "PX4"} and chk["open"] == 3
    assert cb["can_send"] is False
    url = f"/reinsurance/bordereaux/{cb['id']}/decide"

    # A fix must be a number for a money field.
    assert client.post(url, headers=w["ch"], json={"decisions": [
        {"key": items["PX2"]["key"], "kind": "fix", "value": "lots"}]}).status_code == 400
    r = client.post(url, headers=w["ch"], json={"decisions": [
        {"key": items["PX1"]["key"], "kind": "approve"},
        {"key": items["PX2"]["key"], "kind": "fix", "value": "150000"},
        {"key": items["PX4"]["key"], "kind": "dismiss", "reason": "Agreed with the lead"}]})
    assert r.status_code == 200, r.text
    cb = r.json()
    # Recorded, not yet in the statement: Send waits for Fix & Validate.
    assert cb["updated"] == 3 and cb["unapplied"] == 3 and cb["can_send"] is False
    got = {it["policy"]: it for it in next(c for c in cb["checks"] if c["id"] == rule["id"])["items"]}
    assert (got["PX1"]["status"], got["PX1"]["new_value"]) == ("approved", "200000")
    assert (got["PX2"]["status"], got["PX2"]["new_value"]) == ("fixed", "150000")
    assert got["PX4"]["status"] == "dismissed" and got["PX4"]["reason"] == "Agreed with the lead"
    assert client.post(f"/reinsurance/bordereaux/{cb['id']}/send", headers=w["ch"],
                       json={}).status_code == 409

    cb = client.post(f"/reinsurance/bordereaux/{cb['id']}/validate", headers=w["ch"]).json()
    assert cb["unapplied"] == 0 and cb["can_send"] is True and cb["send_blocked_by"] == []
    si = {ln["policy"]: ln["sum_insured"] for ln in cb["lines"]}
    assert (si["PX1"], si["PX2"], si["PX4"]) == (200_000, 150_000, 700_000)
    assert any("corrected" in n for n in cb["notes"])

    # A fixed value that still breaks the rule is open again after Fix & Validate.
    client.post(url, headers=w["ch"], json={"decisions": [
        {"key": items["PX2"]["key"], "kind": "fix", "value": "250000"}]})
    cb = client.post(f"/reinsurance/bordereaux/{cb['id']}/validate", headers=w["ch"]).json()
    px2 = next(it for c in cb["checks"] if c["id"] == rule["id"] for it in c["items"]
               if it["policy"] == "PX2")
    assert px2["status"] == "open" and px2["still_fails"] is True
    assert cb["unapplied"] == 0 and cb["can_send"] is False

    # Decisions survive Prepare again; reopening one is recorded until validated.
    cb = client.post("/reinsurance/bordereaux", headers=w["ch"],
                     json={"contract_id": cid, "period": period}).json()
    assert cb["unapplied"] == 0
    cb = client.post(url, headers=w["ch"], json={"decisions": [
        {"key": items["PX1"]["key"], "kind": "reopen"}]}).json()
    assert cb["unapplied"] == 1


def test_a_treatys_agreed_limits_check_its_reinsurance_bordereau(w, mail):
    """The limits typed on a treaty are the treaty's — only those a treaty has
    are kept, they become rules on its reinsurance bordereau, and they are
    never bound as checks on the brokers' files."""
    period = "2027-05"
    _inward(w, [("PL1", 900_000, 1000, None, None), ("PL2", 100_000, 500, None, None)],
            period=period)
    r = client.post("/contracts", headers=w["ch"], json={
        "program_id": w["p"], "contract_type": "insurer_reinsurer", "name": "QS limits",
        "inception_dt": "2027-01-01", "expiry_dt": "2027-12-31",
        "year_of_account": "2027", "notice_period_days": 90, "create_as": "draft",
        "treaty": _qs(w, 50),
        "agreed_limits": {"max_sum_insured": {"value": 500_000, "severity": "critical"},
                          "brokerage_pct": {"value": 5}}})
    assert r.status_code == 200, r.text
    cid = r.json()["id"]
    with SessionLocal() as s:
        from db import Contract
        assert set(s.get(Contract, cid).commercial_terms) == {"max_sum_insured"}
    assert client.post(f"/contracts/{cid}/bind-checks", headers=w["ch"],
                       json={}).status_code == 409
    _setup(w, cid)
    cb = client.post("/reinsurance/bordereaux", headers=w["ch"],
                     json={"contract_id": cid, "period": period}).json()
    chk = next(c for c in cb["checks"] if c["id"] == "limit_max_sum_insured")
    assert chk["severity"] == "critical" and [it["policy"] for it in chk["items"]] == ["PL1"]
    assert chk["recommendation"] == "500000" and cb["can_send"] is False
