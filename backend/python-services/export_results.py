"""
export_results.py
─────────────────
What happened to EVERY row of one checked file, not only the rows that failed
(plan 21–22: the Results tab, the "Results" sheet in the downloaded Excel and
the full CSV). One counting, used by all three, so they always agree.

A row is exactly one of:
  passed       checked, and no rule raised anything on it
  open         at least one exception on it is still waiting for a decision
  resolved     it had exceptions and every one has been decided
               (approved, fixed, dismissed or rejected)
  not_checked  the rules never looked at it: the whole file could not be
               validated, or the row is blank, a total line or a repeated
               header (the same row classifier the run uses to set rows aside)

Exceptions are counted the way the Exception Triage screen counts them
(broker_tally.countable / settled): row-less notices are left out, and a
file- or policy-level total is counted on the row it is pinned to.

No database and no web imports, so it can be tested on its own.
"""
from __future__ import annotations

import io
import re
from typing import Iterable, Optional

from broker_tally import countable, settled

OUTCOMES = ("passed", "open", "resolved", "not_checked")
OUTCOME_WORDS = {
    "passed": "Passed",
    "open": "Failed — open",
    "resolved": "Resolved",
    "not_checked": "Not checked",
}
# Sheet name of the counts added to the downloaded Excel. Never a template's
# own sheet name in practice; a clash is avoided by suffixing (see _free_name).
RESULTS_SHEET = "Results"
RESULTS_TITLE = "Validation results"
# Report sheets added after the template's own (never data): this module's
# counts and the Previous Submission Reconciliation (version_diff_sheet.TITLE).
REPORT_TITLES = (RESULTS_TITLE, "Previous submission reconciliation")

_STATUS_WORDS = {
    "approved": "approved", "fixed": "fixed", "dismissed": "dismissed",
    "rejected": "rejected", "resolved": "resolved",
}


def _norm(s) -> str:
    return " ".join(str(s or "").split()).strip().lower()


def _row_int(v) -> Optional[int]:
    try:
        return int(v)
    except (TypeError, ValueError):
        return None


_POLICY_RE = re.compile(r"\bpolicy\b.*\b(no|number|num|ref|reference|id)\b")
_CERT_RE = re.compile(r"\b(certificate|cert)\b.*\b(no|number|num|ref|reference|id)\b")


def policy_column(cols: Iterable[str]) -> Optional[str]:
    """The column that holds the policy number, by its heading. None if none."""
    cols = [c for c in cols if c]
    for c in cols:
        if _norm(c) in ("policy number", "policy no", "policy no.", "policy ref",
                        "policy reference", "policy id", "policynumber", "policy"):
            return c
    # A policy reference before a certificate one ("Certificate Ref" comes
    # first in the Lloyd's layout, "Policy or Group Ref" is the policy).
    for rx in (_POLICY_RE, _CERT_RE):
        for c in cols:
            if rx.search(_norm(c).replace(".", " ")):
                return c
    return None


def apply_decisions(exceptions: list, decisions: Iterable[dict]) -> list:
    """Stamp each exception with its saved decision (landing_correction rows:
    output_sheet, output_row, output_field, kind), the way the review screen
    attaches them when it reads an export. Returns new dicts."""
    kinds = {"approve": "approved", "fix": "fixed", "dismiss": "dismissed",
             "reject": "rejected"}
    by_key = {}
    for d in decisions or ():
        by_key[(d.get("output_sheet"), _row_int(d.get("output_row")),
                d.get("output_field"))] = d
    if not by_key:
        return list(exceptions or [])
    out = []
    for e in exceptions or []:
        if not isinstance(e, dict):
            out.append(e)
            continue
        field = e.get("field") or e.get("column") or e.get("field_path")
        hit = by_key.get((e.get("sheet"), _row_int(e.get("row")), field))
        if hit:
            e = dict(e)
            e["status"] = kinds.get(hit.get("kind"), "resolved")
            e["resolution_note"] = hit.get("reason")
        out.append(e)
    return out


def _not_data_rows(records: list[dict], cols: list[str]) -> set[int]:
    """1-based positions of rows the run would set aside (blank, total line,
    repeated header). Fail-safe: nothing is set aside if the classifier fails."""
    if not records or not cols:
        return set()
    try:
        import row_classifier
        _data, excluded = row_classifier.classify_sheet_rows(records, cols)
        return {int(x["position"]) for x in excluded}
    except Exception:  # noqa: BLE001 — a count, never what breaks a download
        return set()


def compute(sheets: dict[str, list[dict]], exceptions: list, *,
            not_validated: bool = False,
            columns: Optional[dict[str, list[str]]] = None) -> dict:
    """Every row's outcome and the counts.

    `sheets` is {sheet name: [row dict, …]} in file order — row N of a sheet is
    the exception's `row` N (the run's __rowid). `columns` gives each sheet's
    column order (defaults to the first row's keys)."""
    excs = countable(exceptions)
    by_row: dict[tuple, list] = {}
    for e in excs:
        r = _row_int(e.get("row"))
        if r is None:
            continue
        by_row.setdefault((_norm(e.get("sheet")), r), []).append(e)

    rows_out: list[dict] = []
    totals = dict.fromkeys(OUTCOMES, 0)
    rules: dict[tuple, dict] = {}
    multi = len([n for n, rs in sheets.items() if rs]) > 1
    for sheet, records in sheets.items():
        cols = list((columns or {}).get(sheet) or (records[0].keys() if records else []))
        pol_col = policy_column(cols)
        set_aside = set() if not_validated else _not_data_rows(records, cols)
        for i, rec in enumerate(records, start=1):
            mine = by_row.get((_norm(sheet), i), [])
            pol = next((e.get("policy_number") for e in mine if e.get("policy_number")), None)
            if pol is None and pol_col:
                v = rec.get(pol_col)
                pol = None if v in (None, "") else str(v)
            if not_validated or (i in set_aside and not mine):
                outcome = "not_checked"
            elif not mine:
                outcome = "passed"
            elif all(settled(e) for e in mine):
                outcome = "resolved"
            else:
                outcome = "open"
            totals[outcome] += 1
            items = []
            for e in mine:
                done = settled(e)
                st = str(e.get("status") or "").lower()
                items.append({
                    "rule": e.get("rule_name") or e.get("code") or "Check",
                    "field": e.get("field") or e.get("column") or e.get("field_path"),
                    "severity": e.get("severity") or "warning",
                    "message": e.get("message") or e.get("reason"),
                    "actual": e.get("actual_value"),
                    "status": _STATUS_WORDS.get(st, "resolved") if done else "open",
                    "decided_by": e.get("decided_by"),
                    "decided_at": e.get("decided_at"),
                })
                key = (items[-1]["rule"], items[-1]["severity"])
                agg = rules.setdefault(key, {"rule": key[0], "severity": key[1],
                                             "exceptions": 0, "open": 0, "resolved": 0})
                agg["exceptions"] += 1
                agg["resolved" if done else "open"] += 1
            rows_out.append({"sheet": sheet if multi else None, "sheet_name": sheet,
                             "row": i, "policy_number": pol, "outcome": outcome,
                             "exceptions": items})

    n_exc = sum(r["exceptions"] for r in rules.values())
    n_open = sum(r["open"] for r in rules.values())
    summary = {
        "total": len(rows_out),
        "passed": totals["passed"],
        "open": totals["open"],
        "resolved": totals["resolved"],
        "not_checked": totals["not_checked"],
        "exceptions": n_exc,
        "exceptions_open": n_open,
        "exceptions_resolved": n_exc - n_open,
    }
    by_rule = sorted(rules.values(), key=lambda r: (-r["open"], -r["exceptions"], r["rule"]))
    return {"summary": summary, "rows": rows_out, "by_rule": by_rule}


def result_text(row: dict) -> str:
    """One cell's worth: what happened to this row, in words."""
    n = len(row["exceptions"])
    if row["outcome"] == "passed":
        return "Passed — no exceptions"
    if row["outcome"] == "not_checked":
        return "Not checked"
    word = "exception" if n == 1 else "exceptions"
    if row["outcome"] == "resolved":
        return f"Resolved — {n} {word}, all decided"
    still = sum(1 for x in row["exceptions"] if x["status"] == "open")
    return f"Failed — {n} {word}, {still} open"


def exceptions_text(row: dict) -> str:
    parts = []
    for x in row["exceptions"]:
        bit = x["rule"]
        if x.get("field"):
            bit += f" ({x['field']})"
        bit += f": {x['status']}"
        parts.append(bit)
    return "; ".join(parts)


# ── the downloads ───────────────────────────────────────────────────────────

def _free_name(existing: list[str]) -> str:
    taken = {n.lower() for n in existing}
    name, n = RESULTS_SHEET, 2
    while name.lower() in taken:
        name, n = f"{RESULTS_SHEET} {n}", n + 1
    return name


# The Results sheet's look: a calm report. Navy header rows, light-grey cell
# borders, faint zebra rows and soft status tints. None of these fills is a
# flag colour (exporter's light red FFC7CE / light orange FFE0B2), which the
# in-site viewers read as exception marks.
_RS_FONT = "Calibri"
_RS_NAVY = "1F3A5F"
_RS_LINE = "D0D7DE"
_RS_ZEBRA = "F6F8FA"
_RS_MUTED = "5F6368"
_RS_TINT = {                       # outcome -> (soft fill, text)
    "passed": ("E6F4EA", "1E7B34"),
    "open": ("FDECEA", "B3261E"),
    "resolved": ("E8F0FE", "1A56DB"),
    "not_checked": ("F1F3F4", "5F6368"),
}
_RS_SEVERITY = {"critical": ("B3261E", True), "warning": ("B26A00", True),
                "info": ("5F6368", False)}


def _rs_width(text) -> float:
    """Rough width of `text` in Excel column units (Calibri 11, where a digit
    is one unit): capitals and digits about one, most other letters less."""
    return sum(1.0 if ch.isupper() or ch.isdigit() or ch in "mw%@" else 0.8
               for ch in str(text or ""))


def _results_about(about) -> list[str]:
    """The grey line(s) under the title, from whatever `about` holds (file
    name, version, when it was checked …): "Label: value" parts joined by
    dots, wrapped to a readable length."""
    from datetime import date, datetime
    parts = []
    for k, v in (about.items() if hasattr(about, "items") else ()):
        if v is None or (isinstance(v, str) and not v.strip()):
            continue
        if isinstance(v, datetime):
            v = v.strftime("%d %b %Y, %H:%M") + (f" {v.tzname()}" if v.tzinfo else "")
        elif isinstance(v, date):
            v = v.strftime("%d %b %Y")
        label = " ".join(str(k or "").replace("_", " ").split())
        if label.islower():
            label = label[:1].upper() + label[1:]
        text = " ".join(str(v).split())
        parts.append(f"{label}: {text}" if label else text)
    lines, cur = [], ""
    for p in parts:
        if cur and len(cur) + len(p) > 100:
            lines.append(cur)
            cur = ""
        cur = f"{cur}   ·   {p}" if cur else p
    return lines + ([cur] if cur else [])


def _results_sheet_styled(ws, res: dict, about) -> None:
    """The counts as a small report: the title, the `about` line, then three
    bordered tables (rows, exceptions, per check), each closed by a total."""
    import math
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.page import PageMargins

    s = res["summary"]
    checks = [(str(r["rule"]), str(r["severity"]).capitalize(), r["exceptions"],
               r["open"], r["resolved"]) for r in (res.get("by_rule") or [])]

    def share(n, whole):           # a fraction for '0.0%'; blank, not an error, on 0
        return n / whole if whole else None

    def font(**kw):
        return Font(name=_RS_FONT, **{"size": 11, **kw})

    def fill(rgb):
        return PatternFill("solid", fgColor=rgb)

    thin = Side(style="thin", color=_RS_LINE)
    head_rule = Side(style="medium", color=_RS_NAVY)
    total_rule = Side(style="double", color=_RS_NAVY)

    def edges(top=thin, bottom=thin):
        # A shared edge is set the same on both cells, so every app draws it alike.
        return Border(left=thin, right=thin, top=top, bottom=bottom)

    formats = {"int": "#,##0", "pct": "0.0%"}
    # The check names set column A; beyond 50 they wrap, so the sheet still
    # prints on one portrait page without shrinking.
    width_a = min(50, max(30, round(max([_rs_width(c[0]) for c in checks], default=0)) + 4))

    def put(r, j, value, kind, **style):
        c = ws.cell(row=r, column=j, value=value)
        words = kind in ("text", "severity")
        wrap = words and j == 1 and _rs_width(value) > width_a - 3
        c.alignment = Alignment(horizontal="left" if words else "right",
                                vertical="center", indent=1, wrap_text=wrap)
        if kind in formats and not isinstance(value, str):
            c.number_format = formats[kind]
        for k, v in style.items():
            setattr(c, k, v)
        return wrap

    def table(r, head, kinds, body, total):
        """One table from row `r`: navy header, body rows [(values, outcome
        tint or None)] striped, a bold total under a double rule. Returns
        the row after it."""
        for j, (h, kind) in enumerate(zip(head, kinds), start=1):
            put(r, j, h, kind, font=font(bold=True, color="FFFFFF"),
                fill=fill(_RS_NAVY), border=edges(bottom=head_rule))
        ws.row_dimensions[r].height = 22
        for i, (values, outcome) in enumerate(body):
            r += 1
            lines = 1
            border = edges(top=head_rule if i == 0 else thin,
                           bottom=total_rule if i == len(body) - 1 else thin)
            for j, (v, kind) in enumerate(zip(values, kinds), start=1):
                style = {"font": font(), "border": border}
                if i % 2:
                    style["fill"] = fill(_RS_ZEBRA)
                if j == 1 and outcome in _RS_TINT:
                    bg, ink = _RS_TINT[outcome]
                    style.update(fill=fill(bg), font=font(color=ink))
                if kind == "severity":
                    ink, bold = _RS_SEVERITY.get(str(v).lower(), (None, False))
                    style["font"] = font(color=ink, bold=bold)
                if put(r, j, v, kind, **style):
                    lines = math.ceil(_rs_width(v) / (width_a - 3))
            ws.row_dimensions[r].height = max(20, 15 * lines + 5)
        r += 1
        for j, (v, kind) in enumerate(zip(total, kinds), start=1):
            put(r, j, v, kind, font=font(bold=True), border=edges(top=total_rule))
        ws.row_dimensions[r].height = 21
        return r + 1

    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = _RS_NAVY
    ws["A1"] = RESULTS_TITLE
    ws["A1"].font = font(size=16, bold=True, color=_RS_NAVY)
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 28
    r = 2
    try:
        about_lines = _results_about(about)
    except Exception:  # noqa: BLE001 — an odd `about` costs the line, not the sheet
        about_lines = []
    for line in about_lines:
        ws.cell(row=r, column=1, value=line).font = font(size=10, color=_RS_MUTED)
        ws.row_dimensions[r].height = 16
        r += 1
    ws.row_dimensions[r].height = 10
    r += 1

    n = s["total"]
    r = table(r, ("Rows", "Count", "% of rows"), ("text", "int", "pct"),
              [((OUTCOME_WORDS[k], s[k], share(s[k], n)), k)
               for k in ("passed", "open", "resolved", "not_checked")],
              ("Total rows", n, share(n, n)))
    r += 1
    n = s["exceptions"]
    r = table(r, ("Exceptions", "Count", "% of exceptions"), ("text", "int", "pct"),
              [(("Open", s["exceptions_open"], share(s["exceptions_open"], n)), "open"),
               (("Resolved", s["exceptions_resolved"],
                 share(s["exceptions_resolved"], n)), "resolved")],
              ("Total", n, share(n, n)))
    if checks:
        r += 1
        table(r, ("Check", "Severity", "Exceptions", "Open", "Resolved"),
              ("text", "severity", "int", "int", "int"),
              [(c, None) for c in checks],
              ("Total", None, sum(c[2] for c in checks), sum(c[3] for c in checks),
               sum(c[4] for c in checks)))

    for col, width in zip("ABCDE", (width_a, 11, 16, 9, 11)):
        ws.column_dimensions[col].width = width
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.gridLines = False
    ws.print_options.horizontalCentered = True
    ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.6, bottom=0.6)
    ws.oddFooter.center.text = "Page &P of &N"
    ws.oddFooter.center.size = 8
    ws.oddFooter.center.color = _RS_MUTED


def _results_sheet_plain(ws, res: dict, about) -> None:
    """The unstyled sheet (bold headings only): the fallback if styling fails."""
    from openpyxl.styles import Font
    bold = Font(bold=True)
    s = res["summary"]
    ws.append([RESULTS_TITLE])
    ws["A1"].font = Font(bold=True, size=13)
    for k, v in (about or {}).items():
        if v not in (None, ""):
            ws.append([k, v])
    ws.append([])
    ws.append(["Rows", "Count"])
    for c in ws[ws.max_row]:
        c.font = bold
    for label, key in (("Total rows", "total"), ("Passed", "passed"),
                       ("Failed — open", "open"), ("Resolved", "resolved"),
                       ("Not checked", "not_checked")):
        ws.append([label, s[key]])
    ws.append([])
    ws.append(["Exceptions", "Count"])
    for c in ws[ws.max_row]:
        c.font = bold
    ws.append(["Total", s["exceptions"]])
    ws.append(["Open", s["exceptions_open"]])
    ws.append(["Resolved", s["exceptions_resolved"]])
    if res.get("by_rule"):
        ws.append([])
        ws.append(["Check", "Severity", "Exceptions", "Open", "Resolved"])
        for c in ws[ws.max_row]:
            c.font = bold
        for r in res["by_rule"]:
            ws.append([r["rule"], str(r["severity"]).capitalize(), r["exceptions"],
                       r["open"], r["resolved"]])
    ws.column_dimensions["A"].width = 48
    for col in "BCDE":
        ws.column_dimensions[col].width = 14


def add_results_sheet(xlsx_bytes: bytes, res: dict, about: Optional[dict] = None) -> bytes:
    """Append a "Results" sheet of COUNTS to a rendered workbook: rows passed,
    failed (open), resolved and not checked, the exceptions, then the same per
    check — laid out as a small report (see _results_sheet_styled). Never
    raises: if the styling fails the plain sheet is written instead, and on
    any other failure the input comes back."""
    for build in (_results_sheet_styled, _results_sheet_plain):
        try:
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(xlsx_bytes))
            build(wb.create_sheet(_free_name(wb.sheetnames)), res, about)
            out = io.BytesIO()
            wb.save(out)
            return out.getvalue()
        except Exception:  # noqa: BLE001 — the counts never cost the file itself
            continue
    return xlsx_bytes


def results_csv(sheets: dict[str, list[dict]], res: dict,
                columns: Optional[dict[str, list[str]]] = None) -> str:
    """Every row of the file with its result: the output columns, then Result
    and Exceptions. Several sheets go one after another with a Sheet column."""
    import csv
    buf = io.StringIO()
    w = csv.writer(buf)
    multi = len([n for n, rs in sheets.items() if rs]) > 1
    by_pos = {(r["sheet_name"], r["row"]): r for r in res["rows"]}
    for sheet, records in sheets.items():
        cols = list((columns or {}).get(sheet) or (records[0].keys() if records else []))
        w.writerow((["Sheet"] if multi else []) + ["Row"] + cols + ["Result", "Exceptions"])
        for i, rec in enumerate(records, start=1):
            r = by_pos.get((sheet, i))
            w.writerow((([sheet] if multi else []) + [i]
                        + [("" if rec.get(c) is None else rec.get(c)) for c in cols]
                        + ([result_text(r), exceptions_text(r)] if r else ["", ""])))
        if multi:
            w.writerow([])
    return buf.getvalue()


def grid_to_sheets(grid: list[dict], data_only: bool = True) -> tuple[dict, dict]:
    """The viewer grid of a stored file ([{sheet, rows:[[header…],[…]…], is_data}])
    as ({sheet: [row dict]}, {sheet: [columns]}). Only data sheets when the grid
    says which ones they are; trailing all-blank rows are dropped."""
    sheets, columns = {}, {}
    flagged = any("is_data" in g for g in grid)
    for g in grid or []:
        if data_only and flagged and not g.get("is_data"):
            continue
        rows = g.get("rows") or []
        if not rows:
            continue
        # The report sheets added to the Excel are not data.
        if rows[0] and str(rows[0][0]).strip() in REPORT_TITLES:
            continue
        header = [str(h) if h not in (None, "") else f"Column {i + 1}"
                  for i, h in enumerate(rows[0])]
        body = [dict(zip(header, r)) for r in rows[1:]]
        while body and all(v in (None, "") for v in body[-1].values()):
            body.pop()
        sheets[g.get("sheet") or "Sheet1"] = body
        columns[g.get("sheet") or "Sheet1"] = header
    return sheets, columns
