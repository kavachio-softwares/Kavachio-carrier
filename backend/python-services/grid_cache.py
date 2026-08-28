"""
grid_cache.py
─────────────
Parse-once / serve-many cache behind the chunked (infinite-scroll) BDX grid.

The problem this solves
───────────────────────
Reading a row window out of a generated xlsx means opening the workbook. The
*styled* open (needed to see which cells carry the light-red invalid fill) is
the expensive part and its cost scales with the whole file, not with the window
requested. Serving a 50k-row sheet ten rows at a time therefore re-parsed the
entire workbook on every scroll — total work grew with how far the user
scrolled, which is the opposite of what chunked delivery is meant to achieve.

So: parse the workbook ONCE into a compact, plain-Python structure, keep it
briefly in process memory, and serve every subsequent chunk by slicing that
structure. The first chunk pays the parse; the rest are effectively free.

Correctness — why the key is the blob's CONTENT, not its id
───────────────────────────────────────────────────────────
`rerender_export` deliberately re-renders IN PLACE (`reuse_export_id`) so an
export's id stays stable across "Fix & re-run". The same export_id can
therefore point at different bytes over time. Keying on export_id alone would
happily serve pre-correction rows after a re-run, so entries are keyed on a
fast digest of the blob itself: new bytes ⇒ new key ⇒ automatic invalidation,
with no cache-busting call sites to remember and nothing to invalidate by hand.

Memory
──────
Row values are cached only while a workbook stays under a cell budget; past
that the entry keeps just the (small) marks/headers/dimensions and row windows
are streamed from the blob in read-only mode. Either way the expensive styled
parse happens once. Entries are bounded by count and age and evicted
least-recently-used, so a long-lived process can't accumulate workbooks.

Everything here is derived from the blob that was passed in — no database
access, no schema assumptions, no per-tenant or per-file special cases.
"""
from __future__ import annotations

import hashlib
import io
import os
import re
import threading
import time
import xml.etree.ElementTree as ET
import zipfile
from collections import OrderedDict

# ── tunables (env-overridable; the defaults suit a normal deployment) ────────
def _env_int(name: str, default: int) -> int:
    try:
        v = int(os.getenv(name, "").strip())
        return v if v > 0 else default
    except (TypeError, ValueError):
        return default


# How many distinct workbooks stay resident at once.
MAX_ENTRIES = _env_int("KAVACHIO_BDX_CACHE_ENTRIES", 8)
# How long an untouched entry survives. A re-render invalidates by content
# anyway; this only bounds idle memory.
TTL_SECONDS = _env_int("KAVACHIO_BDX_CACHE_TTL_SEC", 900)
# Above this many populated cells, ONE workbook's row values are not held in
# memory (marks and dimensions still are) and its windows stream from the blob
# instead. Sized so an ordinary large BDX still gets the fast path.
MAX_CACHED_CELLS = _env_int("KAVACHIO_BDX_CACHE_MAX_CELLS", 2_000_000)
# Ceiling on row values held across ALL entries at once — entry count alone is
# a poor proxy for memory when one workbook can be a thousand times another.
# Least-recently-used entries are evicted until the cache fits.
MAX_TOTAL_CELLS = _env_int("KAVACHIO_BDX_CACHE_TOTAL_CELLS", 4_000_000)

_INVALID_FILL_FALLBACK = "FFC7CE"
_WARNING_FILL_FALLBACK = "FFE0B2"


def _warning_fill_rgb() -> str:
    """The light-orange fill `highlight_exceptions` paints on non-critical cells.
    Flagged cells carrying it are reported as `warn_marks`, so the in-site grid
    can tint them the same way the downloaded workbook does."""
    try:
        from exporter import _WARNING_FILL_RGB as warn  # noqa: PLC0415
        return str(warn).upper()
    except Exception:
        return _WARNING_FILL_FALLBACK


def _invalid_fill_rgb() -> str:
    try:
        from exporter import _INVALID_FILL_RGB as inv  # noqa: PLC0415
        return str(inv).upper()
    except Exception:
        return _INVALID_FILL_FALLBACK


def fingerprint(blob: bytes) -> str:
    """Content digest used as the cache key — a change-detector, not a security
    boundary.

    md5 specifically, because the same value can be obtained WITHOUT the bytes:
    a caller holding the blob in a database can have the database compute it
    (Postgres `md5(blob)`) and hand it to `open_grid_for`, which then skips
    transferring several megabytes on every already-cached request. Both routes
    must agree on the digest for that to be a cache hit rather than a second,
    identical entry — hence one shared function.
    """
    return hashlib.md5(blob).hexdigest()


# ── fast path: find the flagged cells straight from the xlsx XML ────────────
# openpyxl's styled load builds a Python object for every cell in the file just
# to let us ask each one its fill colour — on a large sheet that is the single
# most expensive thing this module does. The same answer is available far more
# cheaply: styles.xml says which style indices carry the invalid fill, and each
# sheet's XML tags every cell with its style index. Scanning for those is ~24x
# faster and needs a fraction of the memory.
#
# Anything unexpected in the file makes this return None, and the caller
# transparently falls back to the original openpyxl scan — so an odd workbook
# loses the speedup, never the correctness.
_NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
_NS_R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
_CELL_RE = re.compile(rb'<c\s+r="([A-Z]+)(\d+)"([^>]*)>')
_STYLE_RE = re.compile(rb'\ss="(\d+)"')
_ANY_CELL_RE = re.compile(rb'<c[\s/>]')


def _col_index(letters: bytes) -> int:
    n = 0
    for ch in letters:
        n = n * 26 + (ch - 64)  # b'A' == 65
    return n - 1


def _flagged_style_ids(zf, targets: dict):
    """{style index (as bytes) -> kind} for styles painted in a flag colour.

    `targets` maps a kind — 'crit' or 'warn' — to that colour's RGB, so one pass
    over styles.xml classifies both instead of scanning the file twice.
    """
    root = ET.fromstring(zf.read("xl/styles.xml"))
    fills = root.find(f"{_NS}fills")
    if fills is None:
        return {}
    fill_kind = {}
    for i, f in enumerate(fills.findall(f"{_NS}fill")):
        pf = f.find(f"{_NS}patternFill")
        if pf is None or pf.get("patternType") != "solid":
            continue
        fg = pf.find(f"{_NS}fgColor")
        rgb = ((fg.get("rgb") if fg is not None else None) or "").upper()
        for kind, target in targets.items():
            if target and rgb.endswith(target):
                fill_kind[i] = kind
                break
    if not fill_kind:
        return {}
    xfs = root.find(f"{_NS}cellXfs")
    if xfs is None:
        return {}
    out = {}
    for i, xf in enumerate(xfs.findall(f"{_NS}xf")):
        try:
            kind = fill_kind.get(int(xf.get("fillId") or 0))
        except (TypeError, ValueError):
            continue
        if kind:
            out[str(i).encode()] = kind
    return out


def _sheet_xml_paths(zf) -> list:
    """[(sheet name, path to its XML)] in workbook order."""
    wbx = ET.fromstring(zf.read("xl/workbook.xml"))
    rels = ET.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
    rid2target = {r.get("Id"): (r.get("Target") or "") for r in rels}
    names = set(zf.namelist())
    out = []
    sheets_el = wbx.find(f"{_NS}sheets")
    if sheets_el is None:
        raise ValueError("workbook.xml has no <sheets>")
    for sh in sheets_el:
        target = rid2target.get(sh.get(f"{_NS_R}id"), "")
        if not target:
            raise ValueError("sheet without a relationship target")
        cand = target[1:] if target.startswith("/") else (
            target if target.startswith("xl/") else "xl/" + target)
        cand = cand.replace("/./", "/")
        if cand not in names:
            tail = target.rsplit("/", 1)[-1]
            matches = [n for n in names if n.endswith("/" + tail)]
            if len(matches) != 1:
                raise ValueError(f"cannot locate sheet xml for {sh.get('name')!r}")
            cand = matches[0]
        out.append((sh.get("name"), cand))
    return out


def _comments_for_sheet(zf, sheet_path: str) -> dict:
    """{(row, col) -> comment text} for one sheet, via its rels."""
    folder, _, fname = sheet_path.rpartition("/")
    rel_path = f"{folder}/_rels/{fname}.rels"
    if rel_path not in zf.namelist():
        return {}
    rels = ET.fromstring(zf.read(rel_path))
    out = {}
    for r in rels:
        if not (r.get("Type") or "").endswith("/comments"):
            continue
        target = r.get("Target") or ""
        if not target:
            continue
        # A relationship Target is either package-absolute ("/xl/comments/…")
        # or relative to the part's own folder ("../comments/…").
        if target.startswith("/"):
            path = target[1:]
        else:
            path = os.path.normpath(os.path.join(folder, target)).replace(os.sep, "/")
        if path not in zf.namelist():
            continue
        croot = ET.fromstring(zf.read(path))
        for cm in croot.iter(f"{_NS}comment"):
            ref = cm.get("ref") or ""
            m = re.fullmatch(r"([A-Z]+)(\d+)", ref)
            if not m:
                continue
            # openpyxl renders a comment as "text" possibly prefixed by author;
            # join the runs the same way its .text does.
            text = "".join(t.text or "" for t in cm.iter(f"{_NS}t"))
            if text:
                out[(int(m.group(2)) - 1, _col_index(m.group(1).encode()))] = text
    return out


def _scan_marks_xml(blob: bytes, targets: dict):
    """{sheet name: (marks, warn_marks, notes, last_row)} read straight from the
    xlsx, or None if the file's shape isn't one this fast path can vouch for.

    `marks` is EVERY flagged cell whatever its colour; `warn_marks` is the
    non-critical subset. Keeping marks whole means row extents, counts and every
    existing caller behave exactly as before — the colour is extra information
    on the side, never a filter applied to the old one.

    `last_row` is the highest 1-based data row on which the sheet has any cell
    AT ALL, empty or not — the same extent a styled openpyxl load reports, and
    deliberately not the sheet's declared dimension, which a producer may pad
    past the cells it actually wrote.
    """
    try:
        with zipfile.ZipFile(io.BytesIO(blob)) as zf:
            if "xl/styles.xml" not in zf.namelist():
                return None
            flagged = _flagged_style_ids(zf, targets)
            paths = _sheet_xml_paths(zf)
            result = {}
            for name, path in paths:
                data = zf.read(path)
                # Every cell must carry its address for positions to be exact.
                # (openpyxl always writes it; a foreign producer might not.)
                # Counted with finditer, not findall — findall would materialise
                # one match object per cell in the file just to take a length,
                # which on a large sheet costs hundreds of MB for no answer that
                # the running count doesn't already give.
                if (sum(1 for _ in _ANY_CELL_RE.finditer(data))
                        != sum(1 for _ in _CELL_RE.finditer(data))):
                    return None
                marks, warn_marks, notes, last_row = [], [], [], 0
                comments = _comments_for_sheet(zf, path) if flagged else {}
                for m in _CELL_RE.finditer(data):
                    r = int(m.group(2)) - 1
                    if r > last_row:
                        last_row = r
                    if not flagged:
                        continue
                    sm = _STYLE_RE.search(m.group(3))
                    kind = flagged.get(sm.group(1)) if sm is not None else None
                    if kind is None:
                        continue
                    c = _col_index(m.group(1))
                    marks.append([r, c])
                    if kind == "warn":
                        warn_marks.append([r, c])
                    text = comments.get((r, c))
                    if text:
                        notes.append({"r": r, "c": c, "text": text})
                result[name] = (marks, warn_marks, notes, last_row)
            return result
    except Exception:
        return None


class _Entry:
    """One parsed workbook.

    `sheets` maps sheet name → {header, total_rows, marks, warn_marks, notes,
    rows|None}
    where `rows` (when present) maps 1-based data-row index → cell values.
    `order` preserves the workbook's own sheet order.
    """
    __slots__ = ("sheets", "order", "rows_cached", "built_at", "cells")

    def __init__(self, sheets: dict, order: list, rows_cached: bool, cells: int = 0):
        self.sheets = sheets
        self.order = order
        self.rows_cached = rows_cached
        # Row values actually retained — drives the cache-wide memory ceiling.
        self.cells = cells if rows_cached else 0
        self.built_at = time.monotonic()


_lock = threading.RLock()
_cache: "OrderedDict[str, _Entry]" = OrderedDict()
# Observability — surfaced by stats() and used by the test-suite to prove the
# expensive parse really is skipped on a hit.
_stats = {"hits": 0, "misses": 0, "evictions": 0, "expiries": 0, "builds": 0,
          "fastpath": 0, "fallback": 0}


def stats() -> dict:
    with _lock:
        return dict(_stats, entries=len(_cache))


def reset(_for_tests: bool = False) -> None:
    """Drop every entry. Safe at any time — the next request simply rebuilds."""
    with _lock:
        _cache.clear()
        if _for_tests:
            for k in _stats:
                _stats[k] = 0


def _purge_expired_locked() -> None:
    if TTL_SECONDS <= 0:
        return
    now = time.monotonic()
    for key in [k for k, e in _cache.items() if now - e.built_at > TTL_SECONDS]:
        del _cache[key]
        _stats["expiries"] += 1


def _get_locked(key: str):
    _purge_expired_locked()
    entry = _cache.get(key)
    if entry is not None:
        _cache.move_to_end(key)
    return entry


def _put_locked(key: str, entry: _Entry) -> None:
    _cache[key] = entry
    _cache.move_to_end(key)
    while len(_cache) > MAX_ENTRIES:
        _cache.popitem(last=False)
        _stats["evictions"] += 1
    # Then trim by retained row values, oldest first. The entry just inserted is
    # the most recently used, so it is never the one dropped here — a single
    # oversized workbook simply ends up alone in the cache.
    while len(_cache) > 1 and sum(e.cells for e in _cache.values()) > MAX_TOTAL_CELLS:
        _cache.popitem(last=False)
        _stats["evictions"] += 1


def _build(blob: bytes, with_marks: bool, header_row_by_sheet: dict | None = None) -> _Entry:
    """The one pass that populates a cache entry: per sheet, the header, the row
    count, every invalid-fill cell (+ its comment) and — budget permitting — the
    row values themselves.

    Where possible the flagged cells come from the XML fast path and the values
    from a read-only streaming parse, which together avoid materialising a
    styled object for every cell in the workbook. If the fast path can't vouch
    for the file, this falls back to exactly the styled walk the endpoint used
    before, so the result is identical either way.

    `header_row_by_sheet` ({sheet name: 0-based row index}, caller-supplied —
    this module still touches no database itself) overrides which row is the
    HEADER; defaults to 0 (the header is the file's first row) so every export
    whose template's own header sits on row 1 is completely unaffected. A
    template that reused a source workbook with a leading annotation/blank row
    (its header lands on row 2, not row 1) needs this — otherwise that blank
    row is mistaken for the header and the real header text is mistaken for a
    data row. Rows BEFORE the header row are pre-header noise and are skipped
    entirely (not shown as data, not shown as header).
    """
    import openpyxl  # noqa: PLC0415

    inv = _invalid_fill_rgb() if with_marks else None
    warn_rgb = _warning_fill_rgb() if with_marks else None
    prescanned = (_scan_marks_xml(blob, {"crit": inv, "warn": warn_rgb})
                  if with_marks else {})
    if prescanned is not None:
        _stats["fastpath"] += 1
    else:
        _stats["fallback"] += 1
    # Styles only have to be materialised when the fast path declined; values
    # alone stream far more cheaply in read-only mode.
    styled = with_marks and prescanned is None
    wb = openpyxl.load_workbook(io.BytesIO(blob), data_only=True,
                                read_only=not styled)
    try:
        sheets: dict = {}
        order: list = []
        total_cells = 0
        # The budget is enforced DURING the walk, not after it: once the
        # workbook proves too large, already-accumulated row values are released
        # immediately and no further ones are retained, so a huge file never
        # materialises a full copy of itself just to have it discarded. The
        # marks/header/dimension work — the part that made the parse expensive —
        # still completes for every sheet.
        rows_cached = True
        for ws in wb.worksheets:
            order.append(ws.title)
            header: list = []
            rows: dict = {}
            marks: list = []
            warn_marks: list = []
            notes: list = []
            max_gi = 0
            scanned_last = 0
            header_row = (header_row_by_sheet or {}).get(ws.title, 0)

            if styled:
                for ri, row in enumerate(ws.iter_rows()):
                    vals: list = []
                    for ci, cell in enumerate(row):
                        vals.append("" if cell.value is None else cell.value)
                        fill = getattr(cell, "fill", None)
                        ftype = (getattr(fill, "patternType", None)
                                 or getattr(fill, "fill_type", None)) if fill else None
                        if ftype == "solid":
                            rgb = getattr(getattr(fill, "fgColor", None), "rgb", None)
                            rgb = rgb.upper() if isinstance(rgb, str) else ""
                            if rgb.endswith(inv) or rgb.endswith(warn_rgb):
                                marks.append([ri, ci])
                                if rgb.endswith(warn_rgb):
                                    warn_marks.append([ri, ci])
                                cmt = getattr(cell, "comment", None)
                                if cmt is not None and getattr(cmt, "text", None):
                                    notes.append({"r": ri, "c": ci, "text": str(cmt.text)})
                    if ri == header_row:
                        header = vals
                    elif ri > header_row:
                        max_gi = ri
                        if rows_cached:
                            rows[ri] = vals
                    total_cells += len(vals)
                    if rows_cached and total_cells > MAX_CACHED_CELLS:
                        rows_cached = False
                        rows = {}
                        for prev in sheets.values():
                            prev["rows"] = None
            else:
                if with_marks:
                    marks, warn_marks, notes, scanned_last = \
                        prescanned.get(ws.title, ([], [], [], 0))
                for ri, vals_t in enumerate(ws.iter_rows(values_only=True)):
                    vals = ["" if v is None else v for v in vals_t]
                    if ri == header_row:
                        header = vals
                    elif ri > header_row:
                        if any(v is not None for v in vals_t):
                            max_gi = ri
                        if rows_cached:
                            rows[ri] = vals
                    total_cells += len(vals)
                    if rows_cached and total_cells > MAX_CACHED_CELLS:
                        rows_cached = False
                        rows = {}
                        for prev in sheets.values():
                            prev["rows"] = None

            # How far the sheet reaches. Three sources, and the extent is the
            # largest, because each one alone can under-report:
            #
            #  · the last row holding a VALUE — misses rows the exporter created
            #    and formatted but left empty, which a template routinely does
            #    for its fixed-size block. The generated file shows those rows,
            #    so the review grid has to as well, or a sheet Excel opens with
            #    37 rows appears here with one.
            #  · the last row the sheet has any CELL on, from the XML scan —
            #    what the styled read reports, and so what this module has to
            #    agree with. Not the declared dimension: a producer may pad that
            #    past the cells it wrote, and the styled read ignores it too.
            #  · every FLAGGED row — a cell can be painted on a row carrying no
            #    values at all, and such a row must stay addressable or the
            #    Exceptions view would ask for it forever and never finish.
            max_gi = max(max_gi, scanned_last)
            if marks:
                max_gi = max(max_gi, max(r for r, _ in marks))

            # Marks/notes were scanned in ABSOLUTE file rows (both the styled
            # walk and the XML fast path); the served coordinate system is DATA
            # rows (1 = first row after the header), so shift them here — and
            # drop any that land on the header or a pre-header noise row, which
            # are not data and are never served as rows. A no-op when
            # header_row is 0 (every mark's row is already its data row).
            if header_row:
                marks = [[r - header_row, c] for r, c in marks if r > header_row]
                warn_marks = [[r - header_row, c] for r, c in warn_marks
                              if r > header_row]
                notes = [dict(n, r=n["r"] - header_row) for n in notes
                         if n["r"] > header_row]

            sheets[ws.title] = {
                "header": header,
                # Data-row COUNT (not an absolute row index) — the rows before
                # header_row are pre-header noise, not data, so they don't count.
                "total_rows": max(0, max_gi - header_row),
                "header_row": header_row,
                "marks": marks,
                "warn_marks": warn_marks,
                "notes": notes,
                "rows": rows if rows_cached else None,
            }
    finally:
        wb.close()

    if not rows_cached:
        # Too big to hold: keep the cheap parts and stream row windows on demand.
        for sh in sheets.values():
            sh["rows"] = None
    _stats["builds"] += 1
    return _Entry(sheets, order, rows_cached, total_cells)


def _stream_rows(blob: bytes, sheet_name: str, wanted: set) -> dict:
    """Read just the requested 1-based data rows, without styles. Used only for
    workbooks too large to keep row values resident."""
    import openpyxl  # noqa: PLC0415

    if not wanted:
        return {}
    stop_after = max(wanted)
    wb = openpyxl.load_workbook(io.BytesIO(blob), data_only=True, read_only=True)
    try:
        ws = None
        for cand in wb.worksheets:
            if cand.title == sheet_name:
                ws = cand
                break
        if ws is None:
            return {}
        out: dict = {}
        for ri, vals_t in enumerate(ws.iter_rows(values_only=True)):
            if ri in wanted:
                out[ri] = ["" if v is None else v for v in vals_t]
                if ri >= stop_after:
                    break
        return out
    finally:
        wb.close()


def _header_row_key(header_row_by_sheet: dict | None) -> str:
    """Deterministic cache-key suffix for a header-row override — empty for the
    (overwhelmingly common) default so existing keys are completely unchanged."""
    if not header_row_by_sheet:
        return ""
    return ":" + ",".join(f"{k}={v}" for k, v in sorted(header_row_by_sheet.items()))


def _resolve(blob: bytes, with_marks: bool, header_row_by_sheet: dict | None = None) -> _Entry:
    """The cache lookup: the entry for these exact bytes, building it if needed.

    Note the cost of the lookup ITSELF — digesting the blob — is proportional to
    the file, so callers that want many windows of one workbook should resolve
    once (see `open_grid`) rather than per window.
    """
    key = f"{fingerprint(blob)}:{int(bool(with_marks))}{_header_row_key(header_row_by_sheet)}"
    with _lock:
        entry = _get_locked(key)
        if entry is not None:
            _stats["hits"] += 1
        else:
            _stats["misses"] += 1

    if entry is None:
        # Built outside the lock so one slow parse can't stall unrelated
        # requests; a concurrent duplicate build is possible but harmless and
        # self-resolving (last writer wins, identical content either way).
        entry = _build(blob, with_marks, header_row_by_sheet)
        with _lock:
            _put_locked(key, entry)
    return entry


class _LazyBlob:
    """The workbook bytes, fetched at most once and only if something needs them.

    A cache hit whose row values are resident never touches this; that is the
    whole point of `open_grid_for`.
    """
    __slots__ = ("_load", "_blob")

    def __init__(self, load):
        self._load = load
        self._blob = None

    def __call__(self) -> bytes:
        if self._blob is None:
            self._blob = self._load() or b""
        return self._blob


class GridHandle:
    """One resolved workbook, reusable across many row windows.

    Streaming a sheet asks for the same workbook once per chunk. Going through
    `grid_page` each time means re-digesting the whole blob to find the entry
    that is already in hand — cheap per call, but paid thousands of times on a
    large sheet, and it grows with file size rather than window size, which is
    exactly the shape chunked delivery exists to avoid.

    The handle keeps a direct reference to the entry, so a stream also sees ONE
    consistent parse from first row to last even if the cache evicts or the
    export is re-rendered underneath it.
    """
    __slots__ = ("_blob", "_entry")

    def __init__(self, blob, entry: _Entry):
        self._blob = blob if callable(blob) else (lambda b=blob: b)
        self._entry = entry

    @property
    def sheet_names(self) -> list:
        return list(self._entry.order)

    def total_rows(self, sheet_name: str) -> int:
        sh = self._entry.sheets.get(sheet_name)
        return int(sh["total_rows"]) if sh else 0

    def page(self, sheet_name=None, offset: int = 0, limit=None,
             row_indices=None) -> list:
        return _page(self._entry, self._blob, sheet_name, offset, limit, row_indices)


def open_grid(blob: bytes, with_marks: bool = False,
             header_row_by_sheet: dict | None = None) -> GridHandle:
    """Resolve a workbook once, then serve any number of windows from it."""
    return GridHandle(blob, _resolve(blob, with_marks, header_row_by_sheet))


def open_grid_for(content_key: str, load_blob, with_marks: bool = False,
                  header_row_by_sheet: dict | None = None) -> GridHandle:
    """Same as `open_grid`, but for a caller that can name the content without
    holding it — `content_key` must be `fingerprint()` of the bytes `load_blob`
    would return, obtained some cheaper way (see `fingerprint`).

    `load_blob` is called only if this workbook isn't cached, or if it is cached
    without its row values and a window has to be read from the file. On the
    common path — a warm entry — the bytes are never fetched at all.
    """
    key = f"{content_key}:{int(bool(with_marks))}{_header_row_key(header_row_by_sheet)}"
    with _lock:
        entry = _get_locked(key)
        _stats["hits" if entry is not None else "misses"] += 1

    lazy = _LazyBlob(load_blob)
    if entry is None:
        entry = _build(lazy(), with_marks, header_row_by_sheet)
        with _lock:
            _put_locked(key, entry)
    return GridHandle(lazy, entry)


def grid_page(blob: bytes, sheet_name=None, offset: int = 0, limit=None,
              row_indices=None, with_marks: bool = False,
              header_row_by_sheet: dict | None = None) -> list:
    """Cached, chunk-friendly equivalent of a windowed workbook read.

    Returns the same per-sheet shape the endpoint already emits: every sheet
    reports `header`/`total_rows`; the ONE target sheet additionally carries
    whole-sheet `marks`/`notes` plus the requested row window and its
    `row_gis`. The window is either the contiguous range
    `(offset, offset + limit]` or exactly the 1-based rows in `row_indices`.

    One-shot convenience over `open_grid`; use the handle for repeated windows.
    """
    return _page(_resolve(blob, with_marks, header_row_by_sheet), lambda: blob,
                 sheet_name, offset, limit, row_indices)


def _page(entry: _Entry, get_blob, sheet_name, offset: int, limit,
          row_indices) -> list:
    names = entry.order
    target = sheet_name if sheet_name in entry.sheets else (names[0] if names else None)

    if row_indices is not None:
        wanted = {int(i) for i in row_indices}
    else:
        lim = limit or 0
        wanted = set(range(offset + 1, offset + lim + 1)) if lim > 0 else set()

    out: list = []
    for name in names:
        sh = entry.sheets[name]
        if name != target:
            out.append({"sheet": name, "rows": [sh["header"]], "marks": [],
                        "warn_marks": [], "notes": [],
                        "total_rows": sh["total_rows"]})
            continue

        # `gi` is the DATA row number (1 = first row after the header) — always
        # has been, for every existing (header_row=0) caller. `header_row` only
        # shifts where that data actually sits in the file, so translate to the
        # absolute file row when touching `sh["rows"]` (keyed by absolute row).
        header_row = sh.get("header_row", 0)
        wanted_here = {gi for gi in wanted if 1 <= gi <= sh["total_rows"]}
        if sh["rows"] is not None:
            picked = {gi: sh["rows"][header_row + gi]
                      for gi in wanted_here if (header_row + gi) in sh["rows"]}
        else:
            # Only this branch — a workbook too large to keep row values for —
            # needs the file itself, so `get_blob` is the one thing that can
            # force a lazily-held blob to actually be fetched.
            abs_wanted = {header_row + gi for gi in wanted_here}
            streamed = _stream_rows(get_blob(), name, abs_wanted)
            picked = {gi: streamed[header_row + gi]
                      for gi in wanted_here if (header_row + gi) in streamed}

        order_gis = sorted(picked.keys())
        out.append({
            "sheet": name,
            "rows": [sh["header"]] + [picked[gi] for gi in order_gis],
            "row_gis": order_gis,
            "marks": sh["marks"],
            "warn_marks": sh["warn_marks"],
            "notes": sh["notes"],
            "total_rows": sh["total_rows"],
        })
    return out
