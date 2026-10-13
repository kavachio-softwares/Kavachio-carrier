"""export_results.add_results_sheet: the Results sheet looks like a report
(navy headers, bordered tables, % columns, totals) and still never breaks the
download. No database: everything here is built in memory."""
import io
from datetime import datetime

import pytest
from openpyxl import Workbook, load_workbook

import export_results as xr

NAVY = "1F3A5F"
CHECKS = [("Broker must be Wani org", 12), ("Carrier must be Mahi Corp pvt ltd", 12),
          ("Policy expiration date must be on or before 2027-06-30", 4),
          ("Sum insured per risk maximum", 1), ("Total premium written maximum", 1)]


def _book() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Policy Number", "Gross Premium"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _res(extra=()):
    """12 rows, none passed, 30 exceptions over 5 critical checks."""
    sheets = {"Data": [{"Policy Number": f"P-{i}", "Gross Premium": str(100 * i)}
                       for i in range(1, 13)]}
    excs = []
    for rule, n in CHECKS:
        excs += [{"sheet": "Data", "row": i, "rule_id": 1, "rule_name": rule,
                  "field": "Gross Premium", "severity": "critical", "message": "bad"}
                 for i in range(1, n + 1)]
    return xr.compute(sheets, excs + list(extra))


def _sheet(res=None, about=None):
    out = xr.add_results_sheet(_book(), _res() if res is None else res, about)
    return load_workbook(io.BytesIO(out))[xr.RESULTS_SHEET]


def _row_of(ws, text):
    return next(c.row for c in ws["A"] if c.value == text)


def _rgb(color):
    return (getattr(color, "rgb", None) or "")[-6:].upper()


def test_sheet_and_title():
    out = load_workbook(io.BytesIO(xr.add_results_sheet(_book(), _res())))
    assert out.sheetnames == ["Data", "Results"]
    ws = out["Results"]
    assert ws["A1"].value == xr.RESULTS_TITLE
    assert ws["A1"].font.b and ws["A1"].font.sz == 16 and _rgb(ws["A1"].font.color) == NAVY
    assert ws.sheet_view.showGridLines is False


@pytest.mark.parametrize("first", ["Rows", "Exceptions", "Check"])
def test_header_rows_are_navy_with_white_bold_text(first):
    ws = _sheet()
    r = _row_of(ws, first)
    cells = [c for c in ws[r] if c.value is not None]
    assert len(cells) == (5 if first == "Check" else 3)
    for c in cells:
        assert c.fill.fill_type == "solid" and _rgb(c.fill.fgColor) == NAVY
        assert c.font.b and _rgb(c.font.color) == "FFFFFF"
        assert c.border.bottom.style == "medium"


def test_every_table_cell_is_bordered():
    ws = _sheet()
    for first, width in (("Rows", 3), ("Exceptions", 3), ("Check", 5)):
        r = _row_of(ws, first) + 1
        while ws.cell(r, 1).value not in (None, ""):
            for j in range(1, width + 1):
                b = ws.cell(r, j).border
                assert all(s.style for s in (b.left, b.right, b.top, b.bottom)), (first, r, j)
            r += 1


def test_percent_columns_are_values_not_formulas():
    ws = _sheet()
    r = _row_of(ws, "Rows")
    assert [c.value for c in ws[r][:3]] == ["Rows", "Count", "% of rows"]
    got = {ws.cell(i, 1).value: (ws.cell(i, 2).value, ws.cell(i, 3).value)
           for i in range(r + 1, r + 6)}
    assert got == {"Passed": (0, 0.0), "Failed — open": (12, 1.0), "Resolved": (0, 0.0),
                   "Not checked": (0, 0.0), "Total rows": (12, 1.0)}
    assert ws.cell(r + 2, 2).number_format == "#,##0"
    assert ws.cell(r + 2, 3).number_format == "0.0%"
    e = _row_of(ws, "Exceptions")
    assert ws.cell(e, 3).value == "% of exceptions"
    assert [ws.cell(e + 1, j).value for j in (1, 2, 3)] == ["Open", 30, 1.0]
    for row in ws.iter_rows():
        for c in row:
            assert c.data_type != "f" and not str(c.value or "").startswith("=")


def test_percent_is_blank_when_there_is_nothing_to_divide_by():
    ws = _sheet(xr.compute({"Data": []}, []))
    r = _row_of(ws, "Rows")
    assert [ws.cell(i, 3).value for i in range(r + 1, r + 6)] == [None] * 5
    assert ws.cell(_row_of(ws, "Exceptions") + 1, 3).value is None
    assert _rgb(ws.cell(r, 1).fill.fgColor) == NAVY      # still the styled sheet


def test_check_table_has_a_bold_total_under_a_double_rule():
    ws = _sheet()
    r = _row_of(ws, "Check")
    names = [ws.cell(i, 1).value for i in range(r + 1, r + 6)]
    assert names == [n for n, _ in CHECKS]
    total = r + 6
    assert [ws.cell(total, j).value for j in range(1, 6)] == ["Total", None, 30, 30, 0]
    for j in range(1, 6):
        assert ws.cell(total, j).font.b
        assert ws.cell(total, j).border.top.style == "double"
        assert ws.cell(total - 1, j).border.bottom.style == "double"   # same edge, both cells


def test_status_and_severity_colours_are_soft_and_never_flag_colours():
    extra = [{"sheet": "Data", "row": 1, "rule_id": 2, "rule_name": "Expiry present",
              "field": "Gross Premium", "severity": "warning", "message": "m"},
             {"sheet": "Data", "row": 2, "rule_id": 3, "rule_name": "Note",
              "field": "Gross Premium", "severity": "info", "message": "m",
              "status": "approved"}]
    ws = _sheet(_res(extra))
    tint = {"Passed": ("E6F4EA", "1E7B34"), "Failed — open": ("FDECEA", "B3261E"),
            "Resolved": ("E8F0FE", "1A56DB"), "Not checked": ("F1F3F4", "5F6368")}
    r = _row_of(ws, "Rows")
    for i in range(r + 1, r + 5):
        c = ws.cell(i, 1)
        assert (_rgb(c.fill.fgColor), _rgb(c.font.color)) == tint[c.value]
    sev = {ws.cell(i, 2).value: ws.cell(i, 2).font for i in range(_row_of(ws, "Check") + 1,
                                                                  ws.max_row)}
    assert _rgb(sev["Critical"].color) == "B3261E" and sev["Critical"].b
    assert _rgb(sev["Warning"].color) == "B26A00" and sev["Warning"].b
    assert _rgb(sev["Info"].color) == "5F6368" and not sev["Info"].b
    # The in-site viewers read these two fills as exception marks.
    for row in ws.iter_rows():
        for c in row:
            if c.fill.fill_type == "solid":
                assert _rgb(c.fill.fgColor) not in ("FFC7CE", "FFE0B2")


def test_about_line_under_the_title():
    ws = _sheet(about={"File": "bdx.xlsx", "version": 2, "checked_at": datetime(2026, 10, 12, 14, 5),
                       "empty": "", "nothing": None})
    line = ws["A2"].value
    assert "File: bdx.xlsx" in line and "Version: 2" in line
    assert "Checked at: 12 Oct 2026, 14:05" in line and "Empty" not in line
    assert _rgb(ws["A2"].font.color) == "5F6368"
    # No `about` (today's call) or an odd one: no line, never an error.
    for odd in (None, {}, ["not", "a", "dict"], {"x": object(), ("t",): [1, 2]}):
        ws = _sheet(about=odd)
        assert ws["A1"].value == xr.RESULTS_TITLE
        assert _rgb(ws.cell(_row_of(ws, "Rows"), 1).fill.fgColor) == NAVY


def test_long_check_names_wrap_and_the_page_fits_portrait():
    long = "Gross written premium net of brokerage and taxes must reconcile to the signed line share"
    extra = [{"sheet": "Data", "row": 3, "rule_id": 9, "rule_name": long,
              "field": "Gross Premium", "severity": "critical", "message": "m"}]
    ws = _sheet(_res(extra))
    c = next(c for c in ws["A"] if c.value == long)
    assert c.alignment.wrap_text and ws.row_dimensions[c.row].height > 20
    assert ws.column_dimensions["A"].width <= 50
    assert ws.page_setup.orientation == "portrait"
    assert ws.sheet_properties.pageSetUpPr.fitToPage
    assert ws.page_setup.fitToWidth == 1 and ws.page_setup.fitToHeight == 0
    assert not ws.print_options.gridLines


def test_the_styled_sheet_is_still_skipped_as_data():
    out = xr.add_results_sheet(_book(), _res())
    wb = load_workbook(io.BytesIO(out), read_only=True, data_only=True)
    grid = [{"sheet": ws.title, "rows": [["" if v is None else v for v in r]
                                         for r in ws.iter_rows(values_only=True)]}
            for ws in wb.worksheets]
    sheets, _cols = xr.grid_to_sheets(grid)
    assert list(sheets) == ["Data"]


def test_a_styling_failure_falls_back_to_the_plain_sheet(monkeypatch):
    def boom(*_a, **_k):
        raise RuntimeError("style failed")
    monkeypatch.setattr(xr, "_results_sheet_styled", boom)
    ws = load_workbook(io.BytesIO(xr.add_results_sheet(_book(), _res())))["Results"]
    assert ws["A1"].value == xr.RESULTS_TITLE
    values = [c for row in ws.iter_rows(values_only=True) for c in row if c is not None]
    assert "Failed — open" in values and "Broker must be Wani org" in values
    assert all(c.fill.fill_type is None for row in ws.iter_rows() for c in row)


def test_counts_the_styling_cannot_add_up_get_the_plain_sheet():
    res = _res()
    res["by_rule"] = [{"rule": "x", "severity": "critical", "exceptions": "3",
                       "open": "3", "resolved": "0"}]          # text, not numbers
    ws = load_workbook(io.BytesIO(xr.add_results_sheet(_book(), res)))["Results"]
    assert ws["A1"].value == xr.RESULTS_TITLE
    assert ["x", "Critical", "3", "3", "0"] in [list(r) for r in ws.iter_rows(values_only=True)]
    assert all(c.fill.fill_type is None for row in ws.iter_rows() for c in row)


@pytest.mark.parametrize("bad", [
    {"summary": {}},
    None,
    {"summary": {"total": 1}},
    {"summary": _res()["summary"], "by_rule": [{"rule": "x"}]},
])
def test_malformed_results_never_cost_the_file(bad):
    book = _book()
    out = xr.add_results_sheet(book, bad)
    assert out == book                       # both layouts refused it: input back
    assert load_workbook(io.BytesIO(out)).sheetnames == ["Data"]
    assert xr.add_results_sheet(b"not a workbook", _res()) == b"not a workbook"
