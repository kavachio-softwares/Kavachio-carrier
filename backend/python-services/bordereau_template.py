"""The blank bordereau a setup reads, and the blank file a template produces.

WHY TWO FILES, AND WHY THE FIRST IS THE ONE TO FILL IN. A Bordereau Setup learns
to read ONE input layout — the sample it was built from — and to write into ONE
output template. Process Bordereau finds each column of an uploaded file by the
name it learned from that sample, on the sheet it learned it from. So the file a
person fills in has to be the INPUT layout. The output template (a Lloyd's
standard, say) is what Kavachio WRITES; filled in and uploaded instead, its
columns are found only where they happen to share a name with the sample's, and
a sheet named differently is not read at all.

Nothing is assumed about which layout that is. A carrier that wants brokers to
report straight in Lloyd's columns builds the setup with that layout as its
Input Template, and then the input template downloaded here IS that layout.

WHICH SAMPLE. Every run also leaves a landing record against the setup's input
format, and a run file can be laid out differently from the sample. The sample
is the record whose fingerprint is the format's own — the fingerprint was taken
from it when the setup was built — so a drifted run file can never become the
template. A setup built before that held falls back to the newest landing, the
same rule the setup editor uses to show its input columns.
"""
from __future__ import annotations

import io
import json
import re
from typing import Any, Optional
from urllib.parse import quote

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

_HEAD_FONT = Font(bold=True)
_HEAD_FILL = PatternFill("solid", fgColor="E8EEF7")


class TemplateUnavailable(LookupError):
    """There is nothing to build a template from. The message is for the user."""


# ── which layout ────────────────────────────────────────────────────────────
def sample_landing(s, fmt):
    """The landing record the setup was learned from (see module docstring)."""
    from db import LandingRecord
    q = s.query(LandingRecord).filter(LandingRecord.format_id == fmt.id)
    if fmt.fingerprint:
        hit = (q.filter(LandingRecord.fingerprint == fmt.fingerprint)
               .order_by(LandingRecord.id.desc()).first())
        if hit is not None:
            return hit
    return q.order_by(LandingRecord.id.desc()).first()


def routed_input_sheets(routing: Optional[dict]) -> set[str]:
    """Input sheets a run actually reads — the same names direct_run keeps."""
    return {str(src["input_sheet"])
            for r in (routing or {}).get("routes") or []
            for src in r.get("sources") or [] if src.get("input_sheet")}


def input_layout(sample_sheets: Optional[dict],
                 routing: Optional[dict]) -> list[tuple[str, list[str]]]:
    """[(sheet, [column, ...]), ...] in the sample's own order.

    Only the sheets the routing reads: a run ignores every other tab, so a blank
    copy of them would be a form nobody's answers are taken from. A routing that
    names none of the sample's sheets keeps them all, as a run does."""
    layout = [(str(name), [str(c) for c in (sheet or {}).get("columns") or []])
              for name, sheet in (sample_sheets or {}).items()]
    layout = [(name, cols) for name, cols in layout if cols]
    wanted = routed_input_sheets(routing)
    return [x for x in layout if x[0] in wanted] or layout


# ── the files ───────────────────────────────────────────────────────────────
def layout_workbook(layout: list[tuple[str, list[str]]]) -> bytes:
    """Headings on row 1 of each sheet and nothing under them."""
    wb = Workbook()
    wb.remove(wb.active)
    for name, cols in layout:
        ws = wb.create_sheet(title=name[:31])
        for i, col in enumerate(cols, start=1):
            cell = ws.cell(row=1, column=i, value=col)
            cell.font, cell.fill = _HEAD_FONT, _HEAD_FILL
            ws.column_dimensions[get_column_letter(i)].width = min(max(len(col) + 2, 12), 48)
        ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def setup_input_template(s, pipe) -> tuple[bytes, str]:
    """The blank bordereau for a setup, and the name to save it under."""
    from db import DirectFormat
    fmt = s.get(DirectFormat, pipe.input_format_id) if pipe.input_format_id else None
    if fmt is None:
        raise TemplateUnavailable(
            "This setup has no Input Template yet, so there is no layout to download.")
    rec = sample_landing(s, fmt)
    layout = input_layout((rec.data or {}).get("sheets") if rec else None,
                          fmt.sheet_routing)
    if not layout:
        raise TemplateUnavailable(
            "The sample this setup was built from is no longer on file, so its "
            "layout cannot be rebuilt. Re-upload the Input Template on the setup.")
    return layout_workbook(layout), download_name(pipe.name or fmt.name,
                                                  "Bordereau Template")


def output_workbook(structure: Any, template_bytes: Optional[bytes]) -> bytes:
    """The file a template produces, with no rows in it — through the SAME
    writer a run uses, so the headings, their order and the sample's styling
    are exactly what Generate BDX would deliver."""
    import direct_render as dr
    if isinstance(structure, str):
        structure = json.loads(structure)
    return dr.render_output(structure or {"sheets": []}, {}, template_bytes)


# ── naming ──────────────────────────────────────────────────────────────────
def download_name(base: Optional[str], what: str) -> str:
    base = re.sub(r'[\\/:*?"<>|]+', " ", base or "").strip() or "Bordereau"
    return f"{base} - {what}.xlsx"


def attachment(filename: str) -> dict[str, str]:
    """Content-Disposition that survives a non-ASCII name (setup names carry an
    em dash). `filename*` comes FIRST: the frontend's downloadFile takes the
    first filename it finds, and the ASCII fallback would otherwise win."""
    fallback = re.sub(r"[^\x20-\x7e]", "_", filename).replace('"', "'")
    return {"Content-Disposition":
            f"attachment; filename*=UTF-8''{quote(filename)}; filename=\"{fallback}\""}
