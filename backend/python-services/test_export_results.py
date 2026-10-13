"""export_results: every row's outcome and the counts (Results tab, Excel sheet, CSV)."""
import io

from openpyxl import Workbook, load_workbook

import export_results as xr


def _exc(row, rule="Premium check", field="Gross Premium", status=None, sheet="Data", **kw):
    e = {"sheet": sheet, "row": row, "rule_id": 1, "rule_name": rule, "field": field,
         "severity": "warning", "message": "bad"}
    if status:
        e["status"] = status
    e.update(kw)
    return e


SHEETS = {"Data": [
    {"Policy Number": "P-1", "Gross Premium": "100"},
    {"Policy Number": "P-2", "Gross Premium": "200"},
    {"Policy Number": "P-3", "Gross Premium": "300"},
    {"Policy Number": "P-4", "Gross Premium": "400"},
]}


def test_each_row_has_one_outcome():
    res = xr.compute(SHEETS, [
        _exc(2),                                   # open
        _exc(3, status="dismissed"),               # resolved
        _exc(4, status="fixed"), _exc(4, rule="Other"),   # one still open → open
    ])
    out = {r["row"]: r["outcome"] for r in res["rows"]}
    assert out == {1: "passed", 2: "open", 3: "resolved", 4: "open"}
    s = res["summary"]
    assert (s["total"], s["passed"], s["open"], s["resolved"], s["not_checked"]) == (4, 1, 2, 1, 0)
    assert (s["exceptions"], s["exceptions_open"], s["exceptions_resolved"]) == (4, 2, 2)
    assert res["rows"][0]["policy_number"] == "P-1"      # read from the row itself


def test_notices_are_not_counted():
    notice = {"error_class": "not_checked", "rule_id": None, "message": "2 rules not run"}
    res = xr.compute(SHEETS, [notice])
    assert res["summary"]["passed"] == 4 and res["summary"]["exceptions"] == 0


def test_not_validated_file_is_not_checked():
    res = xr.compute(SHEETS, [], not_validated=True)
    assert res["summary"]["not_checked"] == 4 and res["summary"]["passed"] == 0


def test_blank_row_is_not_checked():
    sheets = {"Data": SHEETS["Data"] + [{"Policy Number": "", "Gross Premium": ""}]}
    res = xr.compute(sheets, [])
    assert res["rows"][-1]["outcome"] == "not_checked"
    assert res["summary"]["passed"] == 4


def test_decisions_resolve_exceptions():
    excs = xr.apply_decisions([_exc(2)], [{"output_sheet": "Data", "output_row": 2,
                                           "output_field": "Gross Premium", "kind": "approve"}])
    assert excs[0]["status"] == "approved"
    assert xr.compute(SHEETS, excs)["rows"][1]["outcome"] == "resolved"


def test_results_sheet_and_csv():
    wb = Workbook(); ws = wb.active; ws.title = "Data"; ws.append(["Policy Number", "Gross Premium"])
    buf = io.BytesIO(); wb.save(buf)
    res = xr.compute(SHEETS, [_exc(2)])
    out = load_workbook(io.BytesIO(xr.add_results_sheet(buf.getvalue(), res)))
    assert out.sheetnames == ["Data", "Results"]
    cells = [c for row in out["Results"].iter_rows(values_only=True) for c in row if c is not None]
    assert "Passed" in cells and "Failed — open" in cells and "Not checked" in cells
    # never the red/orange fills the viewers read as exception marks
    for row in out["Results"].iter_rows():
        for c in row:
            assert (c.fill.fgColor.rgb or "") not in ("00FFC7CE", "FFFFC7CE", "00FFE0B2", "FFFFE0B2")
    csv_text = xr.results_csv(SHEETS, res)
    lines = csv_text.splitlines()
    assert lines[0] == "Row,Policy Number,Gross Premium,Result,Exceptions"
    assert lines[2] == '2,P-2,200,"Failed — 1 exception, 1 open",Premium check (Gross Premium): open'


def test_results_sheet_never_breaks_the_file():
    assert xr.add_results_sheet(b"not a workbook", {"summary": {}}) == b"not a workbook"


def test_grid_skips_the_results_sheet():
    grid = [{"sheet": "Data", "rows": [["Policy Number"], ["P-1"]]},
            {"sheet": "Results", "rows": [[xr.RESULTS_TITLE], ["Rows", "Count"]]}]
    sheets, cols = xr.grid_to_sheets(grid)
    assert list(sheets) == ["Data"] and sheets["Data"] == [{"Policy Number": "P-1"}]
