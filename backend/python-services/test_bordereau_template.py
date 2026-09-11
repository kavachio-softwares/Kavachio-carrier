"""The blank bordereau, downloadable before Process Bordereau.

THE TRAP THIS GUARDS. A setup reads ONE input layout and writes ONE output
template, and a run finds each column of an uploaded file by the name it learned
from the INPUT sample, on the sheet it learned it from. So the file offered for
filling in has to be that input layout, exactly — a heading spelt differently or
a sheet renamed, and the column is silently not read. The proof used here is the
setup's own: the downloaded file, read back the way a run reads it, carries the
fingerprint the setup was built with.

Also guarded: a run's drifted file never becomes the template; a broker gets it
through their own contract path and nobody else's tenant can; the output
template download is the file a run writes, with no rows.

    python -m pytest test_bordereau_template.py
"""
from __future__ import annotations

import io
import os

import pandas as pd
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

import main  # noqa: F401 — loads .env, which db needs at import
import bordereau_template as bt
from auth_tokens import mint_access_token
from db import (
    AppUser, Contract, DirectFormat, ExportTemplate, LandingRecord, Party,
    Pipeline, Program, ProgramBroker, SessionLocal, Tenant,
)
from fingerprint import signature_hash
from mapper import read_excel_all_sheets, signature_multi

client = TestClient(main.app)

RISK = ["Unique Market Reference (UMR)", "Certificate Ref",
        "Location of Risk – Country", "Gross Written Premium", "Commission %"]


def _fp(sheets: dict[str, list[str]]) -> str:
    """The fingerprint a setup takes of its sample — same functions, same order."""
    return signature_hash(signature_multi(
        {name: pd.DataFrame([["x"] * len(cols)], columns=cols)
         for name, cols in sheets.items()}))


# ══════════════════════════════════════════════════════════════════════════
#  which layout, and the file — no database
# ══════════════════════════════════════════════════════════════════════════
def test_only_the_sheets_a_run_reads_are_offered():
    sample = {"Contract terms used": {"columns": ["Term", "Agreed value"]},
              "Risk BDX": {"columns": RISK},
              "What each row should do": {"columns": ["Row", "Severity"]}}
    routing = {"routes": [{"output_sheet": "Lloyds", "sources": [{"input_sheet": "Risk BDX"}]}]}
    assert bt.input_layout(sample, routing) == [("Risk BDX", RISK)]


def test_a_routing_naming_none_of_the_sheets_keeps_them_all():
    """As a run does — it only narrows when the routing's sheets are present."""
    sample = {"A": {"columns": ["one"]}, "B": {"columns": ["two"]}, "Empty": {"columns": []}}
    routing = {"routes": [{"output_sheet": "X", "sources": [{"input_sheet": "Gone"}]}]}
    assert bt.input_layout(sample, routing) == [("A", ["one"]), ("B", ["two"])]
    assert bt.input_layout(sample, None) == [("A", ["one"]), ("B", ["two"])]


def test_the_blank_file_reads_back_as_the_setup_learned_it():
    """The whole point: headings survive character for character (the en dash
    included), nothing is under them, and the fingerprint is the setup's."""
    data = bt.layout_workbook([("Risk BDX", RISK)])
    sheets = read_excel_all_sheets(data, 0)
    assert list(sheets) == ["Risk BDX"]
    assert [str(c) for c in sheets["Risk BDX"].columns] == RISK
    assert len(sheets["Risk BDX"]) == 0
    assert signature_hash(signature_multi(sheets)) == _fp({"Risk BDX": RISK})


def test_the_output_file_is_what_a_run_writes_with_no_rows():
    """Active columns only, in the user's order, under the user's names."""
    structure = {"sheets": [{"sheet_name": "Lloyds", "header_row": 0, "data_start_row": 1,
                             "columns": [
        {"column_index": 0, "column_name": "UMR", "display_order": 1},
        {"column_index": 1, "column_name": "Old", "active": False},
        {"column_index": 2, "column_name": "Premium", "display_name": "Gross Premium",
         "display_order": 0}]}]}
    ws = load_workbook(io.BytesIO(bt.output_workbook(structure, None))).active
    assert [c.value for c in ws[1]] == ["Gross Premium", "UMR"]
    assert ws.max_row == 1


def test_a_setup_name_with_an_em_dash_downloads_under_its_own_name():
    name = bt.download_name("GIRI — DEMO A1", "Bordereau Template")
    assert name == "GIRI — DEMO A1 - Bordereau Template.xlsx"
    header = bt.attachment(name)["Content-Disposition"]
    header.encode("latin-1")                       # a header that cannot encode is a 500
    assert header.index("filename*=") < header.index('filename="')
    assert bt.download_name('a/b:c*?"<>|', "X") == "a b c - X.xlsx"


# ══════════════════════════════════════════════════════════════════════════
#  the routes — a carrier, its broker, a setup, and a stranger
# ══════════════════════════════════════════════════════════════════════════
@pytest.fixture(scope="module")
def world():
    from ingester import _ensure_carrier_party
    sfx = os.urandom(4).hex()
    made: list[tuple[type, int]] = []    # (model, id) — read while still attached
    with SessionLocal() as s:
        founder = (s.query(AppUser).filter(AppUser.role == "kavachio_admin")
                   .order_by(AppUser.id).first())
        if founder is None:
            pytest.skip("no kavachio_admin on this database to seed an invite chain")

        def add(row):
            s.add(row); s.commit(); made.append((type(row), row.id))
            return row

        carrier = add(Tenant(tenant_name=f"bt-{sfx}", legal_name="Template Carrier"))
        stranger_t = add(Tenant(tenant_name=f"bt-x-{sfx}", legal_name="Someone Else"))
        cpid = _ensure_carrier_party(s, carrier.id); s.commit()
        made.append((Party, cpid))                 # created for the carrier; ours to remove
        broker = add(Party(tenant_id=carrier.id, party_type="broker",
                           legal_name="Template Broker", reference=f"bt-b-{sfx}"))
        prog = add(Program(tenant_id=carrier.id, name=f"BT {sfx}", is_app_managed=True))
        add(ProgramBroker(tenant_id=carrier.id, program_id=prog.id,
                          broker_party_id=broker.id, status="active"))
        contract = add(Contract(tenant_id=carrier.id, program_id=prog.id,
                                broker_party_id=broker.id, name=f"BT contract {sfx}",
                                lifecycle="active"))
        admin = add(AppUser(tenant_id=carrier.id, email=f"bt-a-{sfx}@carrier.test",
                            full_name="Carrier Admin", role="carrier_admin",
                            invited_by_user_id=founder.id))
        seat = add(AppUser(tenant_id=None, broker_party_id=broker.id,
                           email=f"bt-b-{sfx}@broker.test", full_name="Broker Seat",
                           role="broker_admin", invited_by_user_id=admin.id))
        other = add(AppUser(tenant_id=stranger_t.id, email=f"bt-x-{sfx}@else.test",
                            full_name="Stranger", role="carrier_admin",
                            invited_by_user_id=founder.id))

        sample = {"Risk BDX": RISK, "Contract terms used": ["Term", "Agreed value"]}
        fmt = add(DirectFormat(
            tenant_id=carrier.id, name=f"BT setup {sfx}", fingerprint=_fp(sample),
            carrier_party_id=cpid, program_id=prog.id, approved=1,
            sheet_routing={"routes": [{"output_sheet": "Lloyds",
                                       "sources": [{"input_sheet": "Risk BDX"}]}]}))
        add(LandingRecord(tenant_id=carrier.id, format_id=fmt.id, fingerprint=fmt.fingerprint,
                          source_filename="sample.xlsx", row_count=1,
                          data={"sheets": {n: {"columns": c, "rows": []}
                                           for n, c in sample.items()}, "row_count": 1}))
        # A LATER run whose file drifted. Newest, same format — and never the template.
        add(LandingRecord(tenant_id=carrier.id, format_id=fmt.id, fingerprint="drifted",
                          source_filename="drifted.xlsx", row_count=1,
                          data={"sheets": {"Risk BDX": {"columns": ["Something else"],
                                                        "rows": []}}, "row_count": 1}))
        tpl = add(ExportTemplate(tenant_id=carrier.id, name=f"BT Lloyds {sfx}", version=1,
                                 is_active=1, program_id=prog.id,
                                 structure={"sheets": [{"sheet_name": "Lloyds",
                                     "header_row": 0, "data_start_row": 1, "columns": [
                                         {"column_index": 0, "column_name": "UMR"},
                                         {"column_index": 1, "column_name": "Gross Premium"}]}]}))
        pipe = add(Pipeline(tenant_id=carrier.id, name=f"BT — setup {sfx}",
                            carrier_party_id=cpid, program_id=prog.id,
                            input_format_id=fmt.id, output_template_id=tpl.id,
                            status="active"))
        w = {
            "pipe": pipe.id, "tpl": tpl.id, "pipe_name": pipe.name,
            "chain": (f"/carriers/{carrier.id}/programs/{prog.id}"
                      f"/brokers/{broker.id}/contracts/{contract.id}"),
            "carrier": {"Authorization": f"Bearer {mint_access_token(admin.id, carrier.id, 'carrier_admin')}"},
            "broker": {"Authorization": f"Bearer {mint_access_token(seat.id, None, 'broker_admin')}"},
            "stranger": {"Authorization": f"Bearer {mint_access_token(other.id, stranger_t.id, 'carrier_admin')}"},
        }
    yield w
    with SessionLocal() as s:                      # by id only — never a cascade
        for model, pk in reversed(made):
            try:
                obj = s.get(model, pk)
                if obj is not None:
                    s.delete(obj); s.commit()
            except Exception:  # noqa: BLE001 — cleanup must not mask a result
                s.rollback()


def _sheets(resp) -> dict[str, list[str]]:
    assert resp.status_code == 200, resp.text
    assert resp.headers["content-type"].startswith(bt.XLSX)
    return {n: [str(c) for c in df.columns]
            for n, df in read_excel_all_sheets(resp.content, 0).items()}


def test_the_carrier_downloads_the_layout_the_setup_learned(world):
    r = client.get(f"/pipelines/{world['pipe']}/bordereau-template", headers=world["carrier"])
    assert _sheets(r) == {"Risk BDX": RISK}         # not the drifted run's columns
    assert "Bordereau%20Template.xlsx" in r.headers["content-disposition"]


def test_another_carrier_cannot_download_it(world):
    r = client.get(f"/pipelines/{world['pipe']}/bordereau-template", headers=world["stranger"])
    assert r.status_code in (403, 404)


def test_the_broker_downloads_the_same_file_through_their_contract(world):
    r = client.get(f"{world['chain']}/bordereau-template", headers=world["broker"])
    assert _sheets(r) == {"Risk BDX": RISK}


def test_readiness_still_names_the_setup_the_broker_downloads(world):
    """_pipe now delegates to the shared lookup — the answer must not move."""
    r = client.get(f"{world['chain']}/bordereau", headers=world["broker"])
    assert r.status_code == 200, r.text
    assert r.json()["ready"] is True and r.json()["setup"]["id"] == world["pipe"]


def test_a_draft_setup_offers_nothing_to_the_broker(world):
    with SessionLocal() as s:
        s.get(Pipeline, world["pipe"]).status = "draft"; s.commit()
    try:
        r = client.get(f"{world['chain']}/bordereau-template", headers=world["broker"])
        assert r.status_code == 404
        assert "no live bordereau setup" in r.json()["detail"]
    finally:
        with SessionLocal() as s:
            s.get(Pipeline, world["pipe"]).status = "active"; s.commit()


def test_the_output_template_downloads_as_the_file_a_run_writes(world):
    r = client.get(f"/output-template/{world['tpl']}/download", headers=world["carrier"])
    assert _sheets(r) == {"Lloyds": ["UMR", "Gross Premium"]}
    r = client.get(f"/output-template/{world['tpl']}/download", headers=world["stranger"])
    assert r.status_code in (403, 404)
