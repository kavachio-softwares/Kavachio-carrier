#!/usr/bin/env python
"""
make_sample_bordereau.py — a bordereau built to break a contract's own checks.

WHAT IT IS FOR. A contract's terms become checks; the only way to know the
checks are right is to put a file in front of them and see which rows they
catch. This writes that file: one clean row, then one row per rule that breaks
exactly that rule and nothing else, then a row that breaks several at once — so
every check has to fire once and stay silent everywhere else.

NOTHING IN IT IS TYPED BY HAND. Every value is derived from the contract's own
rules: the passing value comes from what the term allows, and the breaking value
is manufactured from the same operand (one over a cap, a currency the contract
does not name, a country on its exclusion list). Change a term, run this again,
and the sample follows — which is the same principle the wording and the checks
already work on, and the reason a sample workbook checked into a repo goes stale
and this does not.

    python scripts/make_sample_bordereau.py <contract_id> [-o out.xlsx]

The companion test is test_contract_sample_bordereau.py, which builds the same
rows and runs them through the real validation engine.
"""
from __future__ import annotations

import argparse
import os
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from dotenv import load_dotenv

load_dotenv(Path(__file__).resolve().parent.parent / ".env", override=False)

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from sample_bordereau import build_rows, contract_rules_for, sheet_of


HEAD_FILL = PatternFill("solid", fgColor="1F3864")
BREAK_FILL = PatternFill("solid", fgColor="FDE7E9")
CLEAN_FILL = PatternFill("solid", fgColor="E9F7EF")


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("contract_id", type=int)
    ap.add_argument("-o", "--out", default=None)
    args = ap.parse_args()

    rules = contract_rules_for(args.contract_id)
    if not rules:
        print(f"contract {args.contract_id} has no checks bound — press "
              f"“Bind checks” on its page first, or there is nothing to test.")
        return 1

    sheet = sheet_of(rules)
    columns, rows = build_rows(rules)

    wb = Workbook()
    ws = wb.active
    ws.title = sheet[:31]           # Excel's own limit, not ours
    ws.append(columns)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    for r in rows:
        ws.append([r["values"].get(c, "") for c in columns])
        fill = CLEAN_FILL if not r["breaks"] else BREAK_FILL
        for cell in ws[ws.max_row]:
            cell.fill = fill
    for i, name in enumerate(columns, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = \
            min(34, max(12, len(name) // 2 + 8))

    # What each row is FOR, beside the file rather than in an email about it.
    key = wb.create_sheet("What each row should do")
    key.append(["Row", "What it is", "Which check should catch it",
                "Severity", "The term it comes from"])
    for cell in key[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEAD_FILL
    for n, r in enumerate(rows, start=2):      # row 1 is the header
        if not r["breaks"]:
            key.append([n, "Clean — every term satisfied",
                        "nothing", "", ""])
            continue
        key.append([
            n, r["note"],
            "; ".join(b["rule_name"] for b in r["breaks"]),
            "; ".join(b["severity"] for b in r["breaks"]),
            "; ".join(f"{b['column']} {b['operator']} {b['operand']}"
                      for b in r["breaks"]),
        ])
    for col, width in zip("ABCDE", (6, 46, 34, 18, 52)):
        key.column_dimensions[col].width = width

    out = args.out or f"sample-bordereau-contract-{args.contract_id}.xlsx"
    wb.save(out)
    print(f"wrote {out}")
    print(f"  sheet     {sheet}")
    print(f"  columns   {len(columns)}")
    print(f"  rows      {len(rows)} "
          f"({sum(1 for r in rows if not r['breaks'])} clean)")
    print(f"  breaches  {sum(len(r['breaks']) for r in rows)} across "
          f"{len(rules)} checks")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
