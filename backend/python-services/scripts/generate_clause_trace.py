"""Contract → clause → rule → SQL → output trace report (standalone).

Runs the existing contract-upload pipeline on a Contract PDF + a BDX workbook and
collects EVERY stage into one Excel workbook (plus a raw JSON dump):

    Clauses          — every clause extracted from the contract (verbatim text,
                       page, type, whether it is rule-bearing)
    Rules            — the validation rule(s) generated from each clause (template,
                       severity, target field, condition, error message)
    Queries (SQL)    — the deterministic DuckDB query compiled from each rule, plus
                       its run status against this BDX
    Column Mapping   — which BDX sheet/column each rule binds to
    SQL Output       — the violating rows produced by running each rule's SQL
    Review & Control — clauses that did NOT become rules (and why)
    Summary          — pipeline counts

The BDX file is used as BOTH the data source and the Output Template
("input sheet = output sheet"), so every rule's field is a real BDX column and the
same workbook can be validated against its own rules.

Usage:
    cd backend/python-services
    source venv/bin/activate            # (or .venv) — needs the project deps
    python -m scripts.generate_clause_trace \
        "../../Demo/Contract - Insurisk Spectrum Transportation - CRC Insurisk - revisions - 2025.pdf" \
        "../../Demo/BDX - Insurisk Spectrum Transportation - Written BDX - 122025.xlsx"

    # Defaults to those two Demo files when no paths are given.

Options:
    --out PATH         Output workbook path  (default: <bdx>_trace.xlsx next to the BDX)
    --json PATH        Output JSON path      (default: same stem as --out, .json)
    --sheets A,B       Only treat these sheet names as data tables (default: auto-detect)
    --min-cols N       A row needs >= N distinct non-empty cells to count as a header (default 3)

Requirements:
    * GEMINI_API_KEY in the environment / .env  (the extraction + synthesis LLM calls)
    * Read access to the rule-template catalog in Postgres (DATABASE_URL).
      This script READS the catalog but never writes to Postgres.
"""
from __future__ import annotations

import argparse
import json
import os
import sys

# Make the backend package importable when run as a module from any cwd.
_BACKEND_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, _BACKEND_DIR)

# Standalone scripts don't go through main.py, which is the only place that loads
# .env — so load it here for GEMINI_API_KEY / DATABASE_URL.
try:
    from dotenv import load_dotenv
    load_dotenv(os.path.join(_BACKEND_DIR, ".env"))
except Exception:
    pass

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

_REPO_ROOT = os.path.dirname(  # backend/python-services -> backend -> repo root
    os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
)
_DEMO = os.path.join(_REPO_ROOT, "Demo")
_DEFAULT_PDF = os.path.join(
    _DEMO, "Contract - Insurisk Spectrum Transportation - CRC Insurisk - revisions - 2025.pdf")
_DEFAULT_BDX = os.path.join(
    _DEMO, "BDX - Insurisk Spectrum Transportation - Written BDX - 122025.xlsx")

_ANALYSIS_SHEETS = [
    "Clauses", "Rules", "Queries (SQL)", "Column Mapping",
    "SQL Output", "Review & Control", "Summary",
]


# =========================================================
# 1) Read the BDX as data + Output Template
# =========================================================

def _s(v):
    """Cell -> trimmed string (or None). DuckDB loads everything as VARCHAR."""
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def _detect_sheet(rows, min_cols):
    """Find the header row in a sheet's rows.

    Returns (header_index, [(col_pos, name), ...]) or (None, None) when the sheet
    has no row that looks like a table header (e.g. a key/value banner sheet).
    The first row within the first 15 whose distinct non-empty cell count reaches
    `min_cols` wins.
    """
    for idx in range(min(15, len(rows))):
        cells = rows[idx]
        names = []
        seen = {}
        for pos, raw in enumerate(cells):
            name = _s(raw)
            if not name:
                continue
            # De-duplicate repeated headers so DuckDB column names stay unique.
            if name in seen:
                seen[name] += 1
                name = f"{name}_{seen[name]}"
            else:
                seen[name] = 1
            names.append((pos, name))
        if len({n for _, n in names}) >= min_cols:
            return idx, names
    return None, None


def read_bdx(path, only_sheets=None, min_cols=3):
    """Return (template_fields, records_by_sheet, schema_cols, skipped).

    template_fields  : [{name, sheet, canonical_field=None, samples:[..3]}]  (Output Template)
    records_by_sheet : [{"sheet": name, "records": [{col: val}, ...]}]        (data for DuckDB)
    schema_cols      : {sheet: [col, ...]}                                    (stable column set)
    skipped          : [sheet names treated as non-tabular]
    """
    wb = load_workbook(path, data_only=True)
    template_fields, records_by_sheet, schema_cols, skipped = [], [], {}, []

    for ws in wb.worksheets:
        if only_sheets is not None and ws.title not in only_sheets:
            skipped.append(ws.title)
            continue

        rows = list(ws.iter_rows(values_only=True))
        header_idx, header = _detect_sheet(rows, min_cols)
        if header_idx is None:
            skipped.append(ws.title)
            continue

        cols = [name for _, name in header]
        schema_cols[ws.title] = cols

        records = []
        for r in rows[header_idx + 1:]:
            rec = {name: _s(r[pos]) if pos < len(r) else None for pos, name in header}
            if any(v is not None for v in rec.values()):
                records.append(rec)
        records_by_sheet.append({"sheet": ws.title, "records": records})

        # Per-column sample values (first 3 distinct) for the Output Template.
        for _, name in header:
            samples, seen = [], set()
            for rec in records:
                v = rec.get(name)
                if v is not None and v not in seen:
                    seen.add(v)
                    samples.append(v)
                if len(samples) >= 3:
                    break
            template_fields.append({
                "name": name, "sheet": ws.title,
                "canonical_field": None, "samples": samples,
            })

    return template_fields, records_by_sheet, schema_cols, skipped


# =========================================================
# 2) Run each rule's compiled SQL against the BDX (DuckDB)
# =========================================================

def run_sql(records_by_sheet, schema_cols, rules):
    """Execute every rule's compiled_sql against the BDX and collect violating rows.

    Returns [{rule_name, template, severity, status, violation_count, error,
              compiled_sql, violations:[{...row...}]}], using the same DuckDB
    connection builder the runtime uses (duckdb_validation.build_connection).
    """
    from duckdb_validation import build_connection

    con, _tables = build_connection(records_by_sheet, schema_cols=schema_cols,
                                    label="clause_trace")
    runs = []
    try:
        for r in rules:
            sql = r.get("compiled_sql") or (r.get("rule_spec") or {}).get("compiled_sql")
            run = {
                "rule_name": r.get("rule_name"),
                "template": r.get("template"),
                "severity": r.get("severity"),
                "source_clause_id": r.get("source_clause_id"),
                "compiled_sql": sql,
                "status": "no_sql",
                "violation_count": 0,
                "error": None,
                "violations": [],
            }
            if sql:
                try:
                    cur = con.execute(sql)
                    headers = [d[0] for d in cur.description]
                    for row in cur.fetchall():
                        run["violations"].append(dict(zip(headers, row)))
                    run["violation_count"] = len(run["violations"])
                    run["status"] = "violations" if run["violation_count"] else "clean"
                except Exception as e:  # noqa: BLE001 — surface, never crash the report
                    run["status"] = "error"
                    run["error"] = str(e)
            runs.append(run)
    finally:
        try:
            con.close()
        except Exception:
            pass
    return runs


# =========================================================
# 3) Shape pipeline output into flat rows for the workbook
# =========================================================

def _ir_of(rule):
    return rule.get("ir") or (rule.get("rule_spec") or {}).get("ir") or {}


def _condition(ir):
    """Compact human-readable condition from an IR's params (max/allowed/...)."""
    p = ir.get("params") or {}
    parts = []
    for key in ("operator", "max", "min", "value", "allowed", "excluded",
                "pattern", "unit", "group_by", "scope"):
        if key in p and p[key] not in (None, [], {}, ""):
            v = p[key]
            parts.append(f"{key}={json.dumps(v, default=str) if isinstance(v, (list, dict)) else v}")
    return "; ".join(parts)


def _target(rule):
    ct = rule.get("canonical_target") or {}
    fields = ct.get("output_fields") or ([ct.get("output_field")] if ct.get("output_field") else [])
    return ", ".join([f for f in fields if f])


def clause_rows(clauses):
    out = []
    for c in clauses:
        clf = c.get("classification") or {}
        out.append({
            "clause_id": c.get("clause_id"),
            "type": c.get("clause_type"),
            "title": c.get("title"),
            "page": c.get("page_number") or c.get("page"),
            "section": c.get("section_header"),
            "rule_bearing": clf.get("is_rule_bearing"),
            "status": c.get("rule_generation_status"),
            "confidence": c.get("extraction_confidence"),
            "clause_text": c.get("text"),
        })
    return out


def rule_rows(rules):
    out = []
    for r in rules:
        ir = _ir_of(r)
        out.append({
            "rule_name": r.get("rule_name"),
            "from_clause": r.get("source_clause_id"),
            "template": r.get("template"),
            "severity": r.get("severity"),
            "target_field": _target(r),
            "condition": _condition(ir),
            "error_message": r.get("error_message"),
            "confidence": r.get("generation_confidence"),
            "status": r.get("rule_status"),
            "description": r.get("rule_description"),
        })
    return out


def mapping_rows(rules, field_to_sheet):
    out = []
    for r in rules:
        ir = _ir_of(r)
        ct = r.get("canonical_target") or {}
        fields = ct.get("output_fields") or ([ct.get("output_field")] if ct.get("output_field") else [])
        sheets = sorted({field_to_sheet.get(f, "") for f in fields if f} - {""})
        out.append({
            "rule_name": r.get("rule_name"),
            "from_clause": r.get("source_clause_id"),
            "sheet": ", ".join(sheets),
            "target_field(s)": _target(r),
            "template": r.get("template"),
            "condition": _condition(ir),
        })
    return out


def query_rows(rules, runs):
    by_name = {(rn.get("rule_name"), rn.get("source_clause_id")): rn for rn in runs}
    out = []
    for r in rules:
        run = by_name.get((r.get("rule_name"), r.get("source_clause_id"))) or {}
        out.append({
            "rule_name": r.get("rule_name"),
            "from_clause": r.get("source_clause_id"),
            "template": r.get("template"),
            "run_status": run.get("status"),
            "violations": run.get("violation_count"),
            "error": run.get("error"),
            "compiled_sql": r.get("compiled_sql") or (r.get("rule_spec") or {}).get("compiled_sql"),
        })
    return out


def output_rows(runs):
    out = []
    for run in runs:
        for v in run.get("violations", []):
            out.append({
                "rule_name": run.get("rule_name"),
                "severity": run.get("severity"),
                "sheet": v.get("sheet"),
                "row": v.get("row_id"),
                "column": v.get("field"),
                "actual_value": v.get("actual_value"),
                "policy_number": v.get("policy_number"),
                "reason": v.get("reason"),
            })
    return out


def routing_rows(review, control):
    out = []
    for bucket, items in (("review", review or []), ("control", control or [])):
        for it in items:
            ir = it.get("ir") or {}
            out.append({
                "bucket": bucket,
                "clause_id": it.get("clause_id"),
                "rule_name": ir.get("rule_name") or it.get("rule_name"),
                "reason": it.get("reason"),
                "clause_text": it.get("clause_text"),
            })
    return out


# =========================================================
# 4) Write the workbook (original BDX sheets + analysis sheets)
# =========================================================

_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_WRAP_COLS = {"clause_text", "compiled_sql", "description", "reason",
              "error_message", "error", "condition"}


def _write_sheet(wb, title, rows):
    ws = wb.create_sheet(title=title[:31])
    if not rows:
        ws["A1"] = "(none)"
        return
    headers = list(rows[0].keys())
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="top")
    for r, row in enumerate(rows, 2):
        for c, h in enumerate(headers, 1):
            v = row.get(h)
            if isinstance(v, (list, dict)):
                v = json.dumps(v, default=str)
            cell = ws.cell(row=r, column=c, value=v)
            if h in _WRAP_COLS:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            else:
                cell.alignment = Alignment(vertical="top")
    # Column widths: wide & wrapped for prose/SQL, snug otherwise.
    for c, h in enumerate(headers, 1):
        letter = get_column_letter(c)
        if h in _WRAP_COLS:
            ws.column_dimensions[letter].width = 70
        else:
            longest = max([len(str(h))] + [len(str(row.get(h) or "")) for row in rows[:200]])
            ws.column_dimensions[letter].width = min(max(longest + 2, 10), 40)
    ws.freeze_panes = "A2"


def write_workbook(bdx_path, out_path, sections, summary):
    # Start from the original BDX so the output workbook literally contains the
    # input sheets ("input sheet = output sheet"), then append analysis sheets.
    wb = load_workbook(bdx_path)
    for title in _ANALYSIS_SHEETS:
        if title in wb.sheetnames:
            del wb[wb.sheetnames.index(title)]  # idempotent re-runs
    _write_sheet(wb, "Clauses", sections["clauses"])
    _write_sheet(wb, "Rules", sections["rules"])
    _write_sheet(wb, "Queries (SQL)", sections["queries"])
    _write_sheet(wb, "Column Mapping", sections["mapping"])
    _write_sheet(wb, "SQL Output", sections["output"])
    _write_sheet(wb, "Review & Control", sections["routing"])
    _write_sheet(wb, "Summary", [{"metric": k, "value": v} for k, v in summary.items()])
    wb.save(out_path)


# =========================================================
# Entry point
# =========================================================

def main(argv=None):
    ap = argparse.ArgumentParser(description="Contract→clause→rule→SQL→output trace report")
    ap.add_argument("pdf", nargs="?", default=_DEFAULT_PDF, help="Contract PDF path")
    ap.add_argument("bdx", nargs="?", default=_DEFAULT_BDX, help="BDX .xlsx path")
    ap.add_argument("--out", default=None, help="Output workbook path")
    ap.add_argument("--json", dest="json_path", default=None, help="Output JSON path")
    ap.add_argument("--sheets", default=None, help="Comma-separated data sheet names (default: auto)")
    ap.add_argument("--min-cols", type=int, default=3, help="Min header columns to treat a row as a table header")
    args = ap.parse_args(argv)

    pdf, bdx = os.path.abspath(args.pdf), os.path.abspath(args.bdx)
    for p in (pdf, bdx):
        if not os.path.exists(p):
            ap.error(f"file not found: {p}")

    out_path = args.out or os.path.splitext(bdx)[0] + "_trace.xlsx"
    json_path = args.json_path or os.path.splitext(out_path)[0] + ".json"
    only_sheets = set(s.strip() for s in args.sheets.split(",")) if args.sheets else None

    # 1) BDX → Output Template + data ----------------------------------------
    print(f"[1/4] Reading BDX: {bdx}")
    template_fields, records_by_sheet, schema_cols, skipped = read_bdx(
        bdx, only_sheets=only_sheets, min_cols=args.min_cols)
    data_sheets = [r["sheet"] for r in records_by_sheet]
    print(f"      data sheets: {data_sheets}  (skipped non-tabular: {skipped})")
    print(f"      output-template columns: {sum(len(c) for c in schema_cols.values())}")

    # 2) Run the contract pipeline (LLM) -------------------------------------
    print(f"[2/4] Processing contract through the pipeline (LLM): {os.path.basename(pdf)}")
    from contract_upload_services.contract_extraction_service import ContractExtractionService
    payload = ContractExtractionService().process_contract(pdf, template_fields=template_fields)
    if isinstance(payload, dict) and payload.get("halted_for_references"):
        print("      pipeline halted for external references; nothing to report.")
        return 2
    clauses = payload.get("clauses_extracted") or []
    rules = payload.get("validation_rules") or []
    review = payload.get("review_queue") or []
    control = payload.get("control_register") or []
    print(f"      clauses={len(clauses)}  rules={len(rules)}  review={len(review)}  control={len(control)}")

    # 3) Execute each rule's SQL against the BDX -----------------------------
    print(f"[3/4] Running {len(rules)} compiled queries against the BDX (DuckDB)")
    runs = run_sql(records_by_sheet, schema_cols, rules)
    total_violations = sum(r["violation_count"] for r in runs)
    errored = [r["rule_name"] for r in runs if r["status"] == "error"]
    print(f"      total violating rows={total_violations}  query errors={len(errored)}")

    # 4) Build the workbook + JSON -------------------------------------------
    print(f"[4/4] Writing report")
    sections = {
        "clauses": clause_rows(clauses),
        "rules": rule_rows(rules),
        "queries": query_rows(rules, runs),
        "mapping": mapping_rows(rules, {f["name"]: f["sheet"] for f in template_fields}),
        "output": output_rows(runs),
        "routing": routing_rows(review, control),
    }
    meta = payload.get("metadata") or {}
    summary = {
        "contract_file": os.path.basename(pdf),
        "bdx_file": os.path.basename(bdx),
        "data_sheets": ", ".join(data_sheets),
        "skipped_sheets": ", ".join(skipped),
        "clauses_extracted": len(clauses),
        "rule_bearing": (meta.get("stage_a_summary") or {}).get("rule_bearing"),
        "rules_generated": len(rules),
        "queries_with_violations": sum(1 for r in runs if r["status"] == "violations"),
        "queries_clean": sum(1 for r in runs if r["status"] == "clean"),
        "queries_errored": len(errored),
        "total_violating_rows": total_violations,
        "review_queue": len(review),
        "control_register": len(control),
    }

    write_workbook(bdx, out_path, sections, summary)
    with open(json_path, "w") as f:
        json.dump({
            "summary": summary,
            "template_fields": template_fields,
            "clauses": clauses,
            "rules": rules,
            "review_queue": review,
            "control_register": control,
            "sql_runs": runs,
        }, f, indent=2, default=str)

    print(f"\n✓ Workbook : {out_path}")
    print(f"✓ JSON     : {json_path}")
    print(f"  {summary['rules_generated']} rules · {summary['total_violating_rows']} violating rows · "
          f"{summary['queries_errored']} query errors")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
