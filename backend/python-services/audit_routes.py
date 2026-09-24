"""Audit Logs — the API behind the screen every seat has in its sidebar.

Three endpoints, all reading through audit_feed.py, which is where the scoping
and the wording live:

  GET /audit/logs      one page of the trail, newest first
  GET /audit/options   what the Role/Actor and Action Type dropdowns may offer
  GET /audit/export    the same rows the current filters select, as CSV or xlsx

There is no role guard on any of them, and that is deliberate: every seat has
an audit trail of its own, and what differs is WHOSE rows come back. A guard
would have to repeat the scope rule in a second place and could only ever
disagree with it. `audit_feed.scope_for` is the single answer, and it is
applied to every query before a row is read — a seat with nothing in scope gets
an empty page, never someone else's.
"""
from __future__ import annotations

import csv
import io
from datetime import datetime
from typing import Optional

from fastapi import APIRouter, Depends, Query
from fastapi.responses import Response

import audit_feed
from auth_deps import Principal, current_principal
from db import SessionLocal

router = APIRouter(prefix="/audit", tags=["audit"])

DEFAULT_PAGE_SIZE = 15          # the reference design's page — 1–15 of N


def _filters(days: Optional[int], since: Optional[str], until: Optional[str],
             actor: Optional[str], action: Optional[str], q: Optional[str],
             ) -> audit_feed.Filters:
    """Turn the query string into the one filter object every read uses."""
    lo, hi = audit_feed.window(days, since, until)
    uids, bids, system = audit_feed.parse_actor((actor or "").split(","))
    group = (action or "all").strip()
    return audit_feed.Filters(
        since=lo, until=hi,
        categories=audit_feed.categories_for_group(group),
        actions=audit_feed.actions_for_group(group),
        actor_user_ids=uids, actor_broker_ids=bids, actor_system=system,
        q=(q or "").strip(),
    )


@router.get("/logs")
def audit_logs(
    page: int = 1,
    page_size: int = DEFAULT_PAGE_SIZE,
    days: Optional[int] = None,
    since: Optional[str] = Query(default=None, alias="from"),
    until: Optional[str] = Query(default=None, alias="to"),
    actor: Optional[str] = None,
    action: Optional[str] = None,
    q: Optional[str] = None,
    principal: Principal = Depends(current_principal),
):
    """One page of this seat's audit trail.

    `from` / `to` are explicit UTC instants: the browser knows the viewer's
    timezone and the server does not, so the day boundaries are computed there
    and sent as instants. `days` is the fallback for a caller that has no
    timezone to work from.
    """
    with SessionLocal() as s:
        v = audit_feed.viewer_for(s, principal)
        return audit_feed.page(s, v, _filters(days, since, until, actor, action, q),
                               page, page_size)


@router.get("/options")
def audit_options(principal: Principal = Depends(current_principal)):
    """The dropdown contents — built from what this seat can actually see, so
    no filter is offered that could only ever return nothing, and no name
    reaches a dropdown that the table itself would mask."""
    with SessionLocal() as s:
        return audit_feed.options(s, audit_feed.viewer_for(s, principal))


# --- the download ----------------------------------------------------------
# "Meaningful format" means the columns a person would keep: when, who, in what
# role, at which organisation, what they did, WHAT EXACTLY (the field, the
# policy, the old and new value), to what, and how it ended. Not the internal
# event name — that rides along in its own column for anyone matching the
# export back to the database.

_COLUMNS = [
    ("at",           "Timestamp (UTC)"),
    ("actor",        "Actor"),
    ("actor_role",   "Role"),
    ("actor_org",    "Organisation"),
    ("carrier",      "Carrier"),
    ("action_label", "Action"),
    ("detail",       "What exactly"),
    ("target",       "Target / file"),
    ("status",       "Status"),
    ("category",     "Log"),
    ("ip",           "IP address"),
    ("action",       "Event name"),
]

_CATEGORY_WORDS = {"activity": "Activity", "auth": "Sign-in", "access": "Access",
                   "decision": "Exception"}


def _cell(row: dict, key: str):
    if key == "category":
        return _CATEGORY_WORDS.get(row.get("category"), row.get("category"))
    if key == "at":
        # Excel and every spreadsheet read "2026-09-24 11:14:03" as a time;
        # they read "2026-09-24T11:14:03Z" as a string.
        text = row.get("at") or ""
        return text.replace("T", " ").replace("Z", "").split(".")[0]
    return row.get(key) or ""


def _filename(fmt: str) -> str:
    return f"kavachio-audit-logs-{datetime.utcnow():%Y%m%d-%H%M}.{fmt}"


@router.get("/export")
def audit_export(
    format: str = "csv",
    days: Optional[int] = None,
    since: Optional[str] = Query(default=None, alias="from"),
    until: Optional[str] = Query(default=None, alias="to"),
    actor: Optional[str] = None,
    action: Optional[str] = None,
    q: Optional[str] = None,
    principal: Principal = Depends(current_principal),
):
    """Every row the current filters select, capped at audit_feed.MAX_EXPORT_ROWS.

    Same scope, same filters, same wording as the screen — downloading must not
    be a second door onto rows the table would not show, and a carrier's export
    masks broker people exactly as the table does.
    """
    with SessionLocal() as s:
        v = audit_feed.viewer_for(s, principal)
        rows = audit_feed.export_rows(
            s, v, _filters(days, since, until, actor, action, q))

    if (format or "csv").lower() in ("xlsx", "excel"):
        return _xlsx(rows)
    return _csv(rows)


def _csv(rows: list[dict]) -> Response:
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\r\n")
    w.writerow([label for _k, label in _COLUMNS])
    for r in rows:
        w.writerow([_cell(r, k) for k, _label in _COLUMNS])
    # A BOM so Excel opens a UTF-8 CSV with the accents intact instead of
    # guessing the local code page and mangling every non-ASCII name.
    body = "﻿" + buf.getvalue()
    return Response(
        content=body.encode("utf-8"),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="{_filename("csv")}"'},
    )


def _xlsx(rows: list[dict]) -> Response:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Audit Logs"
    ws.append([label for _k, label in _COLUMNS])
    head = ws[1]
    for cell in head:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3A5F")
        cell.alignment = Alignment(vertical="center")
    for r in rows:
        ws.append([_cell(r, k) for k, _label in _COLUMNS])
    widths = [20, 26, 15, 26, 22, 30, 56, 30, 20, 12, 15, 26]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    out = io.BytesIO()
    wb.save(out)
    return Response(
        content=out.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{_filename("xlsx")}"'},
    )
