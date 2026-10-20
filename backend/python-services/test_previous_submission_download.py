"""Plan task 23.4 — the downloaded Excel is the bordereau alone (decided 19
Oct 2026): an output made before then loses its "Previous submission" sheet
on the way out (version_diff_store.clean_download, version_diff_sheet.
strip_sheet); and the tab's comparison is worked out once and remembered
(version_diff_store.screen). Pure: no database."""
from __future__ import annotations

import io
import re
import threading
import time
import zipfile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.packaging.custom import StringProperty

import version_diff_sheet as vsh
import version_diff_store as vds
from test_version_diff_sheet import _book, _result


def _old_output(own_sheet: bool = False, ours_at: int | None = None) -> bytes:
    """An output as stored before 19 Oct 2026: BDX + Results + the sheet
    (title in A1, a note, a filter), stamped — and a totals formula whose
    result was written in after the last save (as
    direct_render.inject_summary_formulas does)."""
    wb = Workbook()
    wb.active.title = "BDX"
    wb["BDX"].append(["Policy", "Gross"])
    wb["BDX"].append(["P1", 10])
    wb["BDX"].append(["P2", 20])
    wb["BDX"]["B4"] = "=SUM(B2:B3)"
    wb["BDX"]["A2"].comment = Comment("Failed: policy format", "Kavachio")
    wb.create_sheet("Results")
    if own_sheet:
        wb.create_sheet(vsh.SHEET_NAME)["A1"] = "The template's own notes"
    old = wb.create_sheet(vsh.SHEET_NAME + (" 2" if own_sheet else ""), ours_at)
    old["A1"], old["A3"] = vsh.TITLE, "Outcome"
    old["A4"] = "Restated"
    old["A4"].comment = Comment("Amount changed, and nothing in the file says why.", "Kavachio")
    old.auto_filter.ref = "A3:A4"
    wb.custom_doc_props.append(StringProperty(name=vsh._STAMP, value="3|compared"))
    buf = io.BytesIO()
    wb.save(buf)
    z = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zo:
        for item in z.infolist():
            data = z.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = re.sub(rb"(<f>SUM\(B2:B3\)</f>)<v\s*(?:/>|></v>)", rb"\1<v>30</v>", data)
                assert b"<v>30</v>" in data
            zo.writestr(item, data)
    return out.getvalue()


def test_an_old_outputs_sheet_is_taken_out():
    old = _old_output()
    assert vsh.has_sheet(old)
    out = vds.clean_download(old, "x.xlsx")
    wb = load_workbook(io.BytesIO(out))
    assert wb.sheetnames == ["BDX", "Results"]
    assert not vsh.has_sheet(out)
    assert vsh._STAMP not in wb.custom_doc_props.names
    assert wb["BDX"]["A2"].comment.text == "Failed: policy format"      # the BDX's own notes stay
    names = zipfile.ZipFile(io.BytesIO(out)).namelist()
    assert not any("comment" in n and n.endswith("2.xml") for n in names)  # the sheet's notes go


def test_everything_else_is_served_byte_for_byte():
    old = _old_output()
    out = vds.clean_download(old, "x.xlsx")
    zi, zo = zipfile.ZipFile(io.BytesIO(old)), zipfile.ZipFile(io.BytesIO(out))
    edited = {"xl/workbook.xml", "xl/_rels/workbook.xml.rels", "[Content_Types].xml",
              "docProps/custom.xml"}
    assert all(zo.read(n) == zi.read(n) for n in zo.namelist() if n not in edited)
    # the totals keep the result saved with them
    assert load_workbook(io.BytesIO(out), data_only=True)["BDX"]["B4"].value == 30
    assert load_workbook(io.BytesIO(out))["BDX"]["B4"].value == "=SUM(B2:B3)"


def test_a_templates_own_sheet_of_that_name_is_kept():
    out = vds.clean_download(_old_output(own_sheet=True), "x.xlsx")
    wb = load_workbook(io.BytesIO(out))
    assert wb.sheetnames == ["BDX", "Results", vsh.SHEET_NAME]
    assert wb[vsh.SHEET_NAME]["A1"].value == "The template's own notes"


def test_a_sheet_in_the_middle_leaves_the_others_names_right():
    old = _old_output(ours_at=1)
    wb = load_workbook(io.BytesIO(old))
    assert wb.sheetnames == ["BDX", vsh.SHEET_NAME, "Results"]
    wb["Results"].auto_filter.ref = "A1:B2"            # a name that belongs to the third sheet
    wb.active = 1
    buf = io.BytesIO()
    wb.save(buf)
    out = vsh.strip_sheet(buf.getvalue())
    xml = zipfile.ZipFile(io.BytesIO(out)).read("xl/workbook.xml").decode()
    assert re.findall(r'localSheetId="(\d+)"', xml) == ["1"]
    assert 'activeTab="0"' in xml
    wb = load_workbook(io.BytesIO(out))
    assert wb.sheetnames == ["BDX", "Results"]
    assert wb["Results"].auto_filter.ref == "A1:B2"


def test_files_without_it_are_served_as_stored():
    plain = _book("BDX", "Results")
    assert vds.clean_download(plain, "x.xlsx") is plain           # never had the sheet
    old = _old_output()
    assert vds.clean_download(old, "x.csv") is old                # not a workbook
    assert vds.clean_download(b"not a workbook", "x.xlsx") == b"not a workbook"
    assert vds.clean_download(b"", "x.xlsx") == b""


def test_removing_it_is_quick():
    old = _old_output()
    t = time.perf_counter()
    for _ in range(20):
        vsh.strip_sheet(old)
    assert (time.perf_counter() - t) / 20 < 0.5


# ── the comparison is worked out once ───────────────────────────────────────

@pytest.fixture
def worked(monkeypatch):
    """screen() with the database parts stood in for: `stamp` is what the
    comparison depends on; `calls` counts the comparisons worked out."""
    state = {"stamp": "a", "calls": 0}
    vds._REMEMBERED.clear()
    monkeypatch.setattr(vds, "_prepare", lambda s, **kw: {"done": None})
    monkeypatch.setattr(vds, "_depends_on", lambda s, prep: state["stamp"])

    def compare(s, prep):
        state["calls"] += 1
        time.sleep(0.05)
        return _result()
    monkeypatch.setattr(vds, "_compare", compare)
    yield state
    vds._REMEMBERED.clear()


def test_paging_and_filtering_reuse_the_comparison(worked):
    first = vds.screen(None, 7)
    again = vds.screen(None, 7)
    assert worked["calls"] == 1 and again is first
    assert first["screen"]["total"] == 5
    assert "view" not in first["result"]                # the CSV needs the rows that differ only


def test_a_change_it_depends_on_works_it_out_again(worked):
    vds.screen(None, 7)
    worked["stamp"] = "b"                   # e.g. a fix saved, or the setup's mapping edited
    vds.screen(None, 7)
    assert worked["calls"] == 2


def test_two_requests_at_once_work_it_out_once(worked):
    got = []
    threads = [threading.Thread(target=lambda: got.append(vds.screen(None, 9))) for _ in range(4)]
    for t in threads:
        t.start()
    for t in threads:
        t.join()
    assert worked["calls"] == 1 and len(got) == 4 and all(g is got[0] for g in got)


def test_memory_is_bounded(worked, monkeypatch):
    monkeypatch.setenv("PREVIOUS_SUBMISSION_CACHE_ROWS", "8")      # 5 rows a comparison
    vds.screen(None, 1)
    vds.screen(None, 2)
    assert list(vds._REMEMBERED) == [2]
    vds.screen(None, 1)
    assert worked["calls"] == 3


def test_nothing_to_compare_is_said_and_not_remembered(monkeypatch):
    vds._REMEMBERED.clear()
    import version_diff as vd
    done = vd._empty("no_accepted_version", "Nothing to compare yet.")
    monkeypatch.setattr(vds, "_prepare", lambda s, **kw: {"done": done})
    got = vds.screen(None, 3)
    assert got["screen"]["message"] == "Nothing to compare yet."
    assert not vds._REMEMBERED
