"""
The layout of a reinsurance bordereau when the carrier does not use
Kavachio's own — the reinsurance counterpart of a broker's output template.

THREE WAYS TO LAY ONE OUT
  kavachio   Kavachio's columns (ceded_engine.LINE_COLUMNS), chosen and
             renamed on the setup screen. The default.
  uploaded   The carrier already sends its reinsurers a file — an Excel or CSV
             it (or the reinsurer) designed. It is uploaded once; every
             statement is that file, its own headings and styling, filled in.
  standard   No file of its own, so one is built the way a broker's is: from
             the Lloyd's risk sheet (the bundled Coverholder Reporting
             Standards), with the reinsurance columns added after it.

WHAT A COLUMN IS FILLED FROM
  Each column carries a `key`:
    a ceded_engine line key   ceded premium, commission, recovery, … — the
                              figures this treaty works out
    "share_pct"               the reinsurer's share of the line's layer
    "field:<data-model key>"  passed through from the inward bordereau row the
                              line came from (insured address, risk code, …)
    ""                        nothing — left blank (an uploaded file's column
                              Kavachio has no data for)

  Headings are read the way a broker's sample is: the same parse and the same
  AI pass that ties each column to a data-model field. A heading that names a
  reinsurance figure ("Ceded premium", "Recovery", "RI commission") is matched
  to it first — the data model knows nothing about what is ceded.

Reading never writes; ceded_routes saves.
"""
from __future__ import annotations

import csv
import io
import re
from typing import Any, Optional

import ceded_engine as ce

FIELD = "field:"
# An output column copied from ONE input column, by that column's heading.
COL = "col:"
SHARE = ("share_pct", "Your share %", "pct")
SOURCES = ("kavachio", "uploaded", "standard")
MAX_BYTES = 5 * 1024 * 1024

# The same on every line of one statement — what a Lloyd's sheet asks for in
# every row (the reporting period, the coverholder) and Kavachio already knows.
CONTEXT = [("period", "Statement period"), ("period_start", "Period start"),
           ("period_end", "Period end"), ("treaty", "Treaty"),
           ("reinsured", "Reinsured (you)"), ("reinsurer", "Reinsurer"),
           ("currency", "Currency")]

LINE_KEYS = ({k for k, _h, _kind in ce.LINE_COLUMNS} | {SHARE[0]}
             | {k for k, _h in CONTEXT})

# A data-model field that already IS a line figure — read it from the line, so
# the value is the one the calculation used.
_CANON_TO_LINE = {
    **{f: "gross_premium" for f in ce.PREMIUM},
    **{f: "sum_insured" for f in ce.SUM_INSURED},
    **{f: "policy" for f in ce.POLICY},
    **{f: "insured" for f in ce.INSURED},
    **{f: "transaction" for f in ce.TXN},
    **{f: "claim" for f in ce.CLAIM_NO},
}

# Headings that name what is ceded. Only an explicit reinsurance word counts:
# a Lloyd's "Commission %" is the coverholder's commission, not the ceding
# commission, and must stay the inward figure.
_RI = r"(?:ced(?:ed|ing|e)|cession|re-?insur\w*|\bri\b|treaty)"
_HEADINGS: list[tuple[str, str]] = [
    (rf"{_RI}.*commission|commission.*{_RI}|overrid", "ceding_commission"),
    (rf"\bnet\b.*(?:{_RI}|\bdue\b|payable)|(?:{_RI}).*\bnet\b|balance due", "net_premium"),
    (rf"(?:{_RI}).*premium|premium.*(?:{_RI})", "ceded_premium"),
    (r"recover", "recovery"),
    (rf"(?:{_RI}).*(?:paid)|paid.*(?:{_RI})", "claims_paid"),
    (rf"(?:{_RI}).*(?:outstanding|reserve)|(?:outstanding|reserve).*(?:{_RI})",
     "claims_outstanding"),
    (rf"(?:your|reinsurer'?s?|signed|written)\s*(?:share|line)|share\s*of\s*(?:layer|treaty)",
     "share_pct"),
    (rf"(?:{_RI})\s*(?:%|pct|percent|share|rate)|%\s*(?:{_RI})|cession", "ceded_pct"),
    (r"reporting period.*start|period start", "period_start"),
    (r"reporting period.*end|period end", "period_end"),
    (r"^(?:statement|reporting) period$", "period"),
    (r"^coverholder(?:\s*name)?$", "broker"),
    (r"^treaty(?:\s*(?:name|ref|reference))?$", "treaty"),
    (r"^(?:reinsured|cedant|ceding company)(?:\s*name)?$", "reinsured"),
    (r"^reinsurer(?:\s*name)?$", "reinsurer"),
    (r"^layer|layer\s*(?:name|no|number)?$", "layer"),
    (r"^programme|^program\b|^binder\b", "programme"),
    (r"^(?:producing\s*)?broker(?:\s*name)?$", "broker"),
]


# How each figure Kavachio works out is made, and the inward fields it reads —
# what the setup's Field mapping step shows against the input columns.
_CEDED = "× the cession %, capped by the treaty's limits"
DERIVED: dict[str, tuple[str, tuple]] = {
    "programme": ("The programme the bordereau came in on", ()),
    "broker": ("The broker who sent it", ()),
    "policy": ("Copied", ce.POLICY),
    "insured": ("Copied", ce.INSURED),
    "transaction": ("Copied", ce.TXN),
    "claim": ("Copied", ce.CLAIM_NO),
    "layer": ("The treaty layer the line falls in", ()),
    "gross_premium": ("Copied", ce.PREMIUM),
    "sum_insured": ("Copied", ce.SUM_INSURED),
    "ceded_pct": ("From the treaty's terms (surplus: by sum insured)", ce.SUM_INSURED),
    "ceded_premium": (f"Gross premium {_CEDED}", ce.PREMIUM),
    "ceding_commission": ("Ceded premium × the ceding commission %", ce.PREMIUM),
    "net_premium": ("Ceded premium − ceding commission", ce.PREMIUM),
    "claims_paid": (f"Claims paid {_CEDED}", ce.CLAIM_PAID),
    "claims_outstanding": (f"Outstanding claims {_CEDED}", ce.CLAIM_RESERVE),
    "claim_incurred": ("Incurred, or paid + outstanding",
                       ce.CLAIM_INCURRED + ce.CLAIM_PAID + ce.CLAIM_RESERVE),
    "recovery": ("Each claim above the layer's retention, up to its limit",
                 ce.CLAIM_INCURRED + ce.CLAIM_PAID + ce.CLAIM_RESERVE),
    "share_pct": ("The reinsurer's share of the layer, from the treaty", ()),
    "period": ("The statement's period", ()),
    "period_start": ("First day of the statement's period", ()),
    "period_end": ("Last day of the statement's period", ()),
    "treaty": ("The treaty's name", ()),
    "reinsured": ("Your company", ()),
    "reinsurer": ("The reinsurer the statement is for", ()),
    "currency": ("Copied, else the treaty's currency", ce.CURRENCY),
}


def label(key: str) -> str:
    """How a key reads on the setup screen."""
    if key.startswith(FIELD):
        f = key[len(FIELD):]
        return f.replace("_", " ").strip().capitalize()
    if key == SHARE[0]:
        return SHARE[1]
    ctx = dict(CONTEXT)
    if key in ctx:
        return ctx[key]
    return next((h for k, h, _ in ce.LINE_COLUMNS if k == key), key or "Left blank")


def _norm(s: Any) -> str:
    return re.sub(r"[^a-z0-9%]+", " ", str(s or "").lower()).strip()


def key_for(heading: str, canonical: Optional[str]) -> str:
    """What a column of the carrier's layout is filled from."""
    h = _norm(heading)
    for pat, key in _HEADINGS:
        if re.search(pat, h):
            return key
    if canonical:
        return _CANON_TO_LINE.get(canonical) or f"{FIELD}{canonical}"
    return ""


def fillable(key: str, available: set[str], names: Optional[set] = None) -> bool:
    """Whether Kavachio has something to put in this column: a figure it works
    out, an input column, or a field the programmes' bordereaux carry."""
    if not key:
        return False
    if key.startswith(COL):
        return key[len(COL):] in (names or set())
    if key.startswith(FIELD):
        return key[len(FIELD):] in available
    return key in LINE_KEYS


def source_columns(structures: list[dict]) -> tuple[set[str], dict[str, str], dict[str, str]]:
    """What the inward bordereaux carry: their data-model fields, each of their
    own headings → field (a carrier's layout often uses the same headings as
    the files its brokers send), and every heading as it is written."""
    fields: set[str] = set()
    by_heading: dict[str, str] = {}
    names: dict[str, str] = {}
    for st in structures:
        for sh in (st or {}).get("sheets") or []:
            for col in sh.get("columns") or []:
                name = str(col.get("column_name") or "").strip()
                if name:
                    names.setdefault(_norm(name), name)
                f = col.get("canonical_field")
                if not f:
                    continue
                fields.add(f)
                by_heading.setdefault(_norm(name), f)
    return fields, by_heading, names


# ── reading the carrier's file / the standard ───────────────────────────────
def parse_upload(data: bytes, filename: str) -> dict:
    """An uploaded layout → a parsed structure (sheets → columns), unmapped."""
    name = (filename or "").lower()
    if len(data) > MAX_BYTES:
        raise ValueError("The file is larger than 5 MB — upload just the layout, "
                         "with a few sample rows at most.")
    if name.endswith(".csv"):
        text = data.decode("utf-8-sig", errors="replace")
        rows = [r for r in csv.reader(io.StringIO(text))]
        head_i = next((i for i, r in enumerate(rows) if any(c.strip() for c in r)), None)
        if head_i is None:
            raise ValueError("The CSV is empty — it needs a heading row.")
        heads = rows[head_i]
        cols = [{"column_index": i, "column_name": h.strip(),
                 "samples": [r[i] for r in rows[head_i + 1:head_i + 4] if i < len(r)]}
                for i, h in enumerate(heads) if h.strip()]
        return {"sheets": [{"sheet_name": "CSV", "header_row": head_i,
                            "data_start_row": head_i + 1, "columns": cols}]}
    if not name.endswith((".xlsx", ".xlsm")):
        raise ValueError("Upload an Excel (.xlsx) or CSV file.")
    from exporter import parse_template
    try:
        st = parse_template(data, filename=filename)
    except Exception as e:  # noqa: BLE001
        raise ValueError(f"The workbook could not be read ({e}).")
    if not any(sh.get("columns") for sh in st.get("sheets") or []):
        raise ValueError("No heading row was found in the workbook.")
    return st


def map_structure(structure: dict) -> None:
    """The AI pass a broker's sample gets — each column tied to a data-model
    field. Mutates `structure`; quietly does nothing without a model."""
    try:
        from exporter import propose_template_mapping
        propose_template_mapping(structure)
    except Exception:  # noqa: BLE001 — the heading rules and manual picks still work
        pass


def data_sheet(structure: dict) -> dict:
    """The sheet the lines go into — the one with the most columns, skipping
    reference/lookup tabs."""
    sheets = [sh for sh in structure.get("sheets") or [] if sh.get("columns")]
    data = [sh for sh in sheets if (sh.get("role") or "data") == "data"] or sheets
    return max(data, key=lambda sh: len(sh.get("columns") or []))


def propose(structure: dict, *, source: str, available: set[str],
            by_heading: dict[str, str], treaty_type: Optional[str],
            names: Optional[dict[str, str]] = None) -> tuple[dict, list[dict]]:
    """(layout, columns) for the setup screen to review.

    uploaded — every column of the file is kept (it is the reinsurer's format);
               the ones Kavachio cannot fill are left blank and flagged.
    standard — a mandatory Lloyd's column, and one that can be filled, is
               ticked; the rest are offered unticked. The treaty's reinsurance
               columns follow."""
    sh = data_sheet(structure)
    cols: list[dict] = []
    for col in sh.get("columns") or []:
        heading = str(col.get("column_name") or "").strip()
        if not heading:
            continue
        canonical = by_heading.get(_norm(heading)) or col.get("canonical_field")
        key = key_for(heading, canonical)
        # An input column of the same heading is copied as itself — unless
        # the heading names a reinsurance figure, which is worked out.
        same = (names or {}).get(_norm(heading))
        if same and not key_for(heading, None):
            key = f"{COL}{same}"
        ok = fillable(key, available, set((names or {}).values()))
        required = bool(col.get("required"))
        cols.append({
            "heading": heading, "key": key,
            "column_index": int(col.get("column_index") or 0),
            "keep": True if source == "uploaded" else (required or ok),
            "required": required, "ref": col.get("standard_ref"),
        })
    if source == "standard":
        have = {c["key"] for c in cols}
        for c in ce.default_columns(treaty_type) + [{"key": SHARE[0], "heading": SHARE[1]}]:
            if c["key"] in have or c["key"] in ("programme", "broker", "policy",
                                                "insured", "transaction"):
                continue
            cols.append({"heading": c["heading"], "key": c["key"], "column_index": None,
                         "keep": True, "required": False, "ref": None, "added": True})
    layout = {"source": source, "sheet_name": sh.get("sheet_name"),
              "header_row": int(sh.get("header_row") or 0),
              "data_start_row": int(sh.get("data_start_row") or 1)}
    return layout, cols


# ── the carrier's own INPUT file ────────────────────────────────────────────
def input_proposal(structure: dict, filename: str) -> dict:
    """A sample of the carrier's own input file → the setup's input_layout:
    every column, and the data-model field the AI pass tied it to (the user
    corrects any of them on the setup screen)."""
    sh = data_sheet(structure)
    cols = [{"name": str(c.get("column_name") or "").strip(),
             "column_index": int(c.get("column_index") or 0),
             "field": c.get("canonical_field") or None}
            for c in sh.get("columns") or [] if str(c.get("column_name") or "").strip()]
    return {"source": "own", "filename": filename,
            "sheet_name": sh.get("sheet_name"),
            "header_row": int(sh.get("header_row") or 0),
            "data_start_row": int(sh.get("data_start_row") or 1),
            "columns": cols}


def read_input(layout: dict, data: bytes, filename: str) -> list[dict]:
    """One period's own file → rows keyed by data-model field, read with the
    columns the setup mapped. A sheet named differently from the sample falls
    back to the first sheet; headings are matched by name first, so a column
    that moved still lands in the right field."""
    cols = [c for c in layout.get("columns") or [] if c.get("field")]
    name = (filename or "").lower()
    if name.endswith(".csv"):
        text = data.decode("utf-8-sig", errors="replace")
        table = [r for r in csv.reader(io.StringIO(text))]
    else:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb[layout["sheet_name"]] if layout.get("sheet_name") in wb.sheetnames \
            else wb.worksheets[0]
        table = [list(r) for r in ws.iter_rows(values_only=True)]
        wb.close()
    head_i = int(layout.get("header_row") or 0)
    heads = [_norm(h) for h in (table[head_i] if head_i < len(table) else [])]
    mapped = []
    for c in layout.get("columns") or []:
        n = _norm(c.get("name"))
        idx = heads.index(n) if n in heads else int(c.get("column_index") or 0)
        mapped.append({"column_index": idx, "canonical_field": c.get("field") or None,
                       "column_name": c.get("name")})
    return ce._rows_from_table(table, {"data_start_row": layout.get("data_start_row") or head_i + 1,
                                        "columns": mapped})


# ── filling it ──────────────────────────────────────────────────────────────
def needed_fields(columns: Optional[list[dict]]) -> set[str]:
    return {c["key"][len(FIELD):] for c in columns or []
            if str(c.get("key") or "").startswith(FIELD) and c.get("keep", True)}


def needed_cols(columns: Optional[list[dict]]) -> set[str]:
    return {c["key"][len(COL):] for c in columns or []
            if str(c.get("key") or "").startswith(COL) and c.get("keep", True)}


def value(line: dict, key: str, ctx: Optional[dict] = None) -> Any:
    if not key:
        return None
    if key.startswith(FIELD):
        return (line.get("src") or {}).get(key[len(FIELD):])
    if key.startswith(COL):
        return ((line.get("src") or {}).get("_cols") or {}).get(key[len(COL):])
    if ctx and key in ctx and line.get(key) in (None, ""):
        return ctx[key]
    return line.get(key)


def context(*, period: str, treaty: str, reinsured: Optional[str],
            reinsurer: Optional[str], currency: Optional[str]) -> dict:
    b = ce.period_bounds(period)
    return {"period": period, "period_start": b[0] if b else None,
            "period_end": b[1] if b else None, "treaty": treaty,
            "reinsured": reinsured, "reinsurer": reinsurer, "currency": currency}


def prune(lines: list[dict], fields: set[str], cols: Optional[set[str]] = None) -> None:
    """Keep only the inward fields and columns a layout passes through — a
    line otherwise carries every field and column of its source row."""
    for ln in lines:
        src = ln.pop("src", None) or {}
        kept = {f: _plain(src[f]) for f in fields or () if src.get(f) not in (None, "")}
        by_name = src.get("_cols") or {}
        keep_cols = {n: _plain(by_name[n]) for n in cols or () if by_name.get(n) not in (None, "")}
        if keep_cols:
            kept["_cols"] = keep_cols
        if kept:
            ln["src"] = kept


def _plain(v: Any) -> Any:
    """JSON-safe — the lines are stored."""
    if hasattr(v, "isoformat"):
        return v.isoformat()[:10] if not getattr(v, "hour", 0) else v.isoformat()
    return v


def kept(columns: list[dict]) -> list[dict]:
    return [c for c in columns if c.get("keep", True)]


def fill_csv(columns: list[dict], lines: list[dict], ctx: Optional[dict] = None) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    cols = kept(columns)
    w.writerow([c["heading"] for c in cols])
    for ln in lines:
        w.writerow(["" if value(ln, c["key"], ctx) is None else value(ln, c["key"], ctx)
                    for c in cols])
    return buf.getvalue().encode("utf-8-sig")


def fill_xlsx(*, blob: Optional[bytes], layout: dict, columns: list[dict],
              lines: list[dict], summary: dict, ctx: Optional[dict] = None) -> bytes:
    """The carrier's own workbook with the lines written in, sample rows
    cleared — its sheets, headings and styling untouched — plus a summary sheet.

    A Lloyd's layout keeps only the ticked columns, so it is written fresh
    under the standard's sheet name, headings styled like the standard's."""
    import openpyxl
    from copy import copy

    if layout.get("source") == "uploaded" and blob:
        wb = openpyxl.load_workbook(io.BytesIO(blob))
        ws = wb[layout["sheet_name"]] if layout.get("sheet_name") in wb.sheetnames \
            else wb.worksheets[0]
        start = int(layout.get("data_start_row") or 1) + 1      # 1-based
        style_row = [copy(c._style) for c in ws[start]] if ws.max_row >= start else []
        if ws.max_row >= start:
            ws.delete_rows(start, ws.max_row - start + 1)
        for r, ln in enumerate(lines, start=start):
            for c in columns:
                if c.get("column_index") is None or not c.get("keep", True):
                    continue
                cell = ws.cell(row=r, column=int(c["column_index"]) + 1,
                               value=value(ln, c["key"], ctx))
                i = int(c["column_index"])
                if i < len(style_row):
                    cell._style = copy(style_row[i])
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = (layout.get("sheet_name") or "Lines")[:31]
        head_style = None
        if blob:
            try:
                src = openpyxl.load_workbook(io.BytesIO(blob)).worksheets[0]
                head_style = copy(src.cell(row=int(layout.get("header_row") or 0) + 1,
                                           column=1)._style)
            except Exception:  # noqa: BLE001
                head_style = None
        cols = kept(columns)
        ws.append([c["heading"] for c in cols])
        from openpyxl.styles import Alignment, Font, PatternFill
        for cell in ws[1]:
            if head_style is not None:
                cell._style = copy(head_style)
            else:
                cell.font = Font(bold=True)
                cell.fill = PatternFill("solid", fgColor="E1F1F3")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for ln in lines:
            ws.append([value(ln, c["key"], ctx) for c in cols])
        for i in range(1, len(cols) + 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = 18

    # The figures behind the file, on a sheet of their own after the carrier's.
    sm = wb.create_sheet(_free_name(wb, "Reinsurance summary"))
    for k, v in summary.get("rows") or []:
        sm.append([k, v])
    sm.column_dimensions["A"].width = 28
    sm.column_dimensions["B"].width = 22
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _free_name(wb, name: str) -> str:
    n, i = name, 2
    while n in wb.sheetnames:
        n, i = f"{name} {i}", i + 1
    return n
